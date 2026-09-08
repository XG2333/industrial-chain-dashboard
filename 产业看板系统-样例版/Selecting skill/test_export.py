from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook

from financial_variable_curation.classification.service import BatchClassificationService
from financial_variable_curation.classification.settings import LLMSettings
from financial_variable_curation.database.manager import DatabaseManager
from financial_variable_curation.database.repositories import ExportRunRepository
from financial_variable_curation.database.services import InspectionPersistenceService
from financial_variable_curation.export.orchestrator import FullPipelineOrchestrator
from financial_variable_curation.export.service import SelectionExportError, SelectionExportService
from financial_variable_curation.models.artifacts import RequestRecord
from financial_variable_curation.pipeline.inspect import run_inspection
from financial_variable_curation.rule_workflow.service import RuleParserService
from financial_variable_curation.selection.service import VariableSelectionService


def _persist_source(db_manager: DatabaseManager, input_path: Path, run_id: str) -> None:
    result = run_inspection(input_path, run_id=run_id)
    InspectionPersistenceService(db_manager).persist(
        result,
        request_record=RequestRecord(
            run_id=run_id,
            input_path=str(input_path),
            file_hash=result.workbook_profile.file_hash,
            cli_args={},
        ),
        artifacts_dir=Path("artifacts"),
        run_id=run_id,
    )


def _build_selection_run(
    db_manager: DatabaseManager,
    input_path: Path,
    tmp_path: Path,
) -> str:
    run_id = "run-export"
    _persist_source(db_manager, input_path, run_id)
    BatchClassificationService(
        db_manager, settings=LLMSettings(llm_provider="mock")
    ).classify_run(run_id, provider="mock", limit=20, artifacts_dir=tmp_path / "class")
    rules = tmp_path / "rules.txt"
    rules.write_text(
        "价格 > 库存 > 产量\n缺失率超过40%的变量删除\n同一指标只保留一个\n最终最多保留3个变量\n",
        encoding="utf-8",
    )
    RuleParserService(
        db_manager, settings=LLMSettings(llm_provider="mock")
    ).compile_rules(
        rules,
        provider="mock",
        name="export_rules",
        version="1",
        artifacts_dir=tmp_path / "rules_art",
    )
    summary = VariableSelectionService(db_manager).select(
        run_id,
        rule_set_name="export_rules",
        rule_version="1",
        allow_validated_rules=True,
        artifacts_dir=tmp_path / "selection",
    )
    return summary.selection_run_id


def _make_workbook(path: Path) -> None:
    from datetime import date, timedelta

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "价格"
    ws.append(["date", "沪锡主力收盘价", "上期所锡库存", "高缺失率变量"])
    for i in range(5):
        ws.append([date(2026, 1, 1) + timedelta(days=i), 100 + i, 10 + i, i + 1 if i == 0 else None])
    wb.save(path)


def test_export_service_generates_workbook_and_manifest(
    db_manager: DatabaseManager,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    _make_workbook(source)
    selection_run_id = _build_selection_run(db_manager, source, tmp_path)
    output = tmp_path / "selected.xlsx"
    summary = SelectionExportService(db_manager).export(
        selection_run_id,
        output_path=output,
        overwrite=True,
        artifacts_dir=tmp_path / "export_art",
    )
    assert summary.status == "COMPLETED"
    assert output.exists()
    wb = load_workbook(output, read_only=True)
    assert "selected_data__价格" in wb.sheetnames
    assert "selected_variables" in wb.sheetnames
    assert "rejected_variables" in wb.sheetnames
    assert "needs_review" in wb.sheetnames
    assert "run_summary" in wb.sheetnames
    wb.close()
    with db_manager.session_scope() as session:
        assert ExportRunRepository(session).get(summary.export_run_id) is not None


def test_export_detects_source_file_changed(
    db_manager: DatabaseManager,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    _make_workbook(source)
    selection_run_id = _build_selection_run(db_manager, source, tmp_path)
    with source.open("rb") as handle:
        original = handle.read()
    source.write_bytes(original + b"changed")
    with pytest.raises(SelectionExportError, match="SOURCE_FILE_CHANGED"):
        SelectionExportService(db_manager).export(
            selection_run_id,
            output_path=tmp_path / "selected.xlsx",
            overwrite=True,
        )


def test_pipeline_orchestrator_dry_run_does_not_create_output(
    db_manager: DatabaseManager,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    _make_workbook(source)
    rules = tmp_path / "rules.txt"
    rules.write_text("价格 > 库存 > 产量\n", encoding="utf-8")
    output = tmp_path / "out.xlsx"
    summary = FullPipelineOrchestrator(db_manager).run(
        source,
        rules_path=rules,
        provider="mock",
        output_path=output,
        dry_run=True,
        artifacts_dir=tmp_path / "artifacts",
    )
    assert summary.status == "DRY_RUN"
    assert not output.exists()


def test_pipeline_orchestrator_full_mock_workflow(
    db_manager: DatabaseManager,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    _make_workbook(source)
    rules = tmp_path / "rules.txt"
    rules.write_text(
        "价格 > 库存 > 产量\n缺失率超过40%的变量删除\n同一指标只保留一个\n最终最多保留3个变量\n",
        encoding="utf-8",
    )
    output = tmp_path / "pipeline_output.xlsx"
    summary = FullPipelineOrchestrator(db_manager).run(
        source,
        rules_path=rules,
        provider="mock",
        output_path=output,
        overwrite=True,
        artifacts_dir=tmp_path / "artifacts",
    )
    assert summary.status in {"COMPLETED", "REVIEW_REQUIRED"}
    assert output.exists()
    assert summary.output_file_hash is not None
    digest = hashlib.sha256()
    with output.open("rb") as handle:
        digest.update(handle.read())
    assert digest.hexdigest() == summary.output_file_hash
