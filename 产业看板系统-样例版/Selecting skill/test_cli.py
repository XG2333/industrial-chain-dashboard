from __future__ import annotations

import json

from openpyxl import Workbook

from financial_variable_curation.cli import main


def test_inspect_cli_creates_artifacts(sample_workbook, tmp_path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    code = main(["inspect", "--input", str(sample_workbook), "--artifacts-dir", "artifacts"])
    assert code == 0
    artifact_dirs = list((tmp_path / "artifacts").iterdir())
    assert artifact_dirs
    run_dir = artifact_dirs[0]
    assert (run_dir / "workbook_profile.json").exists()
    assert (run_dir / "variable_profiles.json").exists()
    assert (run_dir / "inspection_summary.json").exists()
    assert (run_dir / "request.json").exists()
    assert (run_dir / "run.log").exists()
    assert (run_dir / "data_quality_report.json").exists()
    assert not (run_dir / "variable_classification.json").exists()


def test_inspect_cli_supports_sheet_filter(sample_workbook, tmp_path, monkeypatch) -> None:
    workbook = Workbook()
    workbook.active.title = "A"
    workbook.active.append(["date", "value"])
    from datetime import date

    workbook.active.append([date(2024, 1, 1), 1])
    workbook.active.append([date(2024, 1, 2), 2])
    workbook.active.append([date(2024, 1, 3), 3])
    workbook.create_sheet("B").append(["date", "other"])
    path = tmp_path / "sheets.xlsx"
    workbook.save(path)
    monkeypatch.chdir(tmp_path)

    code = main(
        [
            "inspect",
            "--input",
            str(path),
            "--sheet",
            "A",
            "--artifacts-dir",
            "artifacts",
        ]
    )
    assert code == 0
    run_dir = next((tmp_path / "artifacts").iterdir())
    summary = json.loads((run_dir / "inspection_summary.json").read_text(encoding="utf-8"))
    assert summary["sheets_processed"] == 1
    request = json.loads((run_dir / "request.json").read_text(encoding="utf-8"))
    assert request["run_id"]
    assert request["cli_args"]["sheet"] == "A"
    assert request["software_version"]


def test_inspect_cli_supports_header_row_and_date_column(sample_workbook, tmp_path, monkeypatch) -> None:
    from datetime import date

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Title row"])
    sheet.append([])
    sheet.append(["date", "value"])
    sheet.append([date(2024, 1, 1), 1])
    sheet.append([date(2024, 1, 2), 2])
    sheet.append([date(2024, 1, 3), 3])
    path = tmp_path / "header_args.xlsx"
    workbook.save(path)
    monkeypatch.chdir(tmp_path)

    code = main(
        [
            "inspect",
            "--input",
            str(path),
            "--header-row",
            "3",
            "--date-column",
            "date",
            "--artifacts-dir",
            "artifacts",
        ]
    )
    assert code == 0
    run_dir = next((tmp_path / "artifacts").iterdir())
    summary = json.loads((run_dir / "inspection_summary.json").read_text(encoding="utf-8"))
    assert summary["status"] == "COMPLETED"


def test_default_run_cli_creates_output_and_audit(sample_workbook, tmp_path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    output = tmp_path / "out.xlsx"
    code = main(
        [
            "run",
            "--input",
            str(sample_workbook),
            "--use-default-rules",
            "--output",
            str(output),
            "--artifacts-dir",
            "artifacts",
        ]
    )
    assert code == 0
    assert output.exists()
    run_dir = next((tmp_path / "artifacts").iterdir())
    audit = json.loads((run_dir / "selection_report.json").read_text(encoding="utf-8"))
    assert audit["rule_source"] == "DEFAULT_PLACEHOLDER"
    assert audit["production_ready"] is False


def test_user_rules_cli_creates_parsed_and_compiled_artifacts(
    sample_workbook, sample_rules_text, tmp_path, monkeypatch
) -> None:
    monkeypatch.chdir(tmp_path)
    output = tmp_path / "user.xlsx"
    code = main(
        [
            "run",
            "--input",
            str(sample_workbook),
            "--rules",
            str(sample_rules_text),
            "--output",
            str(output),
            "--artifacts-dir",
            "artifacts",
        ]
    )
    assert code == 0
    assert output.exists()
    run_dir = next((tmp_path / "artifacts").iterdir())
    assert (run_dir / "parsed_rules.json").exists()
    assert (run_dir / "compiled_rules.json").exists()
    parsed = json.loads((run_dir / "parsed_rules.json").read_text(encoding="utf-8"))
    assert parsed["status"] == "VALIDATED"


def test_empty_user_rules_are_not_executed(sample_workbook, tmp_path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    empty_rules = tmp_path / "empty_rules.txt"
    empty_rules.write_text("", encoding="utf-8")
    code = main(
        [
            "run",
            "--input",
            str(sample_workbook),
            "--rules",
            str(empty_rules),
            "--output",
            str(tmp_path / "empty.xlsx"),
            "--artifacts-dir",
            "artifacts",
        ]
    )
    assert code == 2
    assert not (tmp_path / "empty.xlsx").exists()
