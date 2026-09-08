from __future__ import annotations

import os
import uuid
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_HEADERS = [
    "序号", "指标名称", "所属Sheet", "单位", "频率",
    "大类", "子类", "数据性质", "置信度", "是否选中", "状态说明",
]
NEW_HEADERS = [
    "大类", "子类", "数据性质", "置信度", "是否选中", "状态说明",
]
NAME_ALIASES = {"指标名称", "指标", "名称", "variable", "indicator", "name"}
SHEET_ALIASES = {"所属sheet", "sheet", "所属表", "工作表", "所在sheet", "表"}

CATEGORY_CN = {
    "PRICE": "价格", "QUANTITY": "数量", "INVENTORY": "库存",
    "CAPACITY": "产能", "COST": "成本", "MARGIN": "利润",
    "FINANCIAL": "财务", "FUNDAMENTAL": "基本面", "MACRO": "宏观",
    "INDEX": "指数", "UNKNOWN": "未知", "": "",
}

NATURE_CN = {
    "LEVEL": "水平值", "DIFFERENCE": "差值", "RATIO": "比率",
    "PERCENT_CHANGE": "变化率", "INDEX": "指数", "STOCK": "存量",
    "FLOW": "流量", "AGGREGATE": "汇总", "UNKNOWN": "未知", "": "",
}

FREQ_CN = {
    "daily": "日", "weekly": "周", "monthly": "月",
    "quarterly": "季", "annual": "年", "irregular": "不规则", "unknown": "未知",
}

HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")


def _selection_label(status: str | None) -> str:
    if status == "SELECTED":
        return "是"
    if status == "NEEDS_REVIEW":
        return "复核"
    if status == "REJECTED":
        return "否"
    if status == "DUPLICATE":
        return "去重"
    if status == "FAILED":
        return "失败"
    return "否"


def _fill_for_status(status: str | None) -> PatternFill | None:
    if status == "SELECTED":
        return GREEN_FILL
    if status == "NEEDS_REVIEW":
        return YELLOW_FILL
    if status in ("REJECTED", "DUPLICATE"):
        return RED_FILL
    return None


def _norm_header(value: object) -> str:
    return str(value or "").strip().lower().replace(" ", "")


def _sheet_link(sheet_name: str) -> str:
    return f"'{str(sheet_name).replace(chr(39), chr(39)*2)}'!A1"

CATEGORY2_CN = {
    "Price": "价格", "Spread": "价差", "Cost": "成本", "Margin": "利润",
    "Quantity": "数量", "Inventory": "库存", "Capacity": "产能利用",
    "Financial": "财务", "Fundamental": "基本面", "Macro": "宏观",
    "Index": "指数", "Unknown": "未知", "": "",
    "FUTURES_PRICE": "期货价格", "COST_PROFIT": "成本利润",
    "SUPPLY": "供应", "DEMAND": "需求", "IMPORT": "进口", "EXPORT": "出口",
    "POSITION": "持仓", "VOLUME": "成交量",
    "BASIS": "基差", "BALANCE": "平衡",
}

REASON_TR = {
    "Selected by deterministic rules.": "确定性规则选中",
    "Duplicate in comparison group": "对比组去重",
    "CLASSIFICATION_NEEDS_REVIEW": "分类需复核",
    "UNKNOWN_CATEGORY": "未知大类",
    "UNRESOLVED_REVIEW": "未解决",
    "LOW_CONFIDENCE": "低置信度",
    "COMPARISON_GROUP_MISSING": "对比组缺失",
    "UNCLASSIFIED": "未分类",
    "No blocking rule": "无阻塞规则",
    "Hard filter matched:": "硬过滤命中：",
    "Exceeded": "超出限制",
    "Not classified yet.": "尚未分类",
    "Deterministic": "确定性的",
}

def _translate_reason(text: str) -> str:
    for eng, cn in REASON_TR.items():
        if eng in text:
            text = text.replace(eng, cn)
    return text


def _infer_category_cn(entry: dict) -> str:
    cat = str(entry.get("category_level_1") or "")
    name = str(entry.get("original_name") or entry.get("column_name") or "")
    subcat = str(entry.get("category_level_2") or "")
    if "成交持仓" in name:
        return "价格"
    if "成交量" in name:
        return "其他"
    if "盈亏" in name:
        return "成本利润"
    if "交易者数量" in name or subcat == "POSITION":
        return "价格"
    if "销量" in name:
        return "需求"
    if cat in {"", "UNKNOWN"}:
        if "产能" in name or "开工" in name or "产量" in name:
            return "供应"
        if "平衡" in name:
            return "平衡"
        if "需求" in name or "消费" in name:
            return "需求"
        if "进出口" in name or "进口" in name or "出口" in name:
            return "进出口"
        if "库存" in name:
            return "库存"
    if cat == "PRICE":
        return "价格"
    if cat in {"COST", "MARGIN"}:
        return "成本利润"
    if cat == "INVENTORY":
        return "库存"
    if cat == "CAPACITY":
        return "供应"
    if cat == "QUANTITY":
        if "储量" in name or "reserve" in name.lower():
            return "供应"
        if "平衡" in name:
            return "平衡"
        if "需求" in name or "消费" in name:
            return "需求"
        if "进出口" in name or "进口" in name or "出口" in name:
            return "进出口"
        return "供应"
    if cat == "FUNDAMENTAL" and ("储量" in name or "reserve" in name.lower()):
        return "供应"
    return "其他"


def _infer_subcategory_cn(entry: dict) -> str:
    cat = str(entry.get("category_level_1") or "")
    name = str(entry.get("original_name") or entry.get("column_name") or "")
    lowered = name.lower()
    subcat = str(entry.get("category_level_2") or "")
    if "成交持仓" in name:
        return "期货价格"
    if "成交量" in name:
        return "成交量"
    if "现货升贴水" in name:
        return "基差"
    if "期现价差" in name or "期现" in name:
        return "基差"
    if "交易者数量" in name or subcat == "POSITION":
        return "持仓"
    if "销量" in name:
        return "销量"
    if "盈亏" in name:
        return "利润"
    if cat in {"", "UNKNOWN"}:
        if "产能" in name:
            return "产能"
        if "开工" in name:
            return "开工率"
        if "产量" in name:
            return "产量"
        if "平衡" in name:
            return "平衡"
        if "需求" in name or "消费" in name:
            return "需求"
        if "库存指数" in name:
            return "库存指数"
        if "仓单" in name:
            return "仓单"
        if "库存天数" in name or "天数" in name:
            return "库存天数"
    if cat == "PRICE":
        if "基差" in name:
            return "基差"
        if "月差" in name:
            return "月差"
        if "价差" in name or "升贴水" in name or "premium" in lowered or "spread" in lowered:
            return "现货价差"
        if "期货" in name or "合约" in name or "收盘" in name:
            return "期货价格"
        if "现货" in name:
            return "现货价格"
        return "价格"
    if cat in {"COST", "MARGIN"}:
        if "利润" in name or "毛利" in name or "盈利" in name:
            return "利润"
        if "成本" in name or "加工费" in name:
            return "成本"
        return "成本利润"
    if cat == "INVENTORY":
        if "仓单" in name:
            return "仓单"
        if "库存天数" in name or "天数" in name:
            return "库存天数"
        if "库存指数" in name or "指数" in name:
            return "库存指数"
        return "库存"
    if cat == "QUANTITY":
        if "储量" in name or "reserve" in lowered:
            return "储量"
        if "产能" in name:
            return "产能"
        if "开工率" in name or "开工" in name:
            return "开工率"
        if "产量" in name or "output" in lowered:
            return "产量"
        if "平衡" in name:
            return "平衡"
        if "需求" in name or "消费" in name:
            return "需求"
        if "进口" in name:
            return "进口"
        if "出口" in name:
            return "出口"
        if "成交" in name:
            return "成交量"
        if "持仓" in name:
            return "持仓"
        return "数量"
    if cat == "FUNDAMENTAL" and ("储量" in name or "reserve" in lowered):
        return "储量"
    if cat == "CAPACITY":
        if "开工" in name:
            return "开工率"
        return "产能"
    if cat == "MACRO":
        return "宏观"
    if cat == "INDEX":
        return "指数"
    return "其他"



def _read_existing(ws) -> tuple[list[str], list[dict], dict[tuple[int, int], object]]:
    """Read the first sheet without clearing it."""
    headers: list[str] = []
    data: list[dict] = []
    hyperlinks: dict[tuple[int, int], object] = {}
    if ws.max_row and ws.max_column:
        headers = [
            str(cell.value).strip() if cell.value is not None else ""
            for cell in ws[1]
        ]
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            record: dict[str, object] = {}
            has_value = False
            for idx, cell in enumerate(row):
                if cell.hyperlink is not None:
                    hyperlinks[(cell.row, cell.column)] = cell.hyperlink
                if idx < len(headers) and headers[idx]:
                    record[headers[idx]] = cell.value
                if cell.value is not None and str(cell.value).strip() != "":
                    has_value = True
            if has_value:
                record["_row"] = row[0].row
                data.append(record)
    return headers, data, hyperlinks


def _find_col(
    headers: list[str],
    aliases: set[str],
    substring_aliases: set[str] | None = None,
) -> int:
    substring_aliases = substring_aliases or aliases
    # Exact match first.
    for idx, header in enumerate(headers):
        if header and _norm_header(header) in aliases:
            return idx
    # Substring match using explicit aliases to avoid false positives.
    for idx, header in enumerate(headers):
        h = _norm_header(header)
        if not h:
            continue
        for alias in substring_aliases:
            if not alias:
                continue
            if alias in h:
                return idx
    return -1


def _cell_hyperlink(cell, sheet_name: str) -> None:
    if not sheet_name:
        return
    cell.hyperlink = Hyperlink(
        ref=cell.coordinate,
        location=_sheet_link(sheet_name),
    )
    cell.font = Font(color="0563C1", underline="single")


def update_directory_sheet(
    source_path: str | Path,
    output_path: str | Path,
    rows: list[dict],
) -> Path:
    source_path = Path(source_path)
    output_path = Path(output_path)
    if output_path.resolve() == source_path.resolve():
        raise ValueError("Output path must not equal the source path.")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Load directly from source. Never load and save the same path on Windows.
    wb = load_workbook(source_path)
    ws = wb[wb.sheetnames[0]]

    existing_headers, existing_data, existing_links = _read_existing(ws)
    has_original = any(existing_headers)

    if has_original:
        original_headers = [h for h in existing_headers if h]
        new_headers = [h for h in NEW_HEADERS if h not in original_headers]
        final_headers = original_headers + new_headers
    else:
        original_headers = DEFAULT_HEADERS[:5]
        new_headers = DEFAULT_HEADERS[5:]
        final_headers = DEFAULT_HEADERS

    name_col = _find_col(
        final_headers,
        NAME_ALIASES,
        substring_aliases={"indicator", "variable", "指标", "指标名称"},
    )
    sheet_col = _find_col(
        final_headers,
        SHEET_ALIASES,
        substring_aliases={"sheet", "所属", "工作表"},
    )
    if name_col < 0 and len(final_headers) > 1:
        name_col = 1

    # Build lookup by indicator and by (indicator, sheet).
    by_name: dict[str, dict] = {}
    by_pair: dict[tuple[str, str], dict] = {}
    for entry in rows:
        name = str(entry.get("original_name") or entry.get("column_name") or "")
        sheet = str(entry.get("sheet_name") or "")
        by_name[name] = entry
        by_pair[(name, sheet)] = entry

    header_font = Font(bold=True, size=11)
    for col_idx, header in enumerate(final_headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = HEADER_FILL

    def _append_row(entry: dict, row_idx: int, *, preserve: dict[str, object] | None = None) -> None:
        preserve = preserve or {}
        status = entry.get("final_status") or "UNCLASSIFIED"
        freq_raw = str(entry.get("detected_frequency") or "")
        values: dict[str, object] = {
            "序号": row_idx - 1,
            "指标名称": entry.get("original_name") or entry.get("column_name") or "",
            "所属Sheet": entry.get("sheet_name") or "",
            "单位": entry.get("unit_hint") or "",
            "频率": FREQ_CN.get(freq_raw.lower(), freq_raw),
            "大类": _infer_category_cn(entry),
            "子类": _infer_subcategory_cn(entry),
            "数据性质": NATURE_CN.get(entry.get("data_nature") or "", entry.get("data_nature") or ""),
            "置信度": round(entry.get("confidence") or 0, 2) if entry.get("confidence") is not None else "",
            "是否选中": _selection_label(status),
            "状态说明": _translate_reason(entry.get("reason_text") or ""),
        }
        orig_row = int(preserve.get("_row", 0) or 0) if preserve else 0
        for col_idx, header in enumerate(final_headers, 1):
            value = values.get(header)
            if preserve and header in preserve:
                value = preserve[header]
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="center")
            if col_idx == name_col + 1:
                existing_link = existing_links.get((orig_row, col_idx)) if orig_row else None
                if existing_link is not None:
                    cell.hyperlink = existing_link
                else:
                    sheet_name = entry.get("sheet_name") or (
                        preserve.get("Sheet Name") or preserve.get("所属Sheet") or ""
                    )
                    _cell_hyperlink(cell, str(sheet_name or ""))
        fill = _fill_for_status(status)
        if fill:
            for col_idx in range(1, len(final_headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill

    if has_original:
        # Preserve original rows and order, then append new indicators.
        row_idx = 2
        for record in existing_data:
            name = str(record.get(final_headers[name_col]) or "") if name_col >= 0 else ""
            sheet = str(record.get(final_headers[sheet_col]) or "") if sheet_col >= 0 else ""
            entry = by_pair.get((name, sheet)) or by_name.get(name)
            if entry is None:
                entry = {
                    "original_name": name,
                    "sheet_name": sheet,
                    "final_status": "UNCLASSIFIED",
                }
            _append_row(entry, row_idx, preserve=record)
            row_idx += 1
        for entry in rows:
            name = str(entry.get("original_name") or entry.get("column_name") or "")
            sheet = str(entry.get("sheet_name") or "")
            if name and not any(
                str(r.get(final_headers[name_col]) or "") == name
                and (
                    sheet_col < 0
                    or not sheet
                    or not str(r.get(final_headers[sheet_col]) or "").strip()
                    or str(r.get(final_headers[sheet_col]) or "").strip() == sheet
                )
                for r in existing_data
            ):
                _append_row(entry, row_idx)
                row_idx += 1
    else:
        for row_idx, entry in enumerate(rows, 2):
            _append_row(entry, row_idx)

    # Preserve existing hyperlinks that were not overwritten.
    for (row, col), link in existing_links.items():
        if row <= ws.max_row and col <= len(final_headers):
            cell = ws.cell(row=row, column=col)
            if cell.hyperlink is None:
                cell.hyperlink = link

    col_widths = [6, 55, 30, 10, 10, 14, 14, 14, 8, 8, 40]
    for i, w in enumerate(col_widths, 1):
        if i <= len(final_headers):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(final_headers))}{max(ws.max_row, 2)}"

    tmp_path = output_path.with_name(
        f"{output_path.stem}__tmp_{uuid.uuid4().hex[:8]}{output_path.suffix}"
    )
    wb.save(tmp_path)
    wb.close()
    try:
        os.replace(tmp_path, output_path)
    except PermissionError:
        return tmp_path.resolve()
    return output_path.resolve()
