from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook

from financial_variable_curation.inspection.exceptions import InspectionError
from financial_variable_curation.inspection.inspection_service import InspectionService
from financial_variable_curation.models.artifacts import InspectionSummary
from financial_variable_curation.models.variable import VariableProfile
from financial_variable_curation.models.workbook import WorkbookProfile


def _write_workbook(path: Path, sheets: list[tuple[str, list[list[object]]]]) -> Path:
    workbook = Workbook()
    first = True
    for sheet_name, rows in sheets:
        if first:
            sheet = workbook.active
            sheet.title = sheet_name
            first = False
        else:
            sheet = workbook.create_sheet(sheet_name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    return path


def _daily_rows(days: int = 10) -> list[list[object]]:
    rows = [["date", "value"]]
    start = date(2024, 1, 1)
    for index in range(days):
        rows.append([start + timedelta(days=index), index + 1])
    return rows


def test_standard_single_sheet_daily(tmp_path) -> None:
    path = _write_workbook(tmp_path / "daily.xlsx", [("Daily", _daily_rows())])
    result = InspectionService().inspect(path)
    assert result.inspection_summary.status == "COMPLETED"
    assert result.workbook_profile.file_hash
    assert result.workbook_profile.sheets[0].date_detection.status == "DETECTED"
    profile = result.variable_profiles[0]
    assert profile.column_name == "value"
    assert profile.frequency.detected_frequency == "daily"
    assert profile.is_valid_variable is True


def test_multi_sheet_excel(tmp_path) -> None:
    path = _write_workbook(
        tmp_path / "multi.xlsx",
        [
            ("A", _daily_rows()),
            ("B", _daily_rows()),
        ],
    )
    result = InspectionService().inspect(path)
    assert len(result.workbook_profile.sheets) == 2
    assert len(result.variable_profiles) == 2
    assert {sheet.sheet_name for sheet in result.workbook_profile.sheets} == {"A", "B"}


def test_empty_sheet_is_recorded(tmp_path) -> None:
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.active.append(["date", "value"])
    workbook.active.append([date(2024, 1, 1), 1])
    workbook.create_sheet("Empty")
    path = tmp_path / "empty_sheet.xlsx"
    workbook.save(path)

    result = InspectionService().inspect(path)
    empty_sheet = next(sheet for sheet in result.workbook_profile.sheets if sheet.sheet_name == "Empty")
    assert empty_sheet.empty_sheet is True
    assert result.workbook_profile.empty_sheet_count == 1


def test_explicit_header_row(tmp_path) -> None:
    path = _write_workbook(
        tmp_path / "explicit_header.xlsx",
        [
            (
                "Sheet1",
                [
                    ["Title row"],
                    [],
                    ["date", "value"],
                    [date(2024, 1, 1), 1],
                    [date(2024, 1, 2), 2],
                    [date(2024, 1, 3), 3],
                    [date(2024, 1, 4), 4],
                ],
            )
        ],
    )
    result = InspectionService().inspect(path, header_row=3)
    sheet = result.workbook_profile.sheets[0]
    assert sheet.header_detection.header_row == 3
    assert sheet.date_detection.status == "DETECTED"
    assert result.variable_profiles[0].column_name == "value"


def test_three_row_metadata_header_layout(tmp_path) -> None:
    rows = [
        ["指标名称", "SHFE: 锡: 主力合约: 收盘价", "SHFE: 锡: 主力合约: 成交量"],
        ["单位", "元/吨", "手"],
        ["频率", "日", "日"],
        [date(2024, 1, 1), 420840, 106368],
        [date(2024, 1, 2), 418540, 147740],
        [date(2024, 1, 3), 417750, 136600],
    ]
    path = _write_workbook(tmp_path / "three_row_header.xlsx", [("Sheet1", rows)])
    result = InspectionService().inspect(path)
    sheet = result.workbook_profile.sheets[0]
    assert sheet.header_detection.header_row == 1
    assert sheet.metadata_row_count == 3
    assert sheet.row_count == 3
    assert sheet.selected_date_column == "指标名称"
    assert len(result.variable_profiles) == 2
    assert result.variable_profiles[0].column_name == "SHFE: 锡: 主力合约: 收盘价"
    assert result.variable_profiles[0].unit_hint == "元/吨"
    assert result.variable_profiles[0].declared_frequency == "daily"
    assert result.variable_profiles[0].frequency.detected_frequency == "daily"
    assert result.variable_profiles[1].unit_hint == "手"


def test_header_not_in_first_row_auto_detected(tmp_path) -> None:
    path = _write_workbook(
        tmp_path / "auto_header.xlsx",
        [
            (
                "Sheet1",
                [
                    ["Title row"],
                    [],
                    ["date", "value"],
                    [date(2024, 1, 1), 1],
                    [date(2024, 1, 2), 2],
                ],
            )
        ],
    )
    result = InspectionService().inspect(path)
    assert result.workbook_profile.sheets[0].header_detection.status == "DETECTED"
    assert result.workbook_profile.sheets[0].header_detection.header_row == 3


def test_single_date_column(tmp_path) -> None:
    result = InspectionService().inspect(_write_workbook(tmp_path / "single_date.xlsx", [("S", _daily_rows())]))
    assert result.workbook_profile.sheets[0].date_detection.selected_column == "date"
    assert result.workbook_profile.sheets[0].date_detection.confidence > 0


def test_multiple_suspected_date_columns_requires_review(tmp_path) -> None:
    rows = [["date", "trade_date", "value"]]
    start = date(2024, 1, 1)
    for index in range(10):
        day = start + timedelta(days=index)
        rows.append([day, day, index])
    path = _write_workbook(tmp_path / "multiple_dates.xlsx", [("S", rows)])
    result = InspectionService().inspect(path)
    sheet = result.workbook_profile.sheets[0]
    assert sheet.date_detection.status == "NEEDS_REVIEW"
    assert len(sheet.date_detection.candidates) >= 2
    assert any("multiple" in issue for issue in sheet.issues)


def test_no_date_column(tmp_path) -> None:
    path = _write_workbook(
        tmp_path / "no_date.xlsx",
        [("S", [["value"], [1], [2], [3], [4]])],
    )
    result = InspectionService().inspect(path)
    assert result.workbook_profile.sheets[0].date_detection.status == "NO_DATE_COLUMN"
    assert result.variable_profiles[0].frequency.detected_frequency == "unknown"


def test_mixed_date_format_requires_review(tmp_path) -> None:
    path = _write_workbook(
        tmp_path / "mixed_date.xlsx",
        [
            (
                "S",
                [
                    ["date", "value"],
                    ["2024-01-01", 1],
                    ["2024/01/02", 2],
                    ["20240103", 3],
                    ["2024-01-04", 4],
                    ["2024/01/05", 5],
                ],
            )
        ],
    )
    result = InspectionService().inspect(path)
    assert result.workbook_profile.sheets[0].date_detection.status == "NEEDS_REVIEW"


@pytest.mark.parametrize(
    "frequency,step_days,expected",
    [
        ("daily", 1, "daily"),
        ("weekly", 7, "weekly"),
        ("monthly", None, "monthly"),
        ("quarterly", None, "quarterly"),
        ("annual", None, "annual"),
    ],
)
def test_frequency_detection(tmp_path, frequency, step_days, expected) -> None:
    rows = [["date", "value"]]
    if frequency == "monthly":
        dates = [date(2024, month, 1) for month in range(1, 13)]
    elif frequency == "quarterly":
        dates = [date(2024, month, 1) for month in (1, 4, 7, 10)] + [date(2025, 1, 1)]
    elif frequency == "annual":
        dates = [date(2024 + offset, 1, 1) for offset in range(5)]
    else:
        start = date(2024, 1, 1)
        dates = [start + timedelta(days=index * step_days) for index in range(12)]
    for index, day in enumerate(dates):
        rows.append([day, index + 1])
    result = InspectionService().inspect(_write_workbook(tmp_path / f"{frequency}.xlsx", [("S", rows)]))
    assert result.variable_profiles[0].frequency.detected_frequency == expected


def test_irregular_frequency(tmp_path) -> None:
    dates = [date(2024, 1, 1), date(2024, 1, 2), date(2024, 1, 10), date(2024, 1, 15), date(2024, 2, 1)]
    rows = [["date", "value"]] + [[day, index + 1] for index, day in enumerate(dates)]
    result = InspectionService().inspect(_write_workbook(tmp_path / "irregular.xlsx", [("S", rows)]))
    assert result.variable_profiles[0].frequency.detected_frequency == "irregular"


def test_all_empty_variable_column(tmp_path) -> None:
    rows = [["date", "empty"]] + [[date(2024, 1, 1) + timedelta(days=index), None] for index in range(5)]
    result = InspectionService().inspect(_write_workbook(tmp_path / "empty_col.xlsx", [("S", rows)]))
    profile = result.variable_profiles[0]
    assert profile.all_empty is True
    assert profile.is_valid_variable is False


def test_constant_variable_column(tmp_path) -> None:
    rows = [["date", "constant"]] + [[date(2024, 1, 1) + timedelta(days=index), 7] for index in range(5)]
    result = InspectionService().inspect(_write_workbook(tmp_path / "constant_col.xlsx", [("S", rows)]))
    profile = result.variable_profiles[0]
    assert profile.constant is True
    assert profile.quality_score == 10.0


def test_high_missing_rate(tmp_path) -> None:
    rows = [["date", "high_missing"]] + [
        [date(2024, 1, 1) + timedelta(days=index), 100 if index in {0, 19} else None]
        for index in range(20)
    ]
    result = InspectionService().inspect(_write_workbook(tmp_path / "high_missing.xlsx", [("S", rows)]))
    profile = result.variable_profiles[0]
    assert profile.missing_rate > 0.8
    assert profile.latest_observation_date is not None


def test_duplicate_column_names_are_not_overwritten(tmp_path) -> None:
    rows = [["date", "value", "value"]] + [
        [date(2024, 1, 1) + timedelta(days=index), index, index * 2]
        for index in range(5)
    ]
    result = InspectionService().inspect(_write_workbook(tmp_path / "duplicate_headers.xlsx", [("S", rows)]))
    names = [profile.column_name for profile in result.variable_profiles]
    assert names == ["value", "value__col_3"]
    assert result.variable_profiles[0].variable_id != result.variable_profiles[1].variable_id


def test_duplicate_dates_are_counted(tmp_path) -> None:
    rows = [["date", "value"]]
    day = date(2024, 1, 1)
    for index in range(5):
        rows.append([day if index < 3 else day + timedelta(days=1), index])
    result = InspectionService().inspect(_write_workbook(tmp_path / "duplicate_dates.xlsx", [("S", rows)]))
    profile = result.variable_profiles[0]
    assert profile.duplicate_date_count >= 2


def test_pseudo_daily_detection_is_conservative(tmp_path) -> None:
    rows = [["date", "monthly_value"]]
    start = date(2024, 1, 1)
    for index in range(90):
        day = start + timedelta(days=index)
        rows.append([day, day.month])
    result = InspectionService().inspect(_write_workbook(tmp_path / "pseudo_daily.xlsx", [("S", rows)]))
    profile = result.variable_profiles[0]
    assert profile.frequency.detected_frequency == "daily"
    assert profile.is_pseudo_high_frequency is True
    assert profile.pseudo_frequency_confidence > 0
    assert profile.pseudo_frequency_reason


def test_non_numeric_comment_column_is_flagged(tmp_path) -> None:
    rows = [["date", "note"]] + [
        [date(2024, 1, 1) + timedelta(days=index), f"note-{index}"]
        for index in range(5)
    ]
    result = InspectionService().inspect(_write_workbook(tmp_path / "comment_col.xlsx", [("S", rows)]))
    profile = result.variable_profiles[0]
    assert profile.is_comment_column is True
    assert profile.is_valid_variable is False


def test_nonexistent_input_file_raises(tmp_path) -> None:
    with pytest.raises(InspectionError):
        InspectionService().inspect(tmp_path / "missing.xlsx")


def test_non_xlsx_file_raises(tmp_path) -> None:
    path = tmp_path / "input.txt"
    path.write_text("not excel", encoding="utf-8")
    with pytest.raises(InspectionError, match="Only .xlsx"):
        InspectionService().inspect(path)


def test_corrupted_excel_raises(tmp_path) -> None:
    path = tmp_path / "corrupt.xlsx"
    path.write_bytes(b"this is not a zip file")
    with pytest.raises(InspectionError):
        InspectionService().inspect(path)


def test_unable_to_detect_header_marks_review(tmp_path) -> None:
    path = _write_workbook(tmp_path / "no_header.xlsx", [("S", [[1, 2], [3, 4], [5, 6]])])
    result = InspectionService().inspect(path)
    sheet = result.workbook_profile.sheets[0]
    assert sheet.header_detection.status == "NEEDS_REVIEW"
    assert result.variable_profiles == []


def test_sample_insufficient_frequency_returns_unknown(tmp_path) -> None:
    rows = [["date", "value"], [date(2024, 1, 1), 1], [date(2024, 1, 2), 2]]
    result = InspectionService().inspect(_write_workbook(tmp_path / "insufficient.xlsx", [("S", rows)]))
    assert result.variable_profiles[0].frequency.detected_frequency == "unknown"
    assert result.variable_profiles[0].frequency.frequency_reason_codes == ["INSUFFICIENT_DATES"]


def test_numeric_strings_are_parsed(tmp_path) -> None:
    rows = [["date", "value"], [date(2024, 1, 1), "100"], [date(2024, 1, 2), "200"], [date(2024, 1, 3), "300"]]
    result = InspectionService().inspect(_write_workbook(tmp_path / "numeric_strings.xlsx", [("S", rows)]))
    profile = result.variable_profiles[0]
    assert profile.data_type.value == "NUMERIC"
    assert profile.numeric_parse_success_rate == 1.0
    assert profile.is_numeric_candidate is True


def test_sheet_filter(tmp_path) -> None:
    path = _write_workbook(tmp_path / "sheet_filter.xlsx", [("A", _daily_rows()), ("B", _daily_rows())])
    result = InspectionService().inspect(path, sheet_filter="A")
    assert result.inspection_summary.sheets_processed == 1
    assert result.workbook_profile.sheet_names == ["A"]
    assert len(result.variable_profiles) == 1


def test_json_outputs_roundtrip_through_pydantic(tmp_path) -> None:
    result = InspectionService().inspect(_write_workbook(tmp_path / "roundtrip.xlsx", [("S", _daily_rows())]))
    assert WorkbookProfile.model_validate(result.workbook_profile.model_dump(mode="json"))
    assert VariableProfile.model_validate(result.variable_profiles[0].model_dump(mode="json"))
    assert InspectionSummary.model_validate(result.inspection_summary.model_dump(mode="json"))


def test_variable_id_is_stable(tmp_path) -> None:
    path = _write_workbook(tmp_path / "stable.xlsx", [("S", _daily_rows())])
    first = InspectionService().inspect(path)
    second = InspectionService().inspect(path)
    assert [profile.variable_id for profile in first.variable_profiles] == [
        profile.variable_id for profile in second.variable_profiles
    ]
