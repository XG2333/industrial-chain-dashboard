from __future__ import annotations

import argparse
import shutil
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

from financial_variable_curation.export.directory_writer import (
    _infer_category_cn,
    _infer_subcategory_cn,
)


CAT_CN_TO_EN = {
    "价格": "PRICE",
    "成本利润": "COST",
    "库存": "INVENTORY",
    "供应": "QUANTITY",
    "需求": "QUANTITY",
    "进出口": "QUANTITY",
    "平衡": "QUANTITY",
    "其他": "UNKNOWN",
}

DEFAULT_DB_PATH = Path(r"C:\Users\11\Documents\Selecting skill\data\output_processed_final.db")


def load_classification_map(db_path: Path) -> dict[str, dict]:
    con = sqlite3.connect(str(db_path))
    con.row_factory = sqlite3.Row
    rows = con.execute(
        "SELECT standard_name, metric_name, category_level_1, category_level_2, "
        "unit, data_nature, confidence FROM variable_classifications"
    ).fetchall()
    con.close()
    mapping: dict[str, dict] = {}
    for row in rows:
        name = str(row["standard_name"] or row["metric_name"] or "").strip()
        if name:
            mapping[name] = dict(row)
    return mapping


def reclassify(path: Path, db_path: Path = DEFAULT_DB_PATH) -> None:
    path = Path(path).resolve()
    db_path = Path(db_path).resolve()
    class_map = load_classification_map(db_path)
    tmp = path.with_name(f"{path.stem}__tmp_reclassify.xlsx")
    shutil.copy2(path, tmp)

    wb = load_workbook(tmp)
    ws = wb[wb.sheetnames[0]]
    current_sheet = ""
    changed = 0
    total = 0

    for row in ws.iter_rows(min_row=2):
        c0 = str(row[0].value).strip() if row[0].value is not None else ""
        if not c0:
            continue
        if "sht" in c0:
            continue
        if not c0.isdigit():
            current_sheet = c0
            continue

        c3 = str(row[3].value).strip() if row[3].value is not None else ""
        if not c3.isdigit():
            continue

        name = str(row[4].value).strip() if row[4].value is not None else ""
        if not name:
            continue

        old_major = str(row[7].value).strip() if row[7].value is not None else ""
        old_sub = str(row[8].value).strip() if row[8].value is not None else ""
        entry = {
            "original_name": name,
            "sheet_name": current_sheet,
            "unit_hint": str(row[5].value).strip() if row[5].value is not None else "",
            "category_level_1": CAT_CN_TO_EN.get(old_major, "UNKNOWN"),
            "category_level_2": "POSITION" if old_sub == "持仓" else old_sub,
        }
        db_entry = class_map.get(name)
        if db_entry:
            entry["category_level_1"] = str(db_entry["category_level_1"] or "")
            entry["category_level_2"] = str(db_entry["category_level_2"] or "")
            entry["unit_hint"] = str(db_entry["unit"] or "") or entry["unit_hint"]
        new_major = _infer_category_cn(entry)
        new_sub = _infer_subcategory_cn(entry)
        if new_major != old_major or new_sub != old_sub:
            row[7].value = new_major
            row[8].value = new_sub
            changed += 1
        total += 1

    wb.save(tmp)
    wb.close()
    try:
        tmp.replace(path)
    except PermissionError:
        print(f"Permission error writing {path}; temp file kept at {tmp}")
        raise
    print(f"Reclassified {path}: total={total}, changed={changed}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Reclassify directory rows from a curation database.")
    parser.add_argument("excel", nargs="*", help="Excel files to reclassify.")
    parser.add_argument("--database", default=DEFAULT_DB_PATH, help="Path to curation SQLite database.")
    args = parser.parse_args()
    paths = [Path(arg) for arg in args.excel] if args.excel else [
        Path(r"C:\Users\11\Documents\Selecting skill\data\output_processed_final.xlsx"),
        Path(r"C:\Users\11\Documents\market-ai-dashboard\output_processed_final_v2.xlsx"),
    ]
    for path in paths:
        reclassify(path, db_path=Path(args.database))
