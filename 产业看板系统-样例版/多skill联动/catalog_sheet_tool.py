# -*- coding: utf-8 -*-
"""目录表精简/合并工具：加速大 Excel 流水线。

背景：流水线中大量步骤（校准/频率校验/一致性检查/筛选/排序等）只操作
第一个 sheet（指标目录表），却要对 300+ 个 sheet 的完整文件做 openpyxl
全量 load + save（每次 20-40 秒）。本工具把"完整文件"与"仅目录表"解耦：

- extract：把完整文件的第一个 sheet 提取为精简文件（几十 KB），后续
  只碰目录表的步骤在精简文件上跑（每次 <1 秒）；同时同目录保存一份
  完整文件副本作为合并模板（<基础名>__full_template.xlsx）。
- merge：把处理后的精简文件合并回完整模板（第一个 sheet 覆盖 + 精简中
  新增的 sheet 如"成交持仓比计算"追加），输出完整文件。

约束：extract 的完整文件必须包含所有需要保留的 sheet（净出口计算、
复合指标计算、数据 sheets）；merge 后输出 = 模板 + 处理后的目录表 +
精简中新增的 sheet。所有值、样式、超链接、列宽/行高、合并单元格、
冻结窗格、筛选均保留。
"""

from __future__ import annotations

import argparse
import shutil
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

# 模板文件名按输入基础名区分（三行业并行时互不覆盖）：
# extract 的 full（xxx_workflow_ai_step1）与 merge 的 slim（xxx_workflow_ai_slim /
# xxx_workflow_ai_selected / xxx_workflow_ai_postsel）均取"去步骤后缀的基础名"
# （xxx_workflow_ai）→ xxx_workflow_ai__full_template.xlsx，两端一致。
def _template_name(path: Path) -> str:
    stem = path.stem
    for suffix in ("_step1", "_slim", "_selected", "_postsel", "_workflow_ai"):
        if stem.endswith(suffix):
            stem = stem[: -len(suffix)]
            break
    return f"{stem}__full_template.xlsx"


def _copy_sheet_style(src_ws, dst_ws, src_wb, dst_wb) -> None:
    """把 src_ws 的样式级属性复制到 dst_ws（同工作簿或跨工作簿）。"""
    # 列宽
    for col_letter, dim in src_ws.column_dimensions.items():
        if dim.width is not None:
            dst_ws.column_dimensions[col_letter].width = dim.width
        if dim.hidden:
            dst_ws.column_dimensions[col_letter].hidden = True
    # 行高
    for row_num, dim in src_ws.row_dimensions.items():
        if dim.height is not None:
            dst_ws.row_dimensions[row_num].height = dim.height
        if dim.hidden:
            dst_ws.row_dimensions[row_num].hidden = True
    # 合并单元格
    for merged in src_ws.merged_cells.ranges:
        try:
            dst_ws.merge_cells(str(merged))
        except ValueError:
            pass  # 目标已存在相同合并，忽略
    # 冻结窗格
    if src_ws.freeze_panes:
        dst_ws.freeze_panes = src_ws.freeze_panes
    # 筛选
    if src_ws.auto_filter.ref:
        dst_ws.auto_filter.ref = src_ws.auto_filter.ref


def _copy_cell(src_cell, dst_cell) -> None:
    """复制单元格值 + 完整样式 + 超链接。"""
    # MergedCell 只读且无独立值/样式（合并关系由 merge_cells 维护），跳过
    if isinstance(src_cell, MergedCell) or isinstance(dst_cell, MergedCell):
        return
    dst_cell.value = src_cell.value
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.border = copy(src_cell.border)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy(src_cell.protection)
        dst_cell.alignment = copy(src_cell.alignment)
    if src_cell.hyperlink is not None:
        dst_cell.hyperlink = copy(src_cell.hyperlink)


def extract(full_path: Path, slim_path: Path) -> dict:
    """从完整文件提取第一个 sheet 为精简文件；同目录保存完整模板副本。"""
    full_path = full_path.resolve()
    slim_path = slim_path.resolve()
    template = slim_path.parent / _template_name(full_path)

    wb = load_workbook(full_path)
    first = wb.sheetnames[0]
    sheets_before = len(wb.sheetnames)
    for name in list(wb.sheetnames):
        if name != first:
            del wb[name]
    wb.save(slim_path)
    wb.close()

    # 模板副本（完整文件本身），供 merge 使用
    shutil.copy2(full_path, template)
    return {"sheets_before": sheets_before, "kept_sheet": first, "template": str(template)}


def merge(slim_path: Path, output_path: Path, template_path: Path | None = None) -> dict:
    """把精简文件合并回完整模板，输出完整文件。

    - 精简文件第一个 sheet（目录表）覆盖模板第一个 sheet（值+样式）
    - 精简文件中模板没有的 sheet（如"成交持仓比计算"）追加到模板
    - 模板其余 sheets（数据 sheets、净出口计算、复合指标计算）保留
    """
    slim_path = slim_path.resolve()
    output_path = output_path.resolve()
    template_path = (template_path or slim_path.parent / _template_name(slim_path)).resolve()
    if not template_path.exists():
        raise FileNotFoundError(f"Merge template not found: {template_path}")

    slim_wb = load_workbook(slim_path)
    template_wb = load_workbook(template_path)

    slim_sheet_names = list(slim_wb.sheetnames)
    template_sheet_names = list(template_wb.sheetnames)
    stats = {"slim_sheets": len(slim_sheet_names), "template_sheets": len(template_sheet_names)}

    # 1) 第一个 sheet（目录表）：覆盖
    slim_first = slim_wb[slim_sheet_names[0]]
    tpl_first = template_wb[template_sheet_names[0]]
    # 清空模板 ws0 现有的合并单元格与超链接（处理后的 slim ws0 才是真源；
    # 未清除会导致模板旧合并/超链接残留，如 selection 已删除的合并被加回）
    for merged in list(tpl_first.merged_cells.ranges):
        tpl_first.unmerge_cells(str(merged))
    for row in tpl_first.iter_rows():
        for cell in row:
            if cell.hyperlink is not None:
                cell.hyperlink = None
    # 目标多余行清空（正常情况下 slim 行数 >= 模板行数，保险处理）
    if tpl_first.max_row > slim_first.max_row:
        tpl_first.delete_rows(slim_first.max_row + 1, tpl_first.max_row - slim_first.max_row)
    for row in slim_first.iter_rows():
        for src_cell in row:
            dst_cell = tpl_first.cell(row=src_cell.row, column=src_cell.column)
            _copy_cell(src_cell, dst_cell)
    _copy_sheet_style(slim_first, tpl_first, slim_wb, template_wb)

    # 2) 精简文件中模板没有的 sheet：追加
    added = []
    for name in slim_sheet_names:
        if name not in template_wb.sheetnames:
            src_ws = slim_wb[name]
            dst_ws = template_wb.create_sheet(name)
            for row in src_ws.iter_rows():
                for src_cell in row:
                    dst_cell = dst_ws.cell(row=src_cell.row, column=src_cell.column)
                    _copy_cell(src_cell, dst_cell)
            _copy_sheet_style(src_ws, dst_ws, slim_wb, template_wb)
            added.append(name)
    stats["added_sheets"] = added

    # 3) 同名且非第一个的 sheet：用精简版覆盖（正常情况下不存在）
    replaced = []
    for name in slim_sheet_names:
        if name != slim_sheet_names[0] and name in template_wb.sheetnames:
            src_ws = slim_wb[name]
            dst_ws = template_wb[name]
            for merged in list(dst_ws.merged_cells.ranges):
                dst_ws.unmerge_cells(str(merged))
            for row in dst_ws.iter_rows():
                for cell in row:
                    if cell.hyperlink is not None:
                        cell.hyperlink = None
            if dst_ws.max_row > src_ws.max_row:
                dst_ws.delete_rows(src_ws.max_row + 1, dst_ws.max_row - src_ws.max_row)
            for row in src_ws.iter_rows():
                for src_cell in row:
                    dst_cell = dst_ws.cell(row=src_cell.row, column=src_cell.column)
                    _copy_cell(src_cell, dst_cell)
            _copy_sheet_style(src_ws, dst_ws, slim_wb, template_wb)
            replaced.append(name)
    stats["replaced_sheets"] = replaced

    output_path.parent.mkdir(parents=True, exist_ok=True)
    template_wb.save(output_path)
    template_wb.close()
    slim_wb.close()
    stats["output"] = str(output_path)
    return stats


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="目录表精简/合并工具")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p_extract = sub.add_parser("extract", help="提取第一个 sheet 为精简文件")
    p_extract.add_argument("full", help="完整文件路径")
    p_extract.add_argument("slim", help="输出精简文件路径")

    p_merge = sub.add_parser("merge", help="把精简文件合并回完整模板")
    p_merge.add_argument("slim", help="精简文件路径（其同目录的 <基础名>__full_template.xlsx 为模板）")
    p_merge.add_argument("output", help="输出完整文件路径")
    p_merge.add_argument("--template", default=None, help="显式指定模板路径（默认同目录 <基础名>__full_template.xlsx）")

    args = parser.parse_args(argv)
    if args.cmd == "extract":
        stats = extract(Path(args.full), Path(args.slim))
        print(f"extract: kept={stats['kept_sheet']} sheets_before={stats['sheets_before']}")
        print(f"  slim={args.slim}")
        print(f"  template={stats['template']}")
    elif args.cmd == "merge":
        stats = merge(Path(args.slim), Path(args.output), Path(args.template) if args.template else None)
        print(f"merge: slim_sheets={stats['slim_sheets']} template_sheets={stats['template_sheets']}")
        print(f"  added={stats['added_sheets']} replaced={stats['replaced_sheets']}")
        print(f"  output={stats['output']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
