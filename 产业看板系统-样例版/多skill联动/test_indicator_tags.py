from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.apply_indicator_tags import apply_indicator_tags, extract_indicator_tags
from workflow.config import WorkflowConfig


PROJECT_ROOT = Path(__file__).resolve().parent.parent


def test_extract_indicator_tags_for_spot_metal_silicon() -> None:
    tags = extract_indicator_tags(
        sheet="工业硅现货价格-日",
        name="SMM: 通氧553#硅（新疆）-平均价：日度",
        unit="元/吨",
        frequency="日度",
        major="价格",
        sub="现货价格",
    )
    assert tags["产品"] == "工业硅"
    assert tags["研究主题"] == "价格"
    assert tags["指标类型"] == "现货价格"
    assert tags["规格"] == "553#"
    assert tags["工艺属性"] == "通氧"
    assert tags["地域"] == "新疆"
    assert tags["统计口径"] == "平均价"
    assert tags["状态"] == "实际"
    assert tags["频率"] == "日度"


def test_extract_indicator_tags_for_futures_price() -> None:
    tags = extract_indicator_tags(
        sheet="工业硅期货价格+基差+月差-日",
        name="GFEX: 工业硅: 主力合约: 收盘价: 日度",
        unit="元/吨",
        frequency="日度",
        major="价格",
        sub="期货价格",
    )
    assert tags["产品"] == "工业硅"
    assert tags["研究主题"] == "价格"
    assert tags["指标类型"] == "期货价格"
    assert tags["统计口径"] == "收盘价"
    assert tags["合约"] == "主力合约"
    assert tags["频率"] == "日度"


def test_basis_under_spot_spread_keeps_basis_type() -> None:
    tags = extract_indicator_tags(
        sheet="工业硅期货价格+基差+月差-日",
        name="SMM: 多晶硅期现价差: 当月合约: 收盘价: 日度",
        unit="元/吨",
        frequency="日度",
        major="价格",
        sub="现货价差",
    )
    assert tags["指标类型"] == "基差"


def test_extract_indicator_tags_avoids_overlapping_process_and_region() -> None:
    tags = extract_indicator_tags(
        sheet="工业硅现货价格-日",
        name="SMM: 不通氧553#硅(天津港)-平均价: 日度",
        unit="元/吨",
        frequency="日度",
        major="价格",
        sub="现货价格",
    )
    assert tags["工艺属性"] == "不通氧"
    assert tags["地域"] == "天津港"


def test_imported_adc12_price_is_price_not_import_export() -> None:
    tags = extract_indicator_tags(
        sheet="铝合金进口利润-日",
        name="SMM: 进口ADC12宁波CIF低价: 日度",
        unit="元/吨",
        frequency="日度",
        major="价格",
        sub="现货价格",
    )
    assert tags["研究主题"] == "价格"
    assert tags["指标类型"] == "现货价格"


def test_component_cost_index_is_cost() -> None:
    tags = extract_indicator_tags(
        sheet="组件成本-周",
        name="SMM: Topcon183组件成本指数-一体化 - 平均价: 周度",
        unit="元/瓦",
        frequency="周度",
        major="成本利润",
        sub="成本",
    )
    assert tags["研究主题"] == "成本利润"
    assert tags["指标类型"] == "成本"


def test_cost_model_price_keeps_price_type() -> None:
    tags = extract_indicator_tags(
        sheet="光伏EVA成本模型-日",
        name="SMM: 光伏级EVA成本(油制): EVA粒子价格: 日度",
        unit="元/吨",
        frequency="日度",
        major="成本利润",
        sub="现货价格",
    )
    assert tags["研究主题"] == "成本利润"
    assert tags["指标类型"] == "现货价格"


def test_pmi_subindex_is_macro() -> None:
    tags = extract_indicator_tags(
        sheet="铝合金PMI-月",
        name="SMM: 原生铝合金新出口订单指数: 月度",
        unit="指数",
        frequency="月度",
        major="其他",
        sub="宏观",
    )
    assert tags["研究主题"] == "其他"
    assert tags["指标类型"] == "宏观"


def test_balance_import_is_import_export() -> None:
    tags = extract_indicator_tags(
        sheet="工业硅平衡-月",
        name="SMM: 工业硅平衡新（按照下游消费算）: 进口量: 月度",
        unit="吨",
        frequency="月度",
        major="进出口",
        sub="进出口",
    )
    assert tags["研究主题"] == "进出口"
    assert tags["指标类型"] == "进口"


def test_supply_in_balance_sheet_is_supply() -> None:
    tags = extract_indicator_tags(
        sheet="多晶硅全球平衡-月",
        name="SMM: 全球多晶硅供需平衡表: 供应量: 月度",
        unit="吨",
        frequency="月度",
        major="供应",
        sub="数量",
    )
    assert tags["研究主题"] == "供给"
    assert tags["指标类型"] == "数量"


def test_dmc_average_price_stays_price() -> None:
    tags = extract_indicator_tags(
        sheet="有机硅成本利润-月",
        name="SMM: 有机硅DMC均价: 月度",
        unit="元/吨",
        frequency="月度",
        major="成本利润",
        sub="成本",
    )
    assert tags["研究主题"] == "价格"
    assert tags["指标类型"] == "现货价格"


def test_apply_indicator_tags_adds_column(tmp_path: Path) -> None:
    path = tmp_path / "directory.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "#",
            "Sheet Name",
            "Freq",
            "Col",
            "Indicator Name",
            "Unit",
            "板块",
            "大类",
            "子类",
            "数据性质",
            "置信度",
            "是否选中",
            "状态说明",
        ]
    )
    ws.append(
        [
            1,
            "",
            "日度",
            1,
            "SMM: 通氧553#硅（新疆）-平均价：日度",
            "元/吨",
            "工业硅",
            "价格",
            "现货价格",
            "水平值",
            0.95,
            "是",
            "示例",
        ]
    )
    wb.save(path)

    apply_indicator_tags(path)

    wb2 = load_workbook(path)
    ws2 = wb2.active
    headers = [cell.value for cell in ws2[1]]
    assert headers[-1] == "指标标签"
    tags = json.loads(ws2.cell(2, len(headers)).value)
    assert tags["产品"] == "工业硅"
    assert tags["地域"] == "新疆"
    assert tags["统计口径"] == "平均价"
    wb2.close()


def test_silicon_workflow_config_merges_indicator_tags_into_skill2_runner() -> None:
    config = WorkflowConfig.from_yaml(PROJECT_ROOT / "configs" / "workflow_silicon.yaml")
    ids = [step.id for step in config.steps]
    assert "apply_indicator_tags" not in ids
    assert "check_silicon_output" in ids
