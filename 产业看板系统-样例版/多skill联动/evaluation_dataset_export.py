# -*- coding: utf-8 -*-
"""Phase 2b：历史人工确认案例 → Langfuse Dataset（幂等导入）。

- 输入：Phase 2a 的 human_cases.csv（含恢复的 physical_variable_id）+
  可选 priority_review_pack.xlsx（candidates 上下文 / industry / commodity）。
- Dataset 名称长期稳定（默认 classification/human_confirmed），版本放在
  dataset metadata 与 item metadata（version / schema_version），不随版本换名。
- 每条 item：
    id = golden_<industry>_<commodity>_<physical_variable_id>（稳定、与顺序无关）
    input = physical_variable_id / original_name / industry / commodity / sheet /
            unit / frequency / candidates（仅当历史数据真实存在且非空）
    expected_output = 仅人工实际确认过的字段（大类/子类/是否选中，空值不写，
                      不推断 golden 中原本为 null 的字段）
    metadata = source=human_review / source_case_id / physical_variable_id /
               audit_crosscheck / original_run_id / imported_at / schema_version
    status = ACTIVE
- 幂等：create_dataset_item(id=...) upsert；比较内容（imported_at 不参与比较），
  相同数据 → unchanged（no-op），内容变化 → updated，新 id → created。
- 模式：--dry-run（默认，只生成报告与 items jsonl）/
        --apply（真实写入 Langfuse） / --verify（API 回查）。

用法：
  python scripts/evaluation_dataset_export.py \
      --human-cases workflow_runs/<run>/audit/evaluation/human_cases.csv \
      --pack <priority_review_pack.xlsx> \
      --run-dir workflow_runs/<run> \
      [--apply|--verify] [--dataset-name classification/human_confirmed]
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]

try:
    from dotenv import load_dotenv

    load_dotenv(PROJECT_ROOT / ".env")
except Exception:  # noqa: BLE001
    pass

DEFAULT_DATASET_NAME = "classification/human_confirmed"
DEFAULT_SCHEMA_VERSION = "1.0"
DEFAULT_VERSION = "v1"

ITEM_ID_PREFIX = "golden"  # 与 golden case_id 同构：golden_<industry>_<commodity>_<vid>

EXPECTED_OUTPUT_FIELDS = ("大类", "子类", "是否选中")


def stable_item_id(industry: str, commodity: str, physical_variable_id: str) -> str:
    return f"{ITEM_ID_PREFIX}_{industry}_{commodity}_{physical_variable_id}"


def load_human_cases(csv_path: Path) -> list[dict]:
    with csv_path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def load_pack_context(pack_path: Path | None) -> dict[str, dict]:
    """priority_review_pack.xlsx → {原始指标: {industry, commodity, candidates}}。"""
    if pack_path is None:
        return {}
    from openpyxl import load_workbook

    wb = load_workbook(pack_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    hdr = [str(c or "").strip() for c in rows[0]]
    idx = {name: i for i, name in enumerate(hdr)}
    out: dict[str, dict] = {}
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        name = str(r[idx["原始指标"]] or "").strip()
        candidates = str(r[idx["candidate categories"]] or "").strip()
        parsed = []
        if candidates and candidates != "[]":
            try:
                parsed = json.loads(candidates)
            except Exception:  # noqa: BLE001
                parsed = []
        out[name] = {
            "industry": str(r[idx["industry"]] or "").strip(),
            "commodity": str(r[idx["commodity"]] or "").strip(),
            "candidates": parsed if isinstance(parsed, list) else [],
        }
    return out


def build_items(human_cases: list[dict], pack_context: dict[str, dict]) -> list[dict]:
    items: list[dict] = []
    for hc in human_cases:
        vid = hc.get("physical_variable_id", "")
        name = hc.get("original_name", "")
        if not vid:
            continue
        ctx = pack_context.get(name, {})
        industry = ctx.get("industry") or "unknown"
        commodity = ctx.get("commodity") or "UNKNOWN"
        item_id = stable_item_id(industry, commodity, vid)

        expected: dict[str, str] = {}
        for key, field in (("human_major", "大类"), ("human_sub", "子类"), ("human_selected", "是否选中")):
            val = str(hc.get(key) or "").strip()
            if val:
                expected[field] = val

        inp: dict[str, Any] = {
            "physical_variable_id": vid,
            "original_name": name,
            "industry": industry,
            "commodity": commodity,
        }
        for key in ("sheet", "unit", "frequency"):
            val = str(hc.get(key) or "").strip()
            if val:
                inp[key] = val
        if ctx.get("candidates"):
            inp["candidates"] = ctx["candidates"]

        metadata = {
            "source": "human_review",
            "source_case_id": hc.get("source_case_id", ""),
            "physical_variable_id": vid,
            "audit_crosscheck": hc.get("audit_crosscheck", "-"),
            "original_run_id": hc.get("workflow_run_id", ""),
            "version": DEFAULT_VERSION,
            "schema_version": DEFAULT_SCHEMA_VERSION,
        }
        items.append({
            "id": item_id,
            "input": inp,
            "expected_output": expected,
            "metadata": metadata,
        })
    return items


def _canonical(item: dict) -> str:
    """内容指纹（imported_at 不参与：保证重跑 no-op）。"""
    meta = {k: v for k, v in item["metadata"].items() if k != "imported_at"}
    payload = {"id": item["id"], "input": item["input"],
               "expected_output": item["expected_output"], "metadata": meta}
    return json.dumps(payload, ensure_ascii=False, sort_keys=True)


class LangfuseDatasetProvider:
    """Langfuse Dataset 读写（可注入替身用于测试）。"""

    def __init__(self, client) -> None:
        self._client = client

    def get_dataset_items(self, name: str) -> dict[str, dict] | None:
        """返回 {item_id: {input, expected_output, metadata}}；Dataset 不存在 → None。"""
        try:
            ds = self._client.get_dataset(name)
        except Exception as exc:  # noqa: BLE001 —— Dataset 不存在
            if "404" in str(exc) or "not found" in str(exc).lower():
                return None
            raise
        out: dict[str, dict] = {}
        for item in ds.items:
            out[item.id] = {
                "input": item.input,
                "expected_output": item.expected_output,
                "metadata": item.metadata or {},
            }
        return out

    def ensure_dataset(self, name: str, description: str, metadata: dict) -> None:
        self._client.create_dataset(
            name=name, description=description, metadata=metadata,
            input_schema={"type": "object", "additionalProperties": True},
            expected_output_schema={"type": "object", "additionalProperties": True},
        )

    def upsert_item(self, name: str, item: dict) -> None:
        self._client.create_dataset_item(
            dataset_name=name,
            id=item["id"],
            input=item["input"],
            expected_output=item["expected_output"],
            metadata=item["metadata"],
            status="ACTIVE",
        )


def export_dataset(
    *,
    items: list[dict],
    provider: LangfuseDatasetProvider | None,
    dataset_name: str,
    apply: bool,
) -> dict:
    """幂等导出：created / updated / unchanged / failed。provider=None → 纯 dry-run。"""
    existing: dict[str, dict] = {}
    dataset_exists = False
    if provider is not None:
        fetched = provider.get_dataset_items(dataset_name)
        existing = fetched or {}
        dataset_exists = fetched is not None  # 空 dataset 也存在

    stats = {"total_candidates": len(items), "created": 0, "updated": 0,
             "unchanged": 0, "failed": 0}
    for item in items:
        try:
            prev = existing.get(item["id"])
            if prev is None:
                stats["created"] += 1
                if apply and provider is not None:
                    if not dataset_exists:
                        provider.ensure_dataset(
                            dataset_name,
                            "人工确认（human review）的历史分类案例，用于正确率评估与 Prompt 实验",
                            {"purpose": "human_confirmed_classification", "version": DEFAULT_VERSION,
                             "source": "priority_review_pack"},
                        )
                        dataset_exists = True
                    provider.upsert_item(dataset_name, item)
            else:
                # 与既有内容比较（不含 imported_at）
                prev_item = {"id": item["id"], "input": prev["input"],
                             "expected_output": prev["expected_output"], "metadata": prev["metadata"]}
                if _canonical(item) == _canonical(prev_item):
                    stats["unchanged"] += 1
                else:
                    stats["updated"] += 1
                    if apply and provider is not None:
                        # 保留首次 imported_at
                        if prev["metadata"].get("imported_at"):
                            item["metadata"]["imported_at"] = prev["metadata"]["imported_at"]
                        provider.upsert_item(dataset_name, item)
        except Exception as exc:  # noqa: BLE001
            stats["failed"] += 1
            item["error"] = str(exc)[:200]
    return stats


def write_outputs(out_dir: Path, items: list[dict], report: dict) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "dataset_items.jsonl").write_text(
        "\n".join(json.dumps(i, ensure_ascii=False) for i in items) + "\n",
        encoding="utf-8",
    )
    (out_dir / "dataset_export_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Phase 2b dataset export (幂等)")
    parser.add_argument("--human-cases", required=True, help="Phase 2a human_cases.csv")
    parser.add_argument("--pack", default=None, help="priority_review_pack.xlsx（candidates 上下文）")
    parser.add_argument("--run-dir", required=True, help="输出报告目录（audit/evaluation/ 下）")
    parser.add_argument("--dataset-name", default=DEFAULT_DATASET_NAME)
    parser.add_argument("--dry-run", action="store_true", help="只生成报告与 items jsonl")
    parser.add_argument("--apply", action="store_true", help="真实写入 Langfuse Dataset")
    parser.add_argument("--verify", action="store_true", help="从 Langfuse API 回查")
    args = parser.parse_args(argv)
    if args.apply and args.dry_run:
        parser.error("--apply 与 --dry-run 互斥")

    human_cases = load_human_cases(Path(args.human_cases))
    pack_context = load_pack_context(Path(args.pack) if args.pack else None)
    items = build_items(human_cases, pack_context)
    no_audit = sum(1 for i in items if i["metadata"].get("audit_crosscheck") == "NO_AUDIT")

    provider = None
    mode = "dry_run"
    if args.apply:
        from langfuse import Langfuse

        client = Langfuse(
            public_key=os.getenv("LANGFUSE_PUBLIC_KEY", ""),
            secret_key=os.getenv("LANGFUSE_SECRET_KEY", ""),
            base_url=os.getenv("LANGFUSE_BASE_URL", "https://us.cloud.langfuse.com"),
            environment=os.getenv("LANGFUSE_TRACING_ENVIRONMENT", "development"),
            debug=False,
        )
        provider = LangfuseDatasetProvider(client)
        mode = "apply"

    imported_at = datetime.now(timezone.utc).isoformat()
    for item in items:
        item["metadata"]["imported_at"] = imported_at
    stats = export_dataset(items=items, provider=provider, dataset_name=args.dataset_name, apply=args.apply)
    stats["no_audit_crosscheck"] = no_audit

    report = {
        "generated_at": imported_at,
        "mode": mode,
        "dataset_name": args.dataset_name,
        "version": DEFAULT_VERSION,
        "schema_version": DEFAULT_SCHEMA_VERSION,
        "original_run_id": Path(args.human_cases).resolve().parent.parent.parent.name,
        **stats,
    }
    out_dir = Path(args.run_dir).resolve() / "audit" / "evaluation"
    write_outputs(out_dir, items, report)

    if args.verify and provider is None:
        from langfuse import Langfuse

        client = Langfuse(
            public_key=os.getenv("LANGFUSE_PUBLIC_KEY", ""),
            secret_key=os.getenv("LANGFUSE_SECRET_KEY", ""),
            base_url=os.getenv("LANGFUSE_BASE_URL", "https://us.cloud.langfuse.com"),
            environment=os.getenv("LANGFUSE_TRACING_ENVIRONMENT", "development"),
            debug=False,
        )
        provider = LangfuseDatasetProvider(client)

    if args.verify and provider is not None:
        existing = provider.get_dataset_items(args.dataset_name)
        ids = sorted(existing or {})
        no_audit_remote = sum(
            1 for it in (existing or {}).values()
            if (it.get("metadata") or {}).get("audit_crosscheck") == "NO_AUDIT"
        )
        overfilled = sum(
            1 for it in (existing or {}).values()
            if it.get("expected_output") and set(it["expected_output"]) - set(EXPECTED_OUTPUT_FIELDS)
        )
        dup = len(ids) != len(set(ids))
        report["verify"] = {
            "dataset_exists": existing is not None,
            "item_count": len(ids),
            "ids_unique": not dup,
            "no_audit_crosscheck_remote": no_audit_remote,
            "expected_output_overfilled": overfilled,
            "sample_ids": ids[:5],
        }

    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
