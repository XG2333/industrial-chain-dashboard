# -*- coding: utf-8 -*-
"""Phase 4：golden_coverage.py 测试（不触网、不改生产）。"""

import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from golden_coverage import (  # noqa: E402
    build_coverage_matrix,
    difficulty,
    family_id,
    load_audit_flags,
    load_catalog,
    quality_metrics,
    select_review_pool,
    write_review_pack,
)


def _var(name, industry="lithium", major="价格", sub="现货价格", candidates=1, vid=None):
    cands = [{"大类": major, "子类": sub}] * candidates if candidates else []
    return {"industry": industry, "physical_variable_id": vid or f"v{abs(hash(name)) % 16**16:016x}",
            "name": name, "sheet": "S", "unit": "元/吨", "frequency": "日度",
            "major": major, "sub": sub, "candidates": cands,
            "candidate_count": len(cands)}


def _golden(name, industry="lithium", major="价格", sub="现货价格", candidates=1,
            crosscheck="OK", current_major="价格"):
    cands = [{"大类": major, "子类": sub}] * max(candidates, 0) if candidates else []
    return {
        "id": f"golden_{industry}_{major}_{name}",
        "input": {"physical_variable_id": "x" * 16, "original_name": name,
                  "industry": industry, "commodity": "LI", "sheet": "S", "unit": "元/吨",
                  "frequency": "日度", "current_major": current_major,
                  "candidates": cands},
        "expected_output": {"大类": major, "子类": sub, "是否选中": "是"},
        "metadata": {"audit_crosscheck": crosscheck, "expected_in_candidates": True},
    }


# ── 1. coverage matrix ─────────────────────────────────────

def test_coverage_matrix():
    space = [_var("A"), _var("B", major="进出口", candidates=3), _var("C", industry="tin", major="供给")]
    golden = [_golden("A"), _golden("B", major="进出口", candidates=3)]
    m = build_coverage_matrix(space, golden)
    assert m["totals"]["production_variables"] == 3
    assert m["totals"]["golden_cases"] == 2
    cells = {c["industry"] + "|" + c["major_category"] + "|" + c["difficulty"]: c for c in m["cells"]}
    assert cells["lithium|价格|EASY"]["coverage_rate"] == 1.0
    assert cells["lithium|进出口|HARD"]["coverage_rate"] == 1.0
    assert cells["tin|供给|EASY"]["coverage_rate"] == 0.0


# ── 2. difficulty classification ───────────────────────────

def test_difficulty():
    assert difficulty(1, False) == "EASY"
    assert difficulty(2, False) == "MEDIUM"
    assert difficulty(3, False) == "HARD"
    assert difficulty(1, True) == "HARD"       # verifier flagged
    assert difficulty(1, False, human_corrected=True) == "HARD"
    assert difficulty(0, False) == "UNKNOWN"   # 无候选信息


# ── 3/4/5. P0-P3 priority + family dedup + cap ─────────────

def _pool_space():
    flags = {"vflagged": {"verdicts": ["ERROR"], "dispositions": ["CONFIRMED_ERROR"],
                          "issues": ["X"], "flagged": True}}
    space = [
        _var("SMM: 精炼锡进出口盈亏: 进口盈亏: 日度", industry="tin", major="成本利润",
             candidates=3, vid="vflagged"),          # P0（verifier flagged）
        _var("SMM: 精炼锡进出口盈亏: 出口盈亏: 日度", industry="tin", major="成本利润",
             candidates=3),                          # P1（3+ 候选，同 family）
        _var("SMM: 精炼锡进出口盈亏: 锡价: 日度", industry="tin", major="成本利润",
             candidates=3),                          # P1（同 family）
        _var("SMM: 普通价格A: 日度", major="价格", candidates=1),   # P3
        _var("SMM: 普通价格B: 日度", major="价格", candidates=1),   # P3
    ]
    return space, flags


def test_priority_and_family_dedup():
    space, flags = _pool_space()
    result = select_review_pool(space, flags, golden_vids=set(), pool_size=60, family_cap=2)
    pool = result["pool"]
    prios = {p["priority"] for p in pool}
    assert "P0" in prios and "P1" in prios and "P3" in prios
    # family cap=2：进出口盈亏 family 最多 2 条
    fam_count = sum(1 for p in pool if p["name"].startswith("SMM: 精炼锡进出口盈亏"))
    assert fam_count <= 2
    assert result["family_report"]["max_family_size"] <= 2
    assert result["priority_counts"]["P0"] == 1


def test_family_cap_zero_excludes():
    space, flags = _pool_space()
    result = select_review_pool(space, flags, golden_vids=set(), pool_size=60, family_cap=0)
    assert result["pool"] == []


# ── 6/7. human fields blank + verifier 不成为 expected ─────

def test_review_pack_human_fields_blank(tmp_path):
    space, flags = _pool_space()
    result = select_review_pool(space, flags, golden_vids=set(), pool_size=60, family_cap=2)
    path = tmp_path / "golden_review_pack.xlsx"
    write_review_pack(result["pool"], path)
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    i_dec, i_exp_m, i_exp_s, i_comment = (header.index("Human decision"),
                                          header.index("Expected major"),
                                          header.index("Expected sub"),
                                          header.index("Human comment"))
    for row in ws.iter_rows(min_row=2, values_only=True):
        assert row[i_dec] in (None, "")   # 绝不预填
        assert row[i_exp_m] in (None, "")
        assert row[i_exp_s] in (None, "")
        assert row[i_comment] in (None, "")
    # verifier 标记只出现在 verifier_status/reason 列，不进入 expected
    i_vstatus = header.index("verifier_status")
    i_vreason = header.index("verifier_reason")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    flagged_row = [r for r in rows if r[header.index("physical_variable_id")] == "vflagged"][0]
    assert "ERROR" in (flagged_row[i_vstatus] or "")
    assert "CONFIRMED_ERROR" in (flagged_row[i_vreason] or "")
    wb.close()


# ── 8. existing 25 cases preserved ─────────────────────────

def test_existing_cases_preserved(tmp_path):
    from golden_coverage import load_golden_items

    path = tmp_path / "items.jsonl"
    items = [_golden("A"), _golden("B", major="进出口", candidates=3)]
    path.write_text("\n".join(json.dumps(i, ensure_ascii=False) for i in items) + "\n",
                    encoding="utf-8")
    loaded = load_golden_items(path)
    assert len(loaded) == 2
    assert loaded[0]["expected_output"] == {"大类": "价格", "子类": "现货价格", "是否选中": "是"}
    assert loaded[0]["metadata"]["expected_in_candidates"] is True


# ── 9. only human-confirmed importable ─────────────────────

def test_only_human_confirmed_importable(tmp_path):
    # review pack 无 expected → 与 golden 判定隔离（quality_metrics 只计 golden）
    space, flags = _pool_space()
    pool = select_review_pool(space, flags, golden_vids=set(), pool_size=60, family_cap=2)["pool"]
    golden = [_golden("A")]
    q = quality_metrics(golden, space + pool)
    assert q["golden_total"] == 1  # pool 不计入 golden
    # 任何 pool 行都没有 expected_output 字段（人工未确认不可导入）
    assert all("expected_output" not in p for p in pool)


# ── 10. quality metrics ────────────────────────────────────

def test_quality_metrics():
    space = [_var("A", industry="lithium", major="价格"), _var("B", industry="tin", major="进出口", candidates=3),
             _var("C", industry="silicon", major="供给")]
    golden = [
        _golden("A", major="价格"),
        _golden("B", industry="tin", major="进出口", candidates=3, current_major="其他"),  # human corrected
        _golden("C", industry="silicon", major="供给", candidates=1, current_major="供给"),
    ]
    q = quality_metrics(golden, space)
    assert q["golden_total"] == 3
    assert q["industry_coverage"] == 1.0            # 3/3 产业
    assert q["major_category_coverage"] == 1.0      # 3/3 大类
    assert q["hard_case_ratio"] == round(1 / 3, 4)  # 1 HARD（3 候选 + corrected）
    assert q["multi_candidate_ratio"] == round(1 / 3, 4)
    assert q["human_corrected_ratio"] == round(1 / 3, 4)
    assert q["family_diversity"] == 1.0
    assert q["expected_in_candidates_rate"] == 1.0


# ── 11. production workflow unchanged ──────────────────────

def test_production_unchanged(tmp_path):
    import ai_assisted_curation as aac

    before = aac.AI_DISAMBIGUATION_SYSTEM_PROMPT
    # 目录/审计读取只读
    from openpyxl import Workbook

    xlsx = tmp_path / "catalog.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "指标目录"
    ws.append(["#", "Sheet Name", "Freq", "Col", "Indicator Name", "Unit", "板块",
               "大类", "子类", "数据性质", "是否选中", "状态说明", "指标标签"])
    ws.append([1, "S", "日度", 1, "SMM: 测试指标: 日度", "元/吨", "s", "价格", "现货价格",
               "水平值", "是", "", json.dumps({"候选分类组合": [{"大类": "价格", "子类": "现货价格"}]})])
    wb.save(xlsx)
    rows = load_catalog(xlsx, "lithium", "")
    assert rows[0]["candidate_count"] == 1
    assert rows[0]["major"] == "价格"
    assert aac.AI_DISAMBIGUATION_SYSTEM_PROMPT == before
    assert family_id("tin", "SMM: 精炼锡进出口盈亏: 进口盈亏: 日度") == "tin|SMM: 精炼锡进出口盈亏"


def test_audit_flags_parsing(tmp_path):
    audit = tmp_path / "audit.json"
    audit.write_text(json.dumps({"findings": [
        {"physical_variable_id": "a", "verdict": "ERROR", "final_disposition": "CONFIRMED_ERROR",
         "issue_type": "X"},
        {"physical_variable_id": "b", "verdict": "PASS", "final_disposition": "EXPECTED_TRANSFORMATION",
         "issue_type": "Y"},
    ]}), encoding="utf-8")
    flags = load_audit_flags(audit)
    assert flags["a"]["flagged"] is True
    assert flags["b"]["flagged"] is False
