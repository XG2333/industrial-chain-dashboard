# -*- coding: utf-8 -*-
"""为锡/硅工作簿生成"二次筛选是否保留"列（通用规则）。

与 recompute_secondary_only.py（锂电）不同：锡/硅 Excel 原本没有该列，
脚本会在目录表最右侧追加该列；规则使用 secondary_selection_decision_generic
（未选中 -> 否；频率非日/周度 -> 否；标题含"指数" -> 否）。
输出直接写回 data/ 目录（处理前自动备份 _backup_20260828.xlsx）。
"""
import shutil
from pathlib import Path

from openpyxl import load_workbook

from secondary_selection import secondary_selection_decision_generic

# 部署包相对路径：脚本位于 <包>/多skill联动/scripts/ → <包>/本地可视化dashboard/data
# （支持打包到任意根目录后在新电脑直接运行）
DATA_DIR = Path(__file__).resolve().parents[2] / "本地可视化dashboard" / "data"
FILES = ["锡产业链数据_workflow_ai.xlsx", "硅产业链数据_workflow_ai.xlsx"]
FINAL_COL = "二次筛选是否保留"


def main() -> int:
    for fname in FILES:
        p = DATA_DIR / fname
        if not p.exists():
            print("missing:", p)
            continue
        backup = DATA_DIR / (p.stem + "_backup_20260828.xlsx")
        shutil.copy2(p, backup)
        wb = load_workbook(p)
        ws = wb["指标目录"]
        header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        if FINAL_COL in header:
            final_idx = header.index(FINAL_COL)
        else:
            # 列不存在：追加到目录表最右侧
            final_idx = len(header)
            ws.cell(row=1, column=final_idx + 1, value=FINAL_COL)
        name_idx = header.index("Indicator Name")
        freq_idx = header.index("Freq")
        selected_idx = header.index("是否选中")
        changed = 0
        kept = 0
        for r_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            title = (
                str(row[name_idx].value).strip()
                if row[name_idx].value is not None
                else ""
            )
            freq = (
                str(row[freq_idx].value).strip()
                if row[freq_idx].value is not None
                else ""
            )
            selected = (
                str(row[selected_idx].value).strip()
                if row[selected_idx].value is not None
                else ""
            )
            keep, _reason = secondary_selection_decision_generic(
                title, freq, selected == "是"
            )
            new_val = "是" if keep else "否"
            cur = (
                str(row[final_idx].value).strip()
                if final_idx < len(row) and row[final_idx].value is not None
                else ""
            )
            if cur != new_val:
                changed += 1
            ws.cell(row=r_idx, column=final_idx + 1, value=new_val)
            if keep:
                kept += 1
        wb.save(p)
        print(
            f"{fname}: total={ws.max_row - 1}, kept={kept}, changed={changed}, "
            f"backup={backup.name}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
