# -*- coding: utf-8 -*-
"""Phase 2a evaluation correlation 测试（纯逻辑 + CLI dry-run，不触网不写 score）。"""

import hashlib
import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from evaluation_correlation import (  # noqa: E402
    MATCH_METHOD_AMBIGUOUS_TRIPLE,
    MATCH_METHOD_EXACT,
    MATCH_METHOD_UNMATCHED_TRIPLE,
    classify_evaluability,
    compute_dry_run,
    correlate_human_rows,
    correlate_items,
    evaluation_key,
    index_catalog,
    load_directory_rows,
    parse_generation_input,
    physical_variable_id,
    score_id,
)

RAW_HASH = "ab" * 32  # 64 hex
SHEET, COL, NAME, UNIT = "价格-日", 3, "SMM: 碳酸锂价格: 日度", "元/吨"
VID = physical_variable_id(RAW_HASH, SHEET, COL, NAME)

PROMPT_VERSION = "ai_assisted_curation_v1"
PROMPT_HASH = "h" * 64


def _catalog_rows():
    return [
        {"sheet": SHEET, "col": COL, "name": NAME, "unit": UNIT, "freq": "日度"},
        {"sheet": "进口-月", "col": 1, "name": "中国海关: 碳酸锂进口量: 总计: 月度", "unit": "吨", "freq": "月度"},
    ]


def _item(name=NAME, sheet=SHEET, unit=UNIT):
    return {"variable_id": "x" * 16, "name": name, "sheet": sheet, "unit": unit,
            "frequency": "日度", "sector": "", "current_major": "", "current_sub": ""}


def _gen_input(items):
    return [{"role": "system", "content": "sys"},
            {"role": "user", "content": json.dumps({"provider": "deepseek", "model": "m", "items": items})}]


def _corr(items, catalog=None, raw_hash=RAW_HASH):
    cat = catalog if catalog is not None else _catalog_rows()
    return correlate_items(
        workflow_run_id="run1234567890", industry="lithium", stage="ai_disambiguation",
        observation_id="obs1234567890abcdef", batch_index_start=0,
        prompt_name="ai_disambiguation", prompt_version=PROMPT_VERSION,
        prompt_hash=PROMPT_HASH, model="deepseek-v4-flash", items=items,
        catalog_index=index_catalog(cat), raw_file_hash=raw_hash,
    )


def _finding(vid, verdict="REVIEW_REQUIRED", disposition="NEEDS_HUMAN_REVIEW", issue="FREQUENCY_SOURCE_CONFLICT"):
    return {"finding_id": f"audit_x:V:{issue}:{vid}", "physical_variable_id": vid,
            "original_name": NAME, "verdict": verdict, "final_disposition": disposition,
            "issue_type": issue}


def _golden(vid, status="CONFIRMED", expected=None, rule_ids=None):
    return {"case_id": f"golden_lithium_LI_{vid}", "expected_status": status,
            "expected_category": (expected or {}).get("category"),
            "expected_subcategory": (expected or {}).get("subcategory"),
            "expected_selected": (expected or {}).get("selected"),
            "expected_rule_ids": rule_ids or ["AI_001"]}


def _human(vid, decision="CONFIRMED", major="价格", sub="现货价格"):
    return {"physical_variable_id": vid, "original_name": NAME, "human_decision": decision,
            "human_major": major, "human_sub": sub, "human_selected": "是"}


# ── 1. physical_variable_id 稳定性 ─────────────────────────

def test_physical_variable_id_stable():
    v1 = physical_variable_id(RAW_HASH, SHEET, COL, NAME)
    v2 = physical_variable_id(RAW_HASH, SHEET, COL, NAME)
    assert v1 == v2 == VID
    assert len(v1) == 16 and all(c in "0123456789abcdef" for c in v1)
    # 任一输入变化 → id 变化
    assert physical_variable_id(RAW_HASH, SHEET, COL, NAME + "x") != VID
    assert physical_variable_id(RAW_HASH, SHEET, COL + 1, NAME) != VID
    assert physical_variable_id("c" * 64, SHEET, COL, NAME) != VID
    # 与审计公式一致（已知向量）
    expected = hashlib.sha256(f"{RAW_HASH}|{SHEET}|{COL}|{NAME}".encode()).hexdigest()[:16]
    assert VID == expected


# ── 2. 多变量 → 单 generation 映射 ─────────────────────────

def test_multi_variable_single_generation():
    items = [_item(), _item(name="中国海关: 碳酸锂进口量: 总计: 月度", sheet="进口-月", unit="吨"), _item(name="X指标")]
    rows = _corr(items)
    obs_ids = {r["generation_observation_id"] for r in rows}
    assert obs_ids == {"obs1234567890abcdef"}  # 不拆成虚假 generation
    assert rows[0]["batch_index"] == 0 and rows[1]["batch_index"] == 1
    assert len(rows) == 3


# ── 3. exact run-local correlation ─────────────────────────

def test_exact_run_local_correlation():
    rows = _corr([_item()])
    assert len(rows) == 1
    assert rows[0]["match_method"] == MATCH_METHOD_EXACT
    assert rows[0]["physical_variable_id"] == VID
    assert rows[0]["catalog_col"] == COL


# ── 4. ambiguous 不评分 ────────────────────────────────────

def test_ambiguous_not_scored():
    cat = _catalog_rows() + [{"sheet": SHEET, "col": 9, "name": NAME, "unit": UNIT, "freq": "日度"}]
    rows = _corr([_item()], catalog=cat)
    assert rows[0]["match_method"] == MATCH_METHOD_AMBIGUOUS_TRIPLE
    assert rows[0]["physical_variable_id"] == ""  # 无身份 → 不可评分
    stats = compute_dry_run(correlation_rows=rows, findings_by_vid={}, golden_by_vid={}, human_by_vid={})
    assert stats["matched_generation_cases"] == 0
    assert stats["ambiguous_cases"] == 1
    assert stats["automatic_evaluable"] == 0 and stats["human_evaluable"] == 0


# ── 5. unmatched 不评分 ────────────────────────────────────

def test_unmatched_not_scored():
    rows = _corr([_item(name="不存在的指标")])
    assert rows[0]["match_method"] == MATCH_METHOD_UNMATCHED_TRIPLE
    assert rows[0]["physical_variable_id"] == ""
    stats = compute_dry_run(correlation_rows=rows, findings_by_vid={}, golden_by_vid={}, human_by_vid={})
    assert stats["unmatched_cases"] == 1
    assert stats["matched_generation_cases"] == 0


# ── 6. ERROR 不转 false ────────────────────────────────────

def test_error_verdict_not_derived_as_false():
    rows = _corr([_item()])
    row = rows[0]
    # 即使 golden 有明确期望输出 + 规则 approved，verdict=ERROR 也压制 correctness
    cls = classify_evaluability(
        row,
        [_finding(VID, verdict="ERROR", disposition="CONFIRMED_ERROR", issue="CLASSIFICATION_DRIFT")],
        _golden(VID, expected={"category": "价格", "subcategory": "现货价格"}),
        None,
    )
    assert cls["automatic_evaluable"] is False
    assert cls["automatic_not_evaluable_reason"] == "VERIFIER_FLAGGED_NEEDS_REVIEW"


# ── 7. NEEDS_HUMAN_REVIEW 不转 false ───────────────────────

def test_needs_human_review_not_derived():
    rows = _corr([_item()])
    cls = classify_evaluability(
        rows[0],
        [_finding(VID, verdict="REVIEW_REQUIRED", disposition="NEEDS_HUMAN_REVIEW")],
        _golden(VID, expected={"category": "价格", "subcategory": "现货价格"}),
        None,
    )
    assert cls["automatic_evaluable"] is False
    assert cls["automatic_not_evaluable_reason"] == "VERIFIER_FLAGGED_NEEDS_REVIEW"
    # PENDING_HUMAN_REVIEW golden 同样压制
    cls2 = classify_evaluability(
        rows[0], [], _golden(VID, status="PENDING_HUMAN_REVIEW", expected={"category": "价格", "subcategory": "现货价格"}), None,
    )
    assert cls2["automatic_evaluable"] is False


# ── 8. HUMAN > VERIFIER ────────────────────────────────────

def test_human_over_verifier():
    rows = _corr([_item()])
    cls = classify_evaluability(rows[0], [], None, _human(VID))
    assert cls["human_evaluable"] is True
    assert cls["human_not_evaluable_reason"] == ""
    # 自动侧被压制的 case 人工侧可评 → 人工优先
    cls2 = classify_evaluability(
        rows[0], [_finding(VID, verdict="ERROR")], None, _human(VID, decision="CORRECTED"),
    )
    assert cls2["human_evaluable"] is True
    assert cls2["automatic_evaluable"] is False
    # UNRESOLVED 不计 correctness
    cls3 = classify_evaluability(rows[0], [], None, _human(VID, decision="UNRESOLVED"))
    assert cls3["human_evaluable"] is False
    assert cls3["human_not_evaluable_reason"] == "HUMAN_UNRESOLVED"


# ── 9. evaluation_key 稳定 ─────────────────────────────────

def test_evaluation_key_stable():
    k1 = evaluation_key(VID, PROMPT_VERSION, PROMPT_HASH)
    k2 = evaluation_key(VID, PROMPT_VERSION, PROMPT_HASH)
    assert k1 == k2 and len(k1) == 32
    assert evaluation_key(VID, "v2", PROMPT_HASH) != k1
    assert evaluation_key(VID, PROMPT_VERSION, "g" * 64) != k1
    assert evaluation_key("0" * 16, PROMPT_VERSION, PROMPT_HASH) != k1


# ── 10. score_id 幂等 ──────────────────────────────────────

def test_score_id_idempotent():
    key = evaluation_key(VID, PROMPT_VERSION, PROMPT_HASH)
    s1 = score_id(key, "audit_1", "verifier", "verifier_status")
    s2 = score_id(key, "audit_1", "verifier", "verifier_status")
    assert s1 == s2 and len(s1) == 32
    assert score_id(key, "audit_2", "verifier", "verifier_status") != s1  # 不同 run 保留历史
    assert score_id(key, "audit_1", "human", "human_decision") != s1      # 来源隔离
    assert score_id(key, "audit_1", "verifier", "error_type") != s1


# ── 工具：解析 generation input ────────────────────────────

def test_parse_generation_input():
    items = [_item()]
    assert parse_generation_input(_gen_input(items)) == items
    assert parse_generation_input([{"role": "user", "content": "not json"}]) == []
    assert parse_generation_input(None) == []
    assert parse_generation_input("nope") == []


# ── 工具：目录表读取 + 人工恢复 ────────────────────────────

def test_load_directory_rows_and_human_correlation(tmp_path):
    from openpyxl import Workbook

    xlsx = tmp_path / "dir.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["#", "Sheet Name", "Freq", "Col", "Indicator Name", "Unit"])
    ws.append(["价格-日", None, None, None, None, None])
    ws.append([1, "价格-日", "日度", 3, NAME, UNIT])
    wb.save(xlsx)
    rows = load_directory_rows(xlsx)
    assert rows[0]["sheet"] == SHEET and rows[0]["col"] == COL and rows[0]["unit"] == UNIT

    human = correlate_human_rows(
        workflow_run_id="run20260818", industry="tin",
        human_rows=[{"case_id": "priority_tin_TIN_0", "original_name": NAME,
                     "human_decision": "CONFIRMED", "human_major": "价格", "human_sub": "现货价格",
                     "human_selected": "是", "human_notes": ""}],
        catalog_index=index_catalog(rows), raw_file_hash=RAW_HASH,
    )
    assert human[0]["match_method"] == "NAME_UNIQUE_RUN_LOCAL"
    assert human[0]["physical_variable_id"] == VID
    assert human[0]["human_decision"] == "CONFIRMED"


# ── 11/12. CLI dry-run：disabled Langfuse 不影响 + 不改 workflow 输出 ──

@pytest.fixture
def fixture_run(tmp_path):
    run = tmp_path / "20260831T060000Z_abcdef123456"
    (run / "audit" / "golden").mkdir(parents=True)
    (run / "02_financial_variable_curation").mkdir(parents=True)
    (run / "final").mkdir(parents=True)

    (run / "audit" / "audit_manifest.json").write_text(json.dumps(
        {"raw_file_hash": RAW_HASH, "industry": "lithium", "workflow_run_id": "20260831T060000Z_abcdef123456"},
        ensure_ascii=False), encoding="utf-8")
    (run / "audit" / "audit_report.json").write_text(json.dumps(
        {"manifest": {"audit_run_id": "audit_test"}, "findings": [_finding(VID)]},
        ensure_ascii=False), encoding="utf-8")
    (run / "audit" / "golden" / "golden_cases.jsonl").write_text(
        json.dumps(_golden(VID), ensure_ascii=False) + "\n", encoding="utf-8")

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["#", "Sheet Name", "Freq", "Col", "Indicator Name", "Unit"])
    ws.append(["价格-日", None, None, None, None, None])
    ws.append([1, "价格-日", "日度", COL, NAME, UNIT])
    wb.save(run / "02_financial_variable_curation" / "directory_marked.xlsx")

    final = run / "final" / "final_output.xlsx"
    wb2 = Workbook()
    wb2.active["A1"] = "业务输出"
    wb2.save(final)
    return run, final


def test_cli_dry_run_no_langfuse_and_output_untouched(fixture_run, monkeypatch):
    import evaluation_backfill

    run, final = fixture_run
    before = final.read_bytes()
    monkeypatch.delenv("LANGFUSE_PUBLIC_KEY", raising=False)
    monkeypatch.delenv("LANGFUSE_SECRET_KEY", raising=False)

    rc = evaluation_backfill.main([
        "--run-root", str(run.parent), "--run", run.name, "--no-langfuse",
    ])
    assert rc == 0

    eval_dir = run / "audit" / "evaluation"
    assert (eval_dir / "correlation_index.jsonl").exists()
    assert (eval_dir / "evaluation_dry_run.json").exists()
    assert (eval_dir / "evaluation_dry_run.csv").exists()
    assert (eval_dir / "unmatched_cases.csv").exists()
    assert (eval_dir / "ambiguous_cases.csv").exists()
    summary = json.loads((run.parent / "evaluation_dry_run_summary.json").read_text(encoding="utf-8"))
    assert summary["mode"] == "dry_run"
    assert summary["runs"][0]["trace_status"] == "LANGFUSE_DISABLED"
    # 人工/自动均不可评（无 generation），但不崩溃
    assert summary["totals"]["automatic_evaluable"] == 0
    assert summary["totals"]["human_evaluable"] == 0
    # workflow 输出未被修改
    assert final.read_bytes() == before
    # 未写入任何 score 文件（无 langfuse score 产物）
    assert not (eval_dir / "scores").exists()
