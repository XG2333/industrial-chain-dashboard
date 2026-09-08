# -*- coding: utf-8 -*-
"""Apply industry-specific deterministic selection rules after Skill2.

Skill2 first runs the common selection rules from selection_rules.docx. This
script then applies the industry-specific rule module to override 是否选中.
"""

from __future__ import annotations

import argparse
import importlib
import json
import re
import shutil
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parent))
from import_export_rules import IE_MAJORS, deduplicate_import_export_selection
from selection_utils import (
    apply_common_price_rules,
    enforce_composite_pairing,
    exclude_import_export_amount_rows,
    exclude_import_export_average_rows,
    exclude_yoy_mom_share_rows,
    ensure_unique_selected_rows,
    ensure_volume_position_pairing,
    remove_confidence_column,
)
from secondary_selection import apply_secondary_selection
from volume_position_ratio import append_missing_volume_position_ratios

INDUSTRY_MODULES = {
    "lithium": "lithium_selection_rules",
    "tin": "tin_rules",
    "silicon": "silicon_selection_rules",
}


def is_indicator_row(values: list) -> bool:
    if not values or values[0] is None:
        return False
    raw = str(values[0]).strip()
    return bool(re.fullmatch(r"\d+", raw)) and bool(values[3]) and bool(values[4])


def parse_directory_tags(values: list) -> dict:
    if len(values) <= 13 or not values[13]:
        return {}
    try:
        payload = json.loads(str(values[13]))
    except (TypeError, ValueError):
        return {}
    return payload if isinstance(payload, dict) else {}


def normalize_supply_label(value):
    if isinstance(value, str) and value == "供应":
        return "供给"
    if isinstance(value, list):
        return [normalize_supply_label(item) for item in value]
    if isinstance(value, dict):
        return {key: normalize_supply_label(item) for key, item in value.items()}
    return value



def assert_import_export_balance(rows: list[dict]) -> dict[str, int]:
    counts = {"进口": 0, "出口": 0, "净出口": 0}
    for row in rows:
        if row.get("major") in IE_MAJORS and row.get("selected"):
            sub = row.get("sub")
            if sub in counts:
                counts[sub] += 1
    if len(set(counts.values())) != 1:
        raise RuntimeError(f"进出口三子类选中数量不一致: {counts}")
    return counts


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Apply industry-specific deterministic selection.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--industry", choices=list(INDUSTRY_MODULES), required=True)
    parser.add_argument("--report", default=None)
    parser.add_argument(
        "--data-file",
        default=None,
        help="完整数据文件（含成交量/持仓量等数据 sheets）。仅用于成交持仓比计算的"
        "数据读取（read_only），写入仍在 --input 文件。默认与 --input 相同。",
    )
    args = parser.parse_args(argv)

    input_path = Path(args.input).resolve()
    output_path = Path(args.output).resolve()
    same_path = input_path == output_path
    working_output = output_path
    if same_path:
        working_output = output_path.with_name(f"{output_path.stem}__selection_tmp.xlsx")
    report_path = Path(args.report).resolve() if args.report else output_path.with_suffix(".md")
    working_output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(input_path, working_output)

    wb = load_workbook(working_output)
    ws = wb.worksheets[0]
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    if "置信度" not in headers:
        ws.insert_cols(10)
        ws.cell(1, 10, "置信度")
    current_sheet = ""
    records = []

    for row in ws.iter_rows(min_row=2):
        values = [c.value for c in row]
        col0 = str(values[0]).strip() if values[0] is not None else ""
        if col0.startswith("[统计]"):
            continue
        if not is_indicator_row(values):
            if col0 and not re.fullmatch(r"\d+", col0) and len(values) > 6 and values[6] is not None:
                current_sheet = col0
            continue

        records.append(
            {
                "sheet": current_sheet,
                "title": str(values[4]),
                "unit": str(values[5] or "").strip(),
                "sector": str(values[6] or "").strip(),
                "frequency": str(values[2] or "").strip(),
                "major": str(values[7] or "").strip(),
                "sub": str(values[8] or "").strip(),
                "nature": str(values[9] or "").strip(),
                "tags": parse_directory_tags(values),
                "_row": row,
            }
        )

    module = importlib.import_module(INDUSTRY_MODULES[args.industry])
    processed = module.select_rows(records)
    apply_common_price_rules(processed)
    if args.industry == "lithium":
        for rec in processed:
            if "电解铜" in str(rec.get("title") or "") or "电解铜" in str(rec.get("sheet") or ""):
                rec["selected"] = False
                rec["reason"] = "锂电专属规则：电解铜指标不保留"
    exclude_import_export_average_rows(processed)
    exclude_import_export_amount_rows(processed)
    exclude_yoy_mom_share_rows(processed)
    deduplicate_import_export_selection(processed)
    ensure_volume_position_pairing(processed)
    enforce_composite_pairing(processed)
    ensure_unique_selected_rows(processed)
    apply_secondary_selection(processed)
    ie_counts = assert_import_export_balance(processed)

    headers_now = [str(cell.value or "").strip() for cell in ws[1]]
    if "二次筛选是否保留" in headers_now:
        secondary_col = headers_now.index("二次筛选是否保留") + 1
    else:
        secondary_col = ws.max_column + 1
        ws.cell(1, secondary_col, "二次筛选是否保留")

    selected = 0
    reason_counter: Counter[str] = Counter()
    for rec in processed:
        keep = bool(rec.get("selected"))
        if keep:
            selected += 1
        original_reason = str(rec["_row"][12].value or "").strip()
        reason = rec.get("reason") or ""
        if rec.get("sub") == "成交持仓比" and original_reason.startswith("确定性成交持仓比计算"):
            reason = original_reason
        rec["_row"][11].value = "是" if keep else "否"
        rec["_row"][12].value = reason
        ws.cell(rec["_row"][0].row, secondary_col, "是" if rec.get("secondary_keep") else "否")
        reason_counter[reason or "未覆盖"] += 1

    for row in ws.iter_rows(min_row=2):
        if row[7].value == "供应":
            row[7].value = "供给"
        if row[8].value == "供应":
            row[8].value = "供给"
        raw_tags = row[13].value if len(row) > 13 else None
        if not raw_tags:
            continue
        try:
            tags = json.loads(str(raw_tags))
        except Exception:
            continue
        normalized_tags = normalize_supply_label(tags)
        if normalized_tags != tags:
            row[13].value = json.dumps(normalized_tags, ensure_ascii=False)

    data_workbook = None
    if args.data_file:
        data_path = Path(args.data_file).resolve()
        if not data_path.exists():
            raise SystemExit(f"--data-file does not exist: {data_path}")
        data_workbook = load_workbook(data_path, read_only=True, data_only=False)
    try:
        ratio_count = append_missing_volume_position_ratios(
            wb, processed, secondary_col=secondary_col, data_workbook=data_workbook
        )
    finally:
        if data_workbook is not None:
            data_workbook.close()
    remove_confidence_column(wb)
    wb.save(working_output)
    if same_path:
        working_output.replace(output_path)

    report_lines = [
        "# 产业专属筛选审计报告",
        "",
        f"- 产业：{args.industry}",
        f"- 输入：`{input_path.name}`",
        f"- 输出：`{output_path.name}`",
        f"- 指标总数：{len(processed)}",
        f"- 选中指标数：{selected}",
        f"- 未选中指标数：{len(processed) - selected}",
        f"- 进出口选中：进口 {ie_counts['进口']} / 出口 {ie_counts['出口']} / 净出口 {ie_counts['净出口']}",
        "",
        "## 筛选原因统计",
        "",
    ]
    for reason, count in reason_counter.most_common():
        report_lines.append(f"- {reason}: {count}")
    report_path.write_text("\n".join(report_lines), encoding="utf-8")

    print(f"Industry={args.industry} total={len(processed)} selected={selected} rejected={len(processed) - selected}")
    print(f"Output: {output_path}")
    print(f"Report: {report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
