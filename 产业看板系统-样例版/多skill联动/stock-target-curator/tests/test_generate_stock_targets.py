# -*- coding: utf-8 -*-

from __future__ import annotations

import sys
import tempfile
import shutil
from pathlib import Path

from openpyxl import load_workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

from generate_stock_targets import (  # noqa: E402
    HEADERS,
    build_rows,
    check_segment_limits,
    load_config,
    write_xlsx,
)


def _config() -> dict:
    return load_config(SKILL_ROOT / "config" / "stock_targets.json")


def test_build_rows_returns_expected_columns() -> None:
    rows = build_rows(_config(), "lithium")

    assert rows
    assert set(rows[0]) == {"industry", "segment", "stage", "name", "code", "cap", "reason"}


def test_rows_sorted_by_market_cap_desc_within_segment() -> None:
    rows = build_rows(_config(), "lithium")
    segments = {row["segment"] for row in rows}
    assert segments

    for segment in segments:
        caps = [row["cap"] for row in rows if row["segment"] == segment]
        assert caps == sorted(caps, reverse=True), segment


def test_market_cap_missing_falls_back_to_zero() -> None:
    # 缺失 market_cap 的股票按 0 处理（排在板块最后），不抛异常
    rows = build_rows(
        {
            "industries": {"lithium": {"name": "锂电池产业链", "segments": ["锂盐"]}},
            "stocks": [
                {"industry": "lithium", "segment": "锂盐", "code": "000001", "name": "A"},
                {"industry": "lithium", "segment": "锂盐", "code": "000002", "name": "B", "market_cap": 100},
                {"industry": "lithium", "segment": "锂盐", "code": "000003", "name": "C", "market_cap": "bad"},
            ],
        },
        "lithium",
    )
    assert [row["code"] for row in rows] == ["000002", "000001", "000003"]


def test_build_rows_deduplicates_segment_and_code() -> None:
    rows = build_rows(_config(), "silicon")
    keys = [(row["segment"], row["code"]) for row in rows]

    assert len(keys) == len(set(keys))
    assert any(row["segment"] == "多晶硅" for row in rows)
    assert any(row["code"] == "600438" for row in rows)


def test_write_xlsx_creates_stock_target_sheet() -> None:
    temp_dir = Path(tempfile.mkdtemp(prefix="skill3_", dir=SKILL_ROOT / "tests"))
    try:
        rows = build_rows(_config(), "tin")
        output = write_xlsx(rows, temp_dir / "锡产业链_个股标的.xlsx")

        workbook = load_workbook(output, read_only=True)
        ws = workbook["个股标的"]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        assert headers == HEADERS
        assert ws.max_row == len(rows) + 1
        workbook.close()
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def test_all_segments_meet_minimum_requirements() -> None:
    config = _config()
    for industry in ("silicon", "tin", "lithium"):
        rows = build_rows(config, industry)
        assert check_segment_limits(config, industry, rows) == []
