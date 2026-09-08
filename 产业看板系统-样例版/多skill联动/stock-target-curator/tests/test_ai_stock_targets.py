# -*- coding: utf-8 -*-
"""AI 个股筛选测试：scripts/ai_stock_targets.py。"""

from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

from ai_stock_targets import (  # noqa: E402
    build_prompt,
    load_config,
    mock_selection,
    normalize_stock,
    segment_candidates,
    select_for_segment,
    init_cache,
)
from generate_stock_targets import HEADERS  # noqa: E402


def _config() -> dict:
    return load_config(SKILL_ROOT / "config" / "stock_targets.json")


def test_segments_merged_lithium_ore_and_salt() -> None:
    config = _config()
    segments = config["industries"]["lithium"]["segments"]
    assert "锂矿锂盐" in segments
    assert "锂矿" not in segments
    assert "锂盐" not in segments


def test_mock_selection_returns_candidates_with_reason() -> None:
    config = _config()
    candidates = segment_candidates(config, "lithium", "锂矿锂盐")
    assert candidates
    stocks = mock_selection(candidates)
    codes = {s["code"] for s in stocks}
    assert len(stocks) == len(codes)  # 去重
    assert all(s["reason"].startswith("候选池参考") for s in stocks)
    # 交叉股票（原锂矿+锂盐重复的）只保留一条
    assert len(stocks) <= len(candidates)


def test_normalize_stock_validates_code_and_cap() -> None:
    candidates = {"002466": {"stage": "中游-锂盐", "note": "锂矿及锂盐"}}
    ok = normalize_stock(
        {"name": "天齐锂业", "code": "002466", "market_cap": 700, "reason": "锂资源龙头，拥有格林布什锂辉石矿"},
        candidates,
    )
    assert ok is not None
    assert ok["code"] == "002466"
    assert ok["market_cap"] == 700.0
    assert ok["stage"] == "中游-锂盐"

    # 非法代码/缺失名称被丢弃
    assert normalize_stock({"name": "X", "code": "12", "market_cap": 1, "reason": "r"}, candidates) is None
    assert normalize_stock({"name": "", "code": "002466", "market_cap": 1, "reason": "r"}, candidates) is None
    # 原因过短（<8字）视为无依据，丢弃
    assert normalize_stock({"name": "Z", "code": "002466", "market_cap": 1, "reason": "龙头"}, candidates) is None
    # 非数值市值回退 0
    bad = normalize_stock(
        {"name": "Y", "code": "002466", "market_cap": "nan", "reason": "锂盐产能行业前列，资源自给率高"},
        candidates,
    )
    assert bad is not None and bad["market_cap"] == 0.0


def test_build_prompt_contains_description_and_candidates() -> None:
    config = _config()
    meta = {**config["industries"]["lithium"], "key": "lithium"}
    candidates = segment_candidates(config, "lithium", "三元正极")
    prompt = build_prompt(meta, "三元正极", candidates, 5, 20)
    assert "三元前驱体" in prompt
    assert "20 只" in prompt and "5 只" in prompt
    assert "stocks" in prompt
    assert candidates[0]["name"] in prompt
    # 覆盖全面 + 有理有据：接近上限、覆盖二线/细分公司、原因必须含事实依据
    assert "覆盖尽可能全面" in prompt
    assert "尽量接近 20 只" in prompt
    assert "第二梯队" in prompt
    assert "具体事实依据" in prompt


def test_select_for_segment_mock_uses_cache() -> None:
    config = _config()
    tmp = Path(tempfile.mkdtemp(prefix="ai_stocks_", dir=SKILL_ROOT / "tests"))
    try:
        con = init_cache(tmp / "cache.db")
        first = select_for_segment(config, "lithium", "隔膜", "mock", con, False)
        second = select_for_segment(config, "lithium", "隔膜", "mock", con, False)
        con.close()
        assert first and second
        assert [s["code"] for s in first] == [s["code"] for s in second]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_run_mock_writes_merged_segment_xlsx() -> None:
    from ai_stock_targets import run

    tmp = Path(tempfile.mkdtemp(prefix="ai_stocks_run_", dir=SKILL_ROOT / "tests"))
    try:
        out = tmp / "个股标的_锂电.xlsx"
        rc = run("lithium", "mock", SKILL_ROOT / "config" / "stock_targets.json",
                 out, tmp / "cache.db", False)
        assert rc == 0
        wb = load_workbook(out, read_only=True, data_only=True)
        ws = wb["个股标的"]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        assert list(rows[0]) == HEADERS
        segments = {r[2] for r in rows[1:]}
        assert "锂矿锂盐" in segments
        assert "锂矿" not in segments and "锂盐" not in segments
        assert len(rows) > 1
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
