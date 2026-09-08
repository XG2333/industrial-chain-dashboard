from __future__ import annotations

import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parent))
from selection_utils import ensure_unique_selected_rows, ensure_volume_position_pairing
from silicon_selection_rules import select_rows


def apply_selection(path: Path) -> None:
    path = Path(path).resolve()
    tmp = path.with_name(f"{path.stem}__tmp_sel.xlsx")
    shutil.copy2(path, tmp)

    wb = load_workbook(tmp)
    ws = wb[wb.sheetnames[0]]
    current_sheet = ""
    records = []

    for row in ws.iter_rows(min_row=2):
        c0 = str(row[0].value).strip() if row[0].value is not None else ""
        if not c0:
            continue
        if "sht" in c0:
            continue
        if not c0.isdigit():
            current_sheet = c0
            continue
        c3 = str(row[3].value).strip() if row[3].value is not None else ""
        if not c3.isdigit():
            continue
        title = str(row[4].value).strip() if row[4].value is not None else ""
        if not title:
            continue
        records.append(
            {
                "sheet": current_sheet,
                "title": title,
                "frequency": str(row[2].value).strip() if row[2].value is not None else "",
                "major": str(row[7].value).strip() if row[7].value is not None else "",
                "sub": str(row[8].value).strip() if row[8].value is not None else "",
                "_row": row,
            }
        )

    processed = select_rows(records)
    ensure_volume_position_pairing(processed)
    ensure_unique_selected_rows(processed)
    changed = 0
    for rec in processed:
        row = rec["_row"]
        old = str(row[11].value).strip() if row[11].value is not None else ""
        new = "是" if rec["selected"] else "否"
        if old != new:
            changed += 1
        row[11].value = new
        row[12].value = rec.get("reason", "")

    wb.save(tmp)
    wb.close()
    try:
        tmp.replace(path)
    except PermissionError:
        print(f"Permission error writing {path}; temp file kept at {tmp}")
        raise
    print(f"Applied silicon selection rules to {path}: total={len(processed)}, changed={changed}")


if __name__ == "__main__":
    paths = [Path(arg) for arg in sys.argv[1:] if not arg.startswith("-")]
    if not paths:
        print("Usage: python apply_silicon_selection.py <directory_marked.xlsx>")
        sys.exit(1)
    for path in paths:
        apply_selection(path)
