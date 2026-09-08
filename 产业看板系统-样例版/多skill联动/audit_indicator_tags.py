from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from apply_indicator_tags import extract_indicator_tags


RESEARCH_BY_MAJOR = {
    "价格": "价格",
    "成本利润": "成本利润",
    "库存": "库存",
    "供给": "供给",
    "需求": "需求",
    "进出口": "进出口",
    "平衡": "平衡",
    "其他": "其他",
}

TYPE_BY_SUB = {
    "现货价格": "现货价格",
    "期货价格": "期货价格",
    "基差": "基差",
    "月差": "月差",
    "现货价差": "现货价差",
    "持仓": "持仓",
    "成交量": "成交量",
    "成交持仓比": "成交持仓比",
    "成本": "成本",
    "利润": "利润",
    "库存": "库存",
    "仓单": "仓单",
    "库存天数": "库存天数",
    "库存指数": "库存指数",
    "产能": "产能",
    "产量": "产量",
    "开工率": "开工率",
    "数量": "数量",
    "需求": "需求",
    "销量": "销量",
    "进出口": "进出口",
    "平衡": "平衡",
    "宏观": "宏观",
    "参数": "参数",
    "其他": "其他",
}

STATUS_TOKENS = ("预测", "计划", "目标", "预算", "估算")


def _expected_type(name: str, sub: str) -> str:
    if sub == "进出口":
        if "净出口" in name:
            return "净出口"
        if "出口" in name and "进口" not in name:
            return "出口"
        if "进口" in name and "出口" not in name:
            return "进口"
        return "进出口"
    if sub in ("价差", "现货价差"):
        if "基差" in name or "期现" in name or "升贴水" in name:
            return "基差"
        if "价差" in name or "溢价" in name:
            return "现货价差"
    return TYPE_BY_SUB.get(sub, sub)


def _expected_status(name: str, sheet: str) -> str:
    text = f"{name} {sheet}"
    for status in STATUS_TOKENS:
        if status in text:
            return status
    return "实际"


def _split_tag(value: str) -> list[str]:
    return [part.strip() for part in str(value).split(",") if part.strip()]


def main() -> None:
    parser = argparse.ArgumentParser(description="Audit indicator tags against directory rows.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--report", required=True)
    args = parser.parse_args()

    path = Path(args.input).resolve()
    report_path = Path(args.report).resolve()
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    if "指标标签" not in headers:
        raise SystemExit(f"指标标签 column not found in {path}")
    tag_col = headers.index("指标标签")

    issues: list[tuple[str, str, str, str]] = []
    distributions: dict[str, Counter] = {key: Counter() for key in ("产品", "研究主题", "指标类型", "规格", "工艺属性", "地域", "统计口径", "状态", "频率")}
    total = 0
    current_sheet = ""

    for row in rows[1:]:
        c0 = str(row[0]).strip() if row[0] is not None else ""
        if not c0 or "sht" in c0:
            continue
        if not c0.isdigit():
            current_sheet = c0
            continue
        c3 = str(row[3]).strip() if row[3] is not None else ""
        if not c3.isdigit():
            continue
        name = str(row[4]).strip() if row[4] is not None else ""
        if not name:
            continue
        unit = str(row[5]).strip() if row[5] is not None else ""
        freq = str(row[2]).strip() if row[2] is not None else ""
        major = str(row[7]).strip() if row[7] is not None else ""
        sub = str(row[8]).strip() if row[8] is not None else ""
        raw_tags = str(row[tag_col]).strip() if row[tag_col] is not None else ""
        total += 1

        try:
            tags = json.loads(raw_tags)
        except (json.JSONDecodeError, TypeError):
            issues.append(("JSON异常", current_sheet, name, raw_tags[:200]))
            continue
        if not isinstance(tags, dict):
            issues.append(("结构异常", current_sheet, name, raw_tags[:200]))
            continue

        for dim in distributions:
            distributions[dim][str(tags.get(dim, "未识别"))] += 1

        expected = extract_indicator_tags(
            current_sheet,
            name,
            unit,
            freq,
            major,
            sub,
        )
        for key, expected_value in expected.items():
            actual_value = str(tags.get(key, ""))
            if actual_value != expected_value:
                issues.append(
                    (
                        f"{key}与规则不一致",
                        current_sheet,
                        name,
                        f"{actual_value} != {expected_value}",
                    )
                )

    lines = [
        "# 硅产业链指标标签审计报告",
        "",
        f"- 输入文件：`{path.name}`",
        f"- 指标数：{total}",
        f"- 问题数：{len(issues)}",
        "",
        "## 问题分类",
        "",
    ]
    by_category: Counter = Counter(item[0] for item in issues)
    if by_category:
        for category, count in by_category.most_common():
            lines.append(f"- {category}：{count}")
    else:
        lines.append("- 无")

    lines += ["", "## 问题明细（最多 200 条）", ""]
    for category, sheet, name, detail in issues[:200]:
        lines.append(f"- {category} | {sheet} | {name} | {detail}")

    lines += ["", "## 标签值分布", ""]
    for dim, counter in distributions.items():
        lines += ["", f"### {dim}", ""]
        for value, count in counter.most_common():
            lines.append(f"- {value}：{count}")

    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"Audit complete: {path}, indicators={total}, issues={len(issues)}")


if __name__ == "__main__":
    main()
