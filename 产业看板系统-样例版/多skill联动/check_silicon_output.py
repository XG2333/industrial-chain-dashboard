from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parent))
from reclassify_silicon import classify_row, is_price_unit


EXPECTED_HEADERS = [
    "#",
    "Sheet Name",
    "Freq",
    "Col",
    "Indicator Name",
    "Unit",
    "板块",
    "大类",
    "子类",
    "数据性质",
    "是否选中",
    "状态说明",
    "指标标签",
]

ALLOWED_PAIRS = {
    ("价格", "现货价格"),
    ("价格", "期货价格"),
    ("价格", "月差"),
    ("价格", "现货价差"),
    ("价格", "持仓"),
    ("价格", "成交量"),
    ("价格", "成交持仓比"),
    ("成本利润", "成本"),
    ("成本利润", "利润"),
    ("库存", "库存"),
    ("库存", "仓单"),
    ("库存", "库存天数"),
    ("库存", "库存指数"),
    ("供给", "产能"),
    ("供给", "产量"),
    ("供给", "开工率"),
    ("供给", "数量"),
    ("需求", "需求"),
    ("需求", "销量"),
    ("进出口", "进出口"),
    ("平衡", "平衡"),
    ("其他", "宏观"),
    ("其他", "参数"),
    ("其他", "其他"),
}


def main() -> None:
    parser = argparse.ArgumentParser(description="Check silicon directory classification and selection status.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--report", required=True)
    args = parser.parse_args()

    path = Path(args.input).resolve()
    report_path = Path(args.report).resolve()
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    issues: list[str] = []
    total = 0
    current_sheet = ""
    major_sub = Counter()
    tag_coverage = Counter()

    for r in rows[1:]:
        c0 = str(r[0]).strip() if r[0] is not None else ""
        if not c0 or "sht" in c0:
            continue
        if not c0.isdigit():
            current_sheet = c0
            continue
        c3 = str(r[3]).strip() if r[3] is not None else ""
        if not c3.isdigit():
            continue

        total += 1
        name = str(r[4]).strip() if r[4] is not None else ""
        unit = str(r[5]).strip() if r[5] is not None else ""
        freq = str(r[2]).strip() if r[2] is not None else ""
        major = str(r[7]).strip() if r[7] is not None else ""
        sub = str(r[8]).strip() if r[8] is not None else ""
        selected = str(r[11]).strip() if r[11] is not None else ""
        raw_tags = str(r[13]).strip() if len(r) > 13 and r[13] is not None else ""
        major_sub[(major, sub)] += 1

        if name == "指标名称" and unit == "单位":
            issues.append(f"伪指标行：{current_sheet} / {name} / {unit}")
            continue

        if selected not in ("是", "否"):
            issues.append(f"是否选中异常：{current_sheet} / {name} / {selected}")

        if not raw_tags:
            issues.append(f"指标标签缺失：{current_sheet} / {name}")
        else:
            try:
                tags = json.loads(raw_tags)
            except json.JSONDecodeError:
                issues.append(f"指标标签JSON异常：{current_sheet} / {name} / {raw_tags[:120]}")
                tags = None
            if not isinstance(tags, dict):
                issues.append(f"指标标签结构异常：{current_sheet} / {name}")
            else:
                for required_key in ("产品", "研究主题", "指标类型", "状态", "频率"):
                    if not tags.get(required_key):
                        issues.append(
                            f"指标标签缺少{required_key}：{current_sheet} / {name}"
                        )
                if tags.get("频率") and tags["频率"] != freq:
                    issues.append(
                        f"指标标签频率不一致：{current_sheet} / {name} / "
                        f"{tags['频率']} != {freq}"
                    )
                for key, value in tags.items():
                    if value not in (None, "", "未识别"):
                        tag_coverage[key] += 1

        if (major, sub) not in ALLOWED_PAIRS:
            issues.append(f"非法分类组合：{current_sheet} / {name} / {major} / {sub}")

        expected = classify_row(current_sheet, name, unit, major, sub)
        if (major, sub) != expected:
            issues.append(
                f"分类与规则不一致：{current_sheet} / {name} / {unit} / "
                f"{major}/{sub} -> {expected[0]}/{expected[1]}"
            )

        if major in ("需求", "进出口", "库存", "平衡") and is_price_unit(unit):
            issues.append(f"数量类单位冲突：{current_sheet} / {name} / {unit} / {major}/{sub}")

        if major == "价格" and sub in ("现货价格", "期货价格", "现货价差"):
            if not is_price_unit(unit):
                issues.append(f"价格类单位冲突：{current_sheet} / {name} / {unit} / {major}/{sub}")

    if headers != EXPECTED_HEADERS:
        issues.append(f"目录表头异常：{headers}")

    report_path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# 硅产业链输出检查报告",
        "",
        f"- 输入文件：`{path.name}`",
        f"- 有效指标数：{total}",
        f"- 问题数：{len(issues)}",
        "",
        "## 分类分布",
        "",
    ]
    for (major, sub), count in major_sub.most_common():
        lines.append(f"- {major}/{sub}：{count}")
    lines += ["", "## 指标标签覆盖", ""]
    for key in (
        "产品",
        "研究主题",
        "指标类型",
        "规格",
        "工艺属性",
        "地域",
        "统计口径",
        "状态",
        "频率",
    ):
        lines.append(f"- {key}：{tag_coverage[key]}/{total}")
    lines += ["", "## 问题明细", ""]
    if issues:
        lines.extend(f"- {issue}" for issue in issues[:200])
    else:
        lines.append("- 无")
    report_path.write_text("\n".join(lines), encoding="utf-8")

    if issues:
        print(f"Check failed: {path}, issues={len(issues)}")
        raise SystemExit(1)
    print(f"Check passed: {path}, indicators={total}, issues=0")


if __name__ == "__main__":
    main()
