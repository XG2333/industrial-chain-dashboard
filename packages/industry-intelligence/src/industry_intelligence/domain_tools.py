from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

from .config import IntelligenceSettings
from .rag import HybridRetriever


@dataclass(frozen=True)
class ToolSpec:
    name: str
    description: str
    parameters: dict[str, Any]
    handler: Callable[..., dict[str, Any]]


class DomainToolRegistry:
    """Allow-listed, read-only tools shared by model calls and MCP."""

    def __init__(self, settings: IntelligenceSettings, retriever: HybridRetriever, store) -> None:
        self.settings = settings
        self.retriever = retriever
        self.store = store
        self._tools = {tool.name: tool for tool in self._build_tools()}

    def _build_tools(self) -> list[ToolSpec]:
        return [
            ToolSpec(
                name="search_industry_knowledge",
                description=(
                    "Search approved project rules and documentation. Use this for definitions, "
                    "classification logic, calculation rules, and processing-flow explanations."
                ),
                parameters={
                    "type": "object",
                    "properties": {
                        "query": {"type": "string", "description": "Precise question or search phrase."},
                        "industry": {
                            "type": "string",
                            "enum": ["", "lithium", "tin", "silicon"],
                            "description": "Optional industry filter.",
                        },
                        "limit": {"type": "integer", "minimum": 1, "maximum": 10},
                    },
                    "required": ["query", "industry", "limit"],
                    "additionalProperties": False,
                },
                handler=self._search_knowledge,
            ),
            ToolSpec(
                name="list_rule_documents",
                description="List indexed business-rule documents and their source paths.",
                parameters={
                    "type": "object",
                    "properties": {
                        "industry": {
                            "type": "string",
                            "enum": ["", "lithium", "tin", "silicon"],
                        }
                    },
                    "required": ["industry"],
                    "additionalProperties": False,
                },
                handler=self._list_rules,
            ),
            ToolSpec(
                name="get_indicator_catalog",
                description=(
                    "Read only the catalog worksheet from a processed industry workbook. "
                    "It never writes to the workbook or returns data-sheet observations."
                ),
                parameters={
                    "type": "object",
                    "properties": {
                        "industry": {
                            "type": "string",
                            "enum": ["lithium", "tin", "silicon"],
                        },
                        "sector": {"type": "string"},
                        "limit": {"type": "integer", "minimum": 1, "maximum": 100},
                    },
                    "required": ["industry", "sector", "limit"],
                    "additionalProperties": False,
                },
                handler=self._indicator_catalog,
            ),
            ToolSpec(
                name="get_pipeline_run_status",
                description="Read a workflow summary from runtime/workflow-runs without changing it.",
                parameters={
                    "type": "object",
                    "properties": {
                        "run_id": {
                            "type": "string",
                            "description": "Run directory name, or 'latest'.",
                        }
                    },
                    "required": ["run_id"],
                    "additionalProperties": False,
                },
                handler=self._pipeline_status,
            ),
        ]

    def names(self) -> list[str]:
        return sorted(self._tools)

    def responses_schemas(self) -> list[dict[str, Any]]:
        return [
            {
                "type": "function",
                "name": tool.name,
                "description": tool.description,
                "parameters": tool.parameters,
                "strict": True,
            }
            for tool in self._tools.values()
        ]

    def chat_schemas(self) -> list[dict[str, Any]]:
        return [
            {
                "type": "function",
                "function": {
                    "name": tool.name,
                    "description": tool.description,
                    "parameters": tool.parameters,
                    "strict": True,
                },
            }
            for tool in self._tools.values()
        ]

    def execute(self, name: str, arguments: dict[str, Any]) -> dict[str, Any]:
        tool = self._tools.get(name)
        if tool is None:
            raise ValueError(f"Unknown or disallowed tool: {name}")
        try:
            result = tool.handler(**arguments)
            status = "SUCCESS"
            return result
        except Exception:
            status = "FAILED"
            raise
        finally:
            recorder = getattr(self.store, "record_tool_call", None)
            if callable(recorder):
                recorder(name, arguments, status)

    def _search_knowledge(self, query: str, industry: str, limit: int) -> dict[str, Any]:
        hits = self.retriever.search(query, industry=industry, limit=limit)
        return {"query": query, "results": [hit.as_dict() for hit in hits]}

    def _list_rules(self, industry: str) -> dict[str, Any]:
        documents = self.store.list_documents(
            filters={"industry": industry, "document_type": "rule"}
        )
        return {"industry": industry, "documents": documents}

    def _workbook_for_industry(self, industry: str) -> Path | None:
        aliases = {
            "lithium": ("锂", "lithium"),
            "tin": ("锡", "tin"),
            "silicon": ("硅", "silicon"),
        }[industry]
        root = self.settings.repo_root / "data" / "processed"
        candidates = sorted(root.glob("*.xlsx")) if root.exists() else []
        return next(
            (path for path in candidates if any(alias.lower() in path.name.lower() for alias in aliases)),
            None,
        )

    def _indicator_catalog(self, industry: str, sector: str, limit: int) -> dict[str, Any]:
        path = self._workbook_for_industry(industry)
        if path is None:
            return {"industry": industry, "workbook": None, "rows": []}
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            values = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(values, ())]
            rows = []
            bounded = max(1, min(int(limit), 100))
            for raw in values:
                row = {headers[index]: value for index, value in enumerate(raw) if index < len(headers) and headers[index]}
                if sector and sector not in " ".join(str(value or "") for value in row.values()):
                    continue
                rows.append({key: value for key, value in row.items() if value not in (None, "")})
                if len(rows) >= bounded:
                    break
            return {
                "industry": industry,
                "workbook": path.name,
                "catalog_sheet": sheet.title,
                "rows": rows,
            }
        finally:
            workbook.close()

    def _pipeline_status(self, run_id: str) -> dict[str, Any]:
        root = (self.settings.repo_root / "runtime" / "workflow-runs").resolve()
        if not root.exists():
            return {"run_id": run_id, "status": "NOT_FOUND"}
        if run_id == "latest":
            directories = sorted((path for path in root.iterdir() if path.is_dir()), reverse=True)
            run_dir = directories[0] if directories else None
        else:
            candidate = (root / run_id).resolve()
            candidate.relative_to(root)
            run_dir = candidate if candidate.is_dir() else None
        if run_dir is None:
            return {"run_id": run_id, "status": "NOT_FOUND"}
        summaries = sorted(run_dir.glob("*summary*.json"))
        if not summaries:
            summaries = sorted(run_dir.glob("*.json"))
        if not summaries:
            return {"run_id": run_dir.name, "status": "UNKNOWN", "path": run_dir.name}
        payload = json.loads(summaries[0].read_text(encoding="utf-8-sig"))
        return {"run_id": run_dir.name, "summary_file": summaries[0].name, "summary": payload}
