from __future__ import annotations

import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

from financial_variable_curation.rules.net_export_calculator import (
    compute_net_series,
    find_missing_net_export_pairs,
    net_title,
)


def read_directory_rows(ws) -> list[dict]:
    current_sheet = ""
    rows = []
    for row in ws.iter_rows(min_row=2):
        c0 = str(row[0].value).strip() if row[0].value is not None else ""
        if not c0:
            continue
        if "sht" in c0:
            continue
        if not c0.isdigit():
            current_sheet = c0
            continue
        c3 = str(row[3].value).strip() if row[3].value is not None else ""
        if not c3.isdigit():
            continue
        rows.append(
            {
                "sheet": current_sheet,
                "col": int(c3),
                "title": str(row[4].value) if row[4].value is not None else "",
                "unit": str(row[5].value) if row[5].value is not None else "",
                "sector": str(row[6].value) if row[6].value is not None else "",
                "major": str(row[7].value) if row[7].value is not None else "",
                "sub": str(row[8].value) if row[8].value is not None else "",
                "freq": str(row[2].value) if row[2].value is not None else "",
                "selected": str(row[11].value) if len(row) > 11 and row[11].value is not None else "",
            }
        )
    return rows


def calculate_net_exports(path: Path) -> Path:
    path = Path(path).resolve()
    tmp = path.with_name(f"{path.stem}__tmp_net.xlsx")
    shutil.copy2(path, tmp)

    wb = load_workbook(tmp, data_only=False)
    ws = wb[wb.sheetnames[0]]
    rows = read_directory_rows(ws)
    missing = find_missing_net_export_pairs(rows)
    if not missing:
        print(f"No missing net exports in {path}")
        return path

    sheet_name = "净出口计算"
    if sheet_name not in wb.sheetnames:
        ws_calc = wb.create_sheet(sheet_name)
        ws_calc.cell(1, 1, "日期")
    else:
        ws_calc = wb[sheet_name]

    existing_cols = ws_calc.max_column
    directory_start = ws.max_row + 1
    seq_max = 0
    for row in ws.iter_rows(min_row=2, max_col=1):
        if row[0].value is not None and str(row[0].value).isdigit():
            seq_max = max(seq_max, int(row[0].value))

    row_idx = directory_start
    ws.cell(row_idx, 1, sheet_name)
    ws.cell(row_idx, 2, "")
    ws.cell(row_idx, 3, "月度")
    ws.cell(row_idx, 7, missing[0]["import_row"]["sector"])
    row_idx += 1

    for pair in missing:
        imp = pair["import_row"]
        exp = pair["export_row"]
        series = compute_net_series(wb, pair["sheet"], imp["col"], exp["col"])
        if not series:
            print(f"Skip {pair['key']}: no aligned data")
            continue

        col = existing_cols + 1
        existing_cols += 1
        title = net_title(imp["title"], exp["title"], "月度")
        unit = exp.get("unit") or imp.get("unit") or "吨"

        ws_calc.cell(1, col + 1, title)
        ws_calc.cell(2, col + 1, unit)
        ws_calc.cell(3, col + 1, "月度")
        for i, point in enumerate(series):
            ws_calc.cell(4 + i, 1, point["date"])
            ws_calc.cell(4 + i, col + 1, point["value"])

        seq_max += 1
        ws.cell(row_idx, 1, seq_max)
        ws.cell(row_idx, 2, "")
        ws.cell(row_idx, 3, "月度")
        ws.cell(row_idx, 4, col)
        ws.cell(row_idx, 5, title)
        ws.cell(row_idx, 6, unit)
        ws.cell(row_idx, 7, exp["sector"])
        ws.cell(row_idx, 8, "进出口")
        ws.cell(row_idx, 9, "净出口")
        ws.cell(row_idx, 10, "流量")
        ws.cell(row_idx, 11, 0.8)
        ws.cell(row_idx, 12, "是")
        ws.cell(row_idx, 13, "确定性净出口计算")
        row_idx += 1
        print(f"Computed {title}: {len(series)} points")

    wb.save(tmp)
    wb.close()
    try:
        tmp.replace(path)
    except PermissionError:
        print(f"Permission error writing {path}; temp file kept at {tmp}")
        raise
    print(f"Net exports added to {path}: {len(missing)} pairs")
    return path


if __name__ == "__main__":
    paths = [Path(arg) for arg in sys.argv[1:] if not arg.startswith("-")]
    if not paths:
        paths = [
            Path(r"<local-documents>\Selecting skill\data\output_processed_final.xlsx"),
            Path(r"<local-documents>\market-ai-dashboard\output_processed_final_v4.xlsx"),
        ]
    for path in paths:
        calculate_net_exports(path)
