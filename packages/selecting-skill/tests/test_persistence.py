from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from financial_variable_curation.database.exceptions import PersistenceError
from financial_variable_curation.database.manager import DatabaseManager
from financial_variable_curation.database.repositories import (
    AuditEventRepository,
    PipelineRunRepository,
    SourceFileRepository,
    VariableQualityRepository,
    VariableRepository,
)
from financial_variable_curation.database.services import (
    InspectionPersistenceService,
    RunQueryService,
)
from financial_variable_curation.models.artifacts import RequestRecord
from financial_variable_curation.pipeline.inspect import run_inspection


def _persist(
    db_manager: DatabaseManager,
    input_path: Path,
    run_id: str,
    tmp_path: Path,
    request_record: RequestRecord | None = None,
):
    result = run_inspection(input_path, run_id=run_id)
    request = request_record or RequestRecord(
        run_id=run_id,
        input_path=str(input_path),
        file_hash=result.workbook_profile.file_hash,
        cli_args={"input": str(input_path)},
    )
    return result, request


def test_persist_inspection_result_and_query_run(db_manager: DatabaseManager, sample_workbook: Path, tmp_path: Path) -> None:
    result, request = _persist(db_manager, sample_workbook, "run-persist-1", tmp_path)
    summary = InspectionPersistenceService(db_manager).persist(
        result,
        request_record=request,
        artifacts_dir=tmp_path / "artifacts",
        run_id="run-persist-1",
    )
    assert summary.status == "COMPLETED"
    assert summary.file_count == 1
    assert summary.sheet_count == 1
    assert summary.variable_count > 1
    assert summary.quality_count == summary.variable_count
    assert summary.audit_event_count >= 4

    query = RunQueryService(db_manager).get_run_summary("run-persist-1")
    assert query.run.status == "COMPLETED"
    assert query.file_count == 1
    assert query.sheet_count == 1
    assert query.variable_count == summary.variable_count
    assert query.quality_count == summary.variable_count


def test_repeat_persistence_is_idempotent(
    db_manager: DatabaseManager,
    sample_workbook: Path,
    tmp_path: Path,
) -> None:
    result, request = _persist(db_manager, sample_workbook, "run-repeat", tmp_path)
    service = InspectionPersistenceService(db_manager)
    first = service.persist(result, request_record=request, artifacts_dir=tmp_path / "a", run_id="run-repeat")
    second = service.persist(result, request_record=request, artifacts_dir=tmp_path / "b", run_id="run-repeat")
    assert first.model_dump() == second.model_dump()
    with db_manager.session_scope() as session:
        assert AuditEventRepository(session).count_for_run("run-repeat") == first.audit_event_count
        assert VariableRepository(session).count_for_run("run-repeat") == first.variable_count


def test_same_content_different_paths_reuses_source_file(
    db_manager: DatabaseManager,
    sample_workbook: Path,
    tmp_path: Path,
) -> None:
    copy = tmp_path / "copy.xlsx"
    shutil.copy2(sample_workbook, copy)
    result_a, request_a = _persist(db_manager, sample_workbook, "run-copy-a", tmp_path)
    result_b, request_b = _persist(db_manager, copy, "run-copy-b", tmp_path)
    service = InspectionPersistenceService(db_manager)
    service.persist(result_a, request_record=request_a, artifacts_dir=tmp_path / "a", run_id="run-copy-a")
    service.persist(result_b, request_record=request_b, artifacts_dir=tmp_path / "b", run_id="run-copy-b")
    with db_manager.session_scope() as session:
        file_repo = SourceFileRepository(session)
        assert file_repo.count() == 1
        assert file_repo.count_run_files("run-copy-a") == 1
        assert file_repo.count_run_files("run-copy-b") == 1
        files_a = file_repo.list_for_run("run-copy-a")
        files_b = file_repo.list_for_run("run-copy-b")
        assert files_a[0].file_id == files_b[0].file_id


def test_same_path_with_changed_content_creates_new_source_file(
    db_manager: DatabaseManager,
    sample_workbook: Path,
    tmp_path: Path,
) -> None:
    result_a, request_a = _persist(db_manager, sample_workbook, "run-change-a", tmp_path)
    service = InspectionPersistenceService(db_manager)
    service.persist(result_a, request_record=request_a, artifacts_dir=tmp_path / "a", run_id="run-change-a")

    workbook = load_workbook(sample_workbook)
    sheet = workbook.active
    sheet.cell(row=sheet.max_row, column=2, value=999999)
    workbook.save(sample_workbook)
    workbook.close()

    result_b, request_b = _persist(db_manager, sample_workbook, "run-change-b", tmp_path)
    service.persist(result_b, request_record=request_b, artifacts_dir=tmp_path / "b", run_id="run-change-b")
    with db_manager.session_scope() as session:
        assert SourceFileRepository(session).count() == 2


def test_failed_persistence_rolls_back_and_marks_run_failed(
    db_manager: DatabaseManager,
    sample_workbook: Path,
    tmp_path: Path,
    monkeypatch,
) -> None:
    from financial_variable_curation.database.repositories import VariableRepository

    result, request = _persist(db_manager, sample_workbook, "run-fail", tmp_path)

    def fail_variables(self, variables):
        raise RuntimeError("injected variable persistence failure")

    monkeypatch.setattr(VariableRepository, "bulk_upsert", fail_variables)
    with pytest.raises(PersistenceError, match="run-fail"):
        InspectionPersistenceService(db_manager).persist(
            result,
            request_record=request,
            artifacts_dir=tmp_path / "artifacts",
            run_id="run-fail",
        )

    query = RunQueryService(db_manager).get_run_summary("run-fail")
    assert query.run.status == "FAILED"
    assert query.file_count == 0
    assert query.sheet_count == 0
    assert query.variable_count == 0
    assert query.quality_count == 0
    assert query.audit_event_count >= 2


def test_variable_ids_are_not_regenerated(
    db_manager: DatabaseManager,
    sample_workbook: Path,
    tmp_path: Path,
) -> None:
    result, request = _persist(db_manager, sample_workbook, "run-ids", tmp_path)
    InspectionPersistenceService(db_manager).persist(
        result,
        request_record=request,
        artifacts_dir=tmp_path / "artifacts",
        run_id="run-ids",
    )
    with db_manager.session_scope() as session:
        stored = VariableRepository(session).list_by_run_id("run-ids")
        assert [item.variable_id for item in stored] == [item.variable_id for item in result.variable_profiles]


def test_json_fields_roundtrip_as_structured_data(
    db_manager: DatabaseManager,
    sample_workbook: Path,
    tmp_path: Path,
) -> None:
    result, request = _persist(db_manager, sample_workbook, "run-json", tmp_path)
    InspectionPersistenceService(db_manager).persist(
        result,
        request_record=request,
        artifacts_dir=tmp_path / "artifacts",
        run_id="run-json",
    )
    with db_manager.session_scope() as session:
        run = PipelineRunRepository(session).get("run-json")
        assert run is not None
        assert isinstance(run.input_request_json, dict)
        assert run.input_request_json["cli_args"]["input"] == str(sample_workbook)
        variables = VariableRepository(session).list_by_run_id("run-json")
        assert all(isinstance(item.review_reasons_json, list) for item in variables)


def test_needs_review_and_frequency_stats_are_queryable(
    db_manager: DatabaseManager,
    sample_workbook: Path,
    tmp_path: Path,
) -> None:
    result, request = _persist(db_manager, sample_workbook, "run-stats", tmp_path)
    InspectionPersistenceService(db_manager).persist(
        result,
        request_record=request,
        artifacts_dir=tmp_path / "artifacts",
        run_id="run-stats",
    )
    with db_manager.session_scope() as session:
        review = VariableQualityRepository(session).list_needs_review(run_id="run-stats")
        assert any(item.variable_id in {profile.variable_id for profile in result.needs_review_variables} for item in review)
        stats = VariableQualityRepository(session).count_by_frequency_and_status()
        assert any(item.detected_frequency == "daily" and item.count > 0 for item in stats)
