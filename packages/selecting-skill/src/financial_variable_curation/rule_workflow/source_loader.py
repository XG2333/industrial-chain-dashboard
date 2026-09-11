from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


class RuleSourceDocument:
    def __init__(
        self,
        text: str,
        structure: list[dict[str, Any]],
        metadata: dict[str, Any],
    ) -> None:
        self.text = text
        self.structure = structure
        self.metadata = metadata


def load_rule_source(path: str | Path) -> RuleSourceDocument:
    source_path = Path(path)
    if not source_path.exists():
        raise FileNotFoundError(f"Rule file does not exist: {source_path}")
    if source_path.suffix.lower() == ".docx":
        return _load_docx(source_path)
    if source_path.suffix.lower() in {".txt", ".md"}:
        text = source_path.read_text(encoding="utf-8")
        structure = [{"order": 0, "type": "paragraph", "text": line, "heading_level": None} for line in text.splitlines()]
        metadata = {
            "source_file_name": source_path.name,
            "source_file_hash": _sha256(source_path),
            "source_type": "TEXT" if source_path.suffix.lower() == ".txt" else "MARKDOWN",
            "paragraph_count": len(structure),
            "table_count": 0,
            "extracted_text_length": len(text),
            "warnings": [],
            "unsupported_content": [],
        }
        return RuleSourceDocument(text=text, structure=structure, metadata=metadata)
    raise ValueError(f"Unsupported rule source extension: {source_path.suffix}")


def _load_docx(path: Path) -> RuleSourceDocument:
    document = Document(path)
    structure: list[dict[str, Any]] = []
    texts: list[str] = []
    order = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            style_name = paragraph.style.name if paragraph.style else ""
            heading_level = None
            if style_name.lower().startswith("heading"):
                heading_level = int(style_name.split()[-1]) if style_name.split()[-1].isdigit() else 1
            if text:
                structure.append(
                    {
                        "order": order,
                        "type": "paragraph",
                        "text": text,
                        "heading_level": heading_level,
                        "style": style_name,
                    }
                )
                texts.append(("H" + str(heading_level) + " " if heading_level else "") + text)
                order += 1
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            rows = []
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                rows.append(row_text)
                texts.append(" | ".join(row_text))
            structure.append(
                {
                    "order": order,
                    "type": "table",
                    "text": "\n".join(" | ".join(row) for row in rows),
                    "rows": rows,
                }
            )
            order += 1
    text = "\n".join(item for item in texts if item.strip())
    unsupported = []
    if len(document.inline_shapes):
        unsupported.append(f"{len(document.inline_shapes)} inline shape(s)")
    if len(document.part.package.parts) > 0:
        pass
    metadata = {
        "source_file_name": path.name,
        "source_file_hash": _sha256(path),
        "source_type": "DOCX",
        "paragraph_count": sum(1 for item in structure if item["type"] == "paragraph"),
        "table_count": sum(1 for item in structure if item["type"] == "table"),
        "extracted_text_length": len(text),
        "warnings": [],
        "unsupported_content": unsupported,
    }
    return RuleSourceDocument(text=text, structure=structure, metadata=metadata)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
