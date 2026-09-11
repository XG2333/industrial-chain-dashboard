from __future__ import annotations

import hashlib
import json
from collections import Counter
from pathlib import Path
from uuid import uuid4

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from financial_variable_curation.database.manager import DatabaseManager
from financial_variable_curation.database.repositories import (
    AuditEventRepository,
    ExportedVariableRepository,
    ExportRunRepository,
    ReviewItemRepository,
    RuleSetRepository,
    SelectionResultRepository,
    SelectionRunRepository,
    SourceFileRepository,
    SourceSheetRepository,
    VariableClassificationRepository,
    VariableQualityRepository,
    VariableRepository,
)
from financial_variable_curation.database.schemas import (
    AuditEventDTO,
    ExportedVariableDTO,
    ExportRunDTO,
)
from financial_variable_curation.database.types import utcnow
from financial_variable_curation.export.models import ExportColumnMapping, ExportSummary


class SelectionExportError(Exception):
    pass


class SelectionExportService:
    def __init__(self, manager: DatabaseManager | None = None) -> None:
        self.manager = manager or DatabaseManager()

    def export(
        self,
        selection_run_id: str,
        *,
        output_path: str | Path | None = None,
        overwrite: bool = False,
        include_not_selected: bool = True,
        include_rule_summary: bool = False,
        dry_run: bool = False,
        artifacts_dir: str | Path = "artifacts",
    ) -> ExportSummary:
        with self.manager.session_scope() as session:
            selection_run = SelectionRunRepository(session).get(selection_run_id)
            if selection_run is None:
                raise SelectionExportError(f"Selection run not found: {selection_run_id}")
            if selection_run.status not in {"COMPLETED", "REVIEW_REQUIRED"}:
                raise SelectionExportError(
                    f"Selection run status {selection_run.status} is not exportable."
                )
            results = SelectionResultRepository(session).list_for_run(selection_run_id)
            rule_set = (
                RuleSetRepository(session).get(selection_run.rule_set_id)
                if selection_run.rule_set_id
                else None
            )
            source_run = selection_run.classification_run_id or selection_run.pipeline_run_id
            variables = {
                item.variable_id: item
                for item in VariableRepository(session).list_by_run_id(source_run)
            }
            classifications = {
                item.variable_id: item
                for item in VariableClassificationRepository(session).list_for_run(source_run)
            }
            qualities = {
                item.variable_id: item
                for item in VariableQualityRepository(session).list_by_run_id(source_run)
            }
            source_files = SourceFileRepository(session).list_for_run(source_run)
            if not source_files:
                raise SelectionExportError("No source file record found for selection run.")
            source_file = source_files[0]
            sheets = SourceSheetRepository(session).list_by_file_id(source_file.file_id)
            review_items = {
                item.entity_id: item
                for item in ReviewItemRepository(session).list_for_run(source_run)
                if item.status == "NEEDS_REVIEW"
            }
        if not results:
            raise SelectionExportError("Selection run has no variable results.")
        source_path = Path(source_file.original_path or "")
        if not source_path.exists():
            raise SelectionExportError(f"Original source file missing: {source_path}")
        if self._sha256(source_path) != source_file.file_hash:
            raise SelectionExportError("SOURCE_FILE_CHANGED")
        workbook = load_workbook(source_path, data_only=True)
        if set(sheet.sheet_name for sheet in sheets).issubset(set(workbook.sheetnames)) is False:
            raise SelectionExportError("Source workbook sheets do not match persisted sheet records.")

        selected = [item for item in results if item.final_status == "SELECTED"]
        rejected = [
            item
            for item in results
            if item.final_status in {"REJECTED", "DUPLICATE", "NOT_SELECTED"}
        ]
        needs_review = [
            item
            for item in results
            if item.final_status in {"NEEDS_REVIEW", "FAILED"}
        ]
        selected_sorted = sorted(selected, key=lambda item: item.overall_rank or 9999)
        export_run_id = f"exp_{uuid4().hex[:12]}"
        if dry_run:
            return ExportSummary(
                export_run_id=export_run_id,
                selection_run_id=selection_run_id,
                pipeline_run_id=selection_run.pipeline_run_id,
                status="DRY_RUN",
                selected_variable_count=len(selected),
                started_at=utcnow(),
            )
        resolved_output = self._resolve_output(
            source_path=source_path,
            output_path=output_path,
            rule_set=rule_set,
            overwrite=overwrite,
        )
        output_path_resolved = self._write_workbook(
            source_path=source_path,
            workbook=workbook,
            source_file=source_file,
            sheets=sheets,
            variables=variables,
            classifications=classifications,
            qualities=qualities,
            selected=selected_sorted,
            rejected=rejected,
            needs_review=needs_review,
            review_items=review_items,
            selection_run=selection_run,
            rule_set=rule_set,
            include_rule_summary=include_rule_summary,
            output_path=resolved_output,
        )
        output_hash = self._sha256(output_path_resolved)
        sheet_names = list(load_workbook(output_path_resolved, read_only=True).sheetnames)
        mappings = self._column_mappings(
            selected=selected_sorted,
            variables=variables,
            sheets=sheets,
        )
        export_dir = Path(artifacts_dir) / export_run_id / "export"
        export_dir.mkdir(parents=True, exist_ok=True)
        (export_dir / "export_request.json").write_text(
            json.dumps(
                {
                    "selection_run_id": selection_run_id,
                    "output_path": str(output_path_resolved),
                    "overwrite": overwrite,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        (export_dir / "export_column_mapping.json").write_text(
            json.dumps([mapping.model_dump(mode="json") for mapping in mappings], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        summary = ExportSummary(
            export_run_id=export_run_id,
            selection_run_id=selection_run_id,
            pipeline_run_id=selection_run.pipeline_run_id,
            status="COMPLETED",
            output_path=str(output_path_resolved),
            output_file_hash=output_hash,
            output_file_size_bytes=output_path_resolved.stat().st_size,
            selected_variable_count=len(selected),
            sheet_count=len(sheet_names),
            sheet_names=sheet_names,
            completed_at=utcnow(),
            artifacts_path=str(export_dir),
        )
        (export_dir / "export_summary.json").write_text(
            json.dumps(summary.model_dump(mode="json"), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (export_dir / "output_manifest.json").write_text(
            json.dumps(
                {
                    "output_file": str(output_path_resolved),
                    "file_hash": output_hash,
                    "file_size_bytes": output_path_resolved.stat().st_size,
                    "sheet_names": sheet_names,
                    "selected_variable_count": len(selected),
                    "source_selection_run_id": selection_run_id,
                    "created_at": utcnow().isoformat(),
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        self._persist(
            export_run=summary,
            output_path=output_path_resolved,
            mappings=mappings,
        )
        return summary

    def _write_workbook(
        self,
        *,
        source_path: Path,
        workbook,
        source_file,
        sheets,
        variables,
        classifications,
        qualities,
        selected,
        rejected,
        needs_review,
        review_items,
        selection_run,
        rule_set,
        include_rule_summary: bool,
        output_path: Path,
    ) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        sheet_by_name = {sheet.sheet_name: sheet for sheet in sheets}
        variable_by_id = variables
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            for sheet_profile in sheets:
                raw = pd.read_excel(
                    source_path,
                    sheet_name=sheet_profile.sheet_name,
                    header=None,
                )
                metadata_row_count = max(
                    1,
                    getattr(sheet_profile, "metadata_row_count", 1) or 1,
                )
                header_index = max(0, (sheet_profile.detected_header_row or 1) - 1)
                header = (
                    raw.iloc[header_index]
                    if header_index < len(raw)
                    else pd.Series(dtype=object)
                )
                data_frame = raw.iloc[metadata_row_count:].copy()
                data_frame.columns = [
                    str(value) if value is not None and str(value).strip() else f"Column{index + 1}"
                    for index, value in enumerate(header)
                ]
                selected_sheet = [
                    item
                    for item in selected
                    if variable_by_id[item.variable_id].sheet_id == sheet_profile.sheet_id
                ]
                if sheet_profile.selected_date_column and sheet_profile.selected_date_column in data_frame.columns:
                    columns = [sheet_profile.selected_date_column]
                else:
                    columns = []
                for item in selected_sheet:
                    variable = variable_by_id[item.variable_id]
                    if variable.column_index < len(data_frame.columns):
                        columns.append(data_frame.columns[variable.column_index])
                columns = list(dict.fromkeys(columns))
                if columns:
                    data_frame.loc[:, columns].to_excel(
                        writer,
                        sheet_name=self._safe_sheet_name(f"selected_data__{sheet_profile.sheet_name}"),
                        index=False,
                    )
            selected_rows = [
                self._selected_row(item, variable_by_id, classifications, qualities)
                for item in selected
            ]
            pd.DataFrame(selected_rows).to_excel(writer, sheet_name="selected_variables", index=False)
            rejected_rows = [
                self._rejected_row(item, variable_by_id, classifications, qualities)
                for item in rejected
            ]
            pd.DataFrame(rejected_rows).to_excel(writer, sheet_name="rejected_variables", index=False)
            review_rows = [
                self._review_row(item, variable_by_id, classifications, review_items)
                for item in needs_review
            ]
            pd.DataFrame(review_rows).to_excel(writer, sheet_name="needs_review", index=False)
            pd.DataFrame([self._run_summary_row(selection_run, rule_set, source_file, output_path, len(selected))]).to_excel(
                writer, sheet_name="run_summary", index=False
            )
            if include_rule_summary:
                pd.DataFrame([{"rule_type": "N/A", "note": "Rule summary requires compiled plan"}]).to_excel(
                    writer, sheet_name="rule_summary", index=False
                )
        styled = load_workbook(output_path)
        for sheet in styled.worksheets:
            if sheet.max_row > 0:
                for cell in sheet[1]:
                    cell.font = Font(bold=True)
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for column_cells in sheet.columns:
                    lengths = [
                        len(str(cell.value or ""))
                        for cell in column_cells
                        if cell.value is not None
                    ]
                    if not lengths:
                        continue
                    max_length = max(lengths)
                    sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
                        max(max_length + 2, 10), 50
                    )
        styled.save(output_path)
        return output_path.resolve()

    @staticmethod
    def _safe_sheet_name(name: str) -> str:
        cleaned = "".join(char for char in name if char not in r"[]:*?/\\")
        return (cleaned or "Sheet")[:31]

    def _resolve_output(
        self,
        *,
        source_path: Path,
        output_path: str | Path | None,
        rule_set,
        overwrite: bool,
    ) -> Path:
        if output_path:
            resolved = Path(output_path).resolve()
        else:
            rule_name = rule_set.rule_set_name if rule_set else "rules"
            version = rule_set.version if rule_set else "1"
            resolved = (
                source_path.parent
                / f"{source_path.stem}__selected__{rule_name}__v{version}.xlsx"
            )
        if resolved == source_path.resolve():
            raise SelectionExportError("Output path must not overwrite the original source file.")
        if resolved.exists() and not overwrite:
            resolved = resolved.with_name(f"{resolved.stem}__{uuid4().hex[:8]}{resolved.suffix}")
        return resolved

    def _column_mappings(self, *, selected, variables, sheets) -> list[ExportColumnMapping]:
        sheet_by_id = {sheet.sheet_id: sheet for sheet in sheets}
        mappings = []
        for index, item in enumerate(selected):
            variable = variables[item.variable_id]
            sheet = sheet_by_id[variable.sheet_id]
            mappings.append(
                ExportColumnMapping(
                    variable_id=item.variable_id,
                    original_sheet_name=sheet.sheet_name,
                    original_column_index=variable.column_index,
                    original_column_name=variable.original_name,
                    output_sheet_name=self._safe_sheet_name(f"selected_data__{sheet.sheet_name}"),
                    output_column_index=index + 1,
                    output_column_name=variable.original_name,
                )
            )
        return mappings

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def _selected_row(item, variables, classifications, qualities) -> dict:
        variable = variables[item.variable_id]
        classification = classifications.get(item.variable_id)
        quality = qualities.get(item.variable_id)
        return {
            "overall_rank": item.overall_rank,
            "variable_id": item.variable_id,
            "file_name": variable.original_name,
            "sheet_name": variable.sheet_id,
            "original_name": variable.original_name,
            "standard_name": classification.standard_name if classification else variable.original_name,
            "category_level_1": classification.category_level_1 if classification else None,
            "category_level_2": classification.category_level_2 if classification else None,
            "commodity": classification.commodity if classification else None,
            "market": classification.market if classification else None,
            "region": classification.region if classification else None,
            "detected_frequency": quality.detected_frequency if quality else None,
            "unit_standard": classification.unit if classification else None,
            "classification_confidence": classification.confidence if classification else None,
            "missing_rate": quality.missing_rate if quality else None,
            "coverage_days": quality.coverage_days if quality else None,
            "total_score": item.total_score,
            "category_rank": item.category_rank,
            "frequency_rank": item.frequency_rank,
            "group_rank": item.group_rank,
            "comparison_group_key": item.comparison_group_key,
            "primary_reason_code": item.primary_reason_code,
            "primary_reason_text": item.primary_reason_text,
            "matched_rule_ids": item.matched_rule_ids_json,
            "selection_status": item.final_status,
        }

    @staticmethod
    def _rejected_row(item, variables, classifications, qualities) -> dict:
        variable = variables[item.variable_id]
        classification = classifications.get(item.variable_id)
        quality = qualities.get(item.variable_id)
        return {
            "variable_id": item.variable_id,
            "original_name": variable.original_name,
            "standard_name": classification.standard_name if classification else variable.original_name,
            "final_status": item.final_status,
            "primary_reason_code": item.primary_reason_code,
            "primary_reason_text": item.primary_reason_text,
            "matched_rule_ids": item.matched_rule_ids_json,
            "missing_rate": quality.missing_rate if quality else None,
            "classification_confidence": classification.confidence if classification else None,
            "replacement_variable_id": item.replacement_variable_id,
            "comparison_group_key": item.comparison_group_key,
        }

    @staticmethod
    def _review_row(item, variables, classifications, review_items) -> dict:
        variable = variables[item.variable_id]
        classification = classifications.get(item.variable_id)
        review = review_items.get(item.variable_id)
        reason_code = item.primary_reason_code or (review.reason_code if review else None)
        reason_text = item.primary_reason_text or (review.reason_text if review else None)
        return {
            "variable_id": item.variable_id,
            "original_name": variable.original_name,
            "standard_name": classification.standard_name if classification else variable.original_name,
            "review_type": review.review_type if review else "CLASSIFICATION_REVIEW",
            "reason_code": reason_code,
            "reason_text": reason_text,
            "classification_confidence": classification.confidence if classification else None,
            "category_level_1": classification.category_level_1 if classification else None,
            "comparison_group_key": item.comparison_group_key,
            "blocking": True,
            "current_status": item.final_status,
            "suggested_action": "Review before production use.",
        }

    @staticmethod
    def _run_summary_row(selection_run, rule_set, source_file, output_path, selected_count) -> dict:
        return {
            "pipeline_run_id": selection_run.pipeline_run_id,
            "selection_run_id": selection_run.selection_run_id,
            "rule_set_id": selection_run.rule_set_id,
            "rule_set_name": rule_set.rule_set_name if rule_set else None,
            "rule_set_version": rule_set.version if rule_set else None,
            "source_file_name": source_file.file_name,
            "source_file_hash": source_file.file_hash,
            "selected_count": selected_count,
            "output_file": str(output_path),
            "production_ready": False,
        }

    def _persist(self, *, export_run: ExportSummary, output_path: Path, mappings: list[ExportColumnMapping]) -> None:
        with self.manager.unit_of_work() as uow:
            ExportRunRepository(uow.session).save(
                ExportRunDTO(
                    export_run_id=export_run.export_run_id,
                    pipeline_run_id=export_run.pipeline_run_id,
                    selection_run_id=export_run.selection_run_id,
                    status="COMPLETED",
                    output_path=str(output_path),
                    output_file_hash=export_run.output_file_hash,
                    output_file_size_bytes=export_run.output_file_size_bytes,
                    selected_variable_count=export_run.selected_variable_count,
                    sheet_count=export_run.sheet_count,
                    started_at=export_run.started_at,
                    completed_at=export_run.completed_at,
                    created_at=utcnow(),
                    updated_at=utcnow(),
                )
            )
            ExportedVariableRepository(uow.session).bulk_save(
                export_run.export_run_id,
                [
                    ExportedVariableDTO(
                        export_run_id=export_run.export_run_id,
                        variable_id=mapping.variable_id,
                        output_sheet_name=mapping.output_sheet_name,
                        output_column_index=mapping.output_column_index,
                        output_column_name=mapping.output_column_name,
                        created_at=utcnow(),
                    )
                    for mapping in mappings
                ],
            )
            AuditEventRepository(uow.session).add(
                AuditEventDTO(
                    event_id=str(uuid4()),
                    run_id=export_run.pipeline_run_id,
                    event_type="EXPORT_COMPLETED",
                    entity_type="export_run",
                    entity_id=export_run.export_run_id,
                    event_payload_json={"output_path": str(output_path)},
                    created_at=utcnow(),
                )
            )
