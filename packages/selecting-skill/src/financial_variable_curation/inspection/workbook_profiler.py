from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from financial_variable_curation.inspection.exceptions import InspectionError


@dataclass
class RawSheet:
    sheet_name: str
    rows: list[list[object]]
    row_count: int
    column_count: int
    empty_sheet: bool


@dataclass
class RawWorkbook:
    file_name: str
    file_path: str
    file_hash: str | None
    file_size_bytes: int
    sheets: list[RawSheet] = field(default_factory=list)


@dataclass
class ColumnData:
    index: int
    header: str
    values: list[object]


class WorkbookProfiler:
    def profile(self, path: str | Path, file_hash: str | None) -> RawWorkbook:
        input_path = Path(path)
        try:
            workbook = load_workbook(input_path, data_only=True, read_only=True)
        except Exception as exc:
            raise InspectionError(f"Cannot open workbook as xlsx: {exc}") from exc

        sheets: list[RawSheet] = []
        try:
            for worksheet in workbook.worksheets:
                rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
                max_len = max((len(row) for row in rows), default=0)
                sheets.append(
                    RawSheet(
                        sheet_name=worksheet.title,
                        rows=rows,
                        row_count=len(rows),
                        column_count=max_len,
                        empty_sheet=len(rows) == 0,
                    )
                )
        finally:
            workbook.close()

        return RawWorkbook(
            file_name=input_path.name,
            file_path=str(input_path.resolve()),
            file_hash=file_hash,
            file_size_bytes=input_path.stat().st_size,
            sheets=sheets,
        )
