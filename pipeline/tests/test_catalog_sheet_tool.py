# -*- coding: utf-8 -*-
"""catalog_sheet_tool.py（目录表精简/合并）测试。"""

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink
from pathlib import Path

from catalog_sheet_tool import extract, merge


def _build_full(path: Path) -> None:
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "指标目录"
    ws0.append(["#", "Sheet Name", "Freq", "Col", "Indicator Name"])
    for i in range(1, 6):
        ws0.append([i, f"s{i}", "日度", i, f"指标{i}"])
    ws0["A1"].font = Font(bold=True, size=11)
    ws0["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws0["E1"].hyperlink = Hyperlink(ref="E1", location="s1", display="跳转")
    ws0.column_dimensions["A"].width = 12
    ws0.row_dimensions[1].height = 20
    ws0.merge_cells("C2:D2")
    ws0.freeze_panes = "A2"

    ws1 = wb.create_sheet("数据1")
    ws1.append(["日期", "值"])
    ws1.append(["2026-01-01", 1.5])
    ws2 = wb.create_sheet("净出口计算")
    ws2.append(["日期", "净出口"])
    ws2.append(["2026-01-01", 100])
    wb.save(path)
    wb.close()


def _read_values(path: Path, sheet: str = "指标目录") -> list[list]:
    wb = load_workbook(path, data_only=False)
    ws = wb[sheet]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    wb.close()
    return rows


def test_extract_keeps_only_first_sheet(tmp_path):
    full = tmp_path / "xxx_workflow_ai_step1.xlsx"
    slim = tmp_path / "xxx_workflow_ai_slim.xlsx"
    _build_full(full)
    stats = extract(full, slim)
    assert stats["kept_sheet"] == "指标目录"
    assert slim.exists()
    assert (tmp_path / "xxx_workflow_ai__full_template.xlsx").exists()
    wb = load_workbook(slim)
    assert wb.sheetnames == ["指标目录"]
    assert wb["指标目录"].max_row == 6
    wb.close()


def test_extract_preserves_ws0_styles(tmp_path):
    full = tmp_path / "xxx_workflow_ai_step1.xlsx"
    slim = tmp_path / "xxx_workflow_ai_slim.xlsx"
    _build_full(full)
    extract(full, slim)
    wb = load_workbook(slim)
    ws = wb["指标目录"]
    assert ws["A1"].font.bold is True
    assert ws["A1"].fill.fgColor.rgb.endswith("FFFF00")
    assert ws["E1"].hyperlink is not None
    assert ws.column_dimensions["A"].width == 12
    assert ws.row_dimensions[1].height == 20
    assert "C2:D2" in [str(r) for r in ws.merged_cells.ranges]
    assert ws.freeze_panes == "A2"
    wb.close()


def test_merge_roundtrip_preserves_full_workbook(tmp_path):
    full = tmp_path / "xxx_workflow_ai_step1.xlsx"
    slim = tmp_path / "xxx_workflow_ai_slim.xlsx"
    out = tmp_path / "merged.xlsx"
    _build_full(full)
    extract(full, slim)
    # 不改 slim，直接 merge：输出应等于原完整文件（ws0 值/样式 + 全部 sheets）
    stats = merge(slim, out)
    assert stats["added_sheets"] == []
    assert out.exists()
    wb = load_workbook(out)
    assert wb.sheetnames == ["指标目录", "数据1", "净出口计算"]
    ws = wb["指标目录"]
    assert ws.max_row == 6
    assert ws["A1"].font.bold is True
    assert ws["A1"].fill.fgColor.rgb.endswith("FFFF00")
    assert ws["E1"].hyperlink is not None
    assert ws.freeze_panes == "A2"
    assert wb["数据1"]["B2"].value == 1.5
    assert wb["净出口计算"]["B2"].value == 100
    # ws0 值逐格一致
    assert _read_values(out) == _read_values(full)
    wb.close()


def test_merge_removes_template_merges_removed_in_slim(tmp_path):
    # slim 中删除的合并单元格，merge 后不得从模板残留
    full = tmp_path / "xxx_workflow_ai_step1.xlsx"
    slim = tmp_path / "xxx_workflow_ai_slim.xlsx"
    out = tmp_path / "merged.xlsx"
    _build_full(full)
    extract(full, slim)
    wb = load_workbook(slim)
    ws = wb["指标目录"]
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    wb.save(slim)
    wb.close()
    stats = merge(slim, out)
    wb = load_workbook(out)
    merged = [str(m) for m in wb["指标目录"].merged_cells.ranges]
    assert merged == [], merged
    assert wb["指标目录"]["C2"].value == "日度"  # 合并删除后恢复普通 cell 值
    wb.close()


def test_merge_no_template_hyperlink_residue(tmp_path):
    # slim 中已删除的超链接，merge 后不得从模板残留
    full = tmp_path / "xxx_workflow_ai_step1.xlsx"
    slim = tmp_path / "xxx_workflow_ai_slim.xlsx"
    out = tmp_path / "merged.xlsx"
    _build_full(full)
    extract(full, slim)
    wb = load_workbook(slim)
    ws = wb["指标目录"]
    ws["E1"].hyperlink = None  # 删除 slim 中唯一超链接
    wb.save(slim)
    wb.close()
    merge(slim, out)
    wb = load_workbook(out)
    hls = [c.hyperlink for row in wb["指标目录"].iter_rows() for c in row if c.hyperlink]
    assert hls == [], hls
    wb.close()


def test_merge_adds_new_sheets_and_updated_ws0(tmp_path):
    full = tmp_path / "xxx_workflow_ai_step1.xlsx"
    slim = tmp_path / "xxx_workflow_ai_slim.xlsx"
    out = tmp_path / "merged.xlsx"
    _build_full(full)
    extract(full, slim)

    # 修改 slim 的 ws0（追加行 + 改值）并新增一个 sheet
    wb = load_workbook(slim)
    ws = wb["指标目录"]
    ws.cell(7, 1, 6)
    ws.cell(7, 5, "新增指标")
    ws["E1"] = "跳转2"
    ws2 = wb.create_sheet("成交持仓比计算")
    ws2.append(["日期", "比率"])
    ws2.append(["2026-01-01", 0.95])
    wb.save(slim)
    wb.close()

    stats = merge(slim, out)
    assert "成交持仓比计算" in stats["added_sheets"]
    wb = load_workbook(out)
    assert wb.sheetnames == ["指标目录", "数据1", "净出口计算", "成交持仓比计算"]
    ws = wb["指标目录"]
    assert ws.max_row == 7
    assert ws.cell(7, 5).value == "新增指标"
    assert ws["E1"].value == "跳转2"
    assert wb["成交持仓比计算"]["B2"].value == 0.95
    wb.close()


def test_merge_postsel_suffix_matches_template(tmp_path):
    # 生产 --ai 路径：slim 最终名是 xxx_workflow_ai_postsel.xlsx，
    # _template_name 必须能剥离 _postsel 找到与 step1 同源的模板
    full = tmp_path / "xxx_workflow_ai_step1.xlsx"
    slim = tmp_path / "xxx_workflow_ai_slim.xlsx"
    postsel = tmp_path / "xxx_workflow_ai_postsel.xlsx"
    out = tmp_path / "merged.xlsx"
    _build_full(full)
    extract(full, slim)
    import shutil as _shutil

    _shutil.copy2(slim, postsel)
    stats = merge(postsel, out)
    assert stats["added_sheets"] == []
    assert out.exists()
    wb = load_workbook(out)
    assert wb.sheetnames == ["指标目录", "数据1", "净出口计算"]
    assert wb["指标目录"].max_row == 6
    wb.close()
