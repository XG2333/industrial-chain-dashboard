from __future__ import annotations

from datetime import date
import importlib.util
from pathlib import Path
import sys

import pytest
import yaml
from openpyxl import Workbook, load_workbook

from workflow.cli import run_workflow
from workflow.config import StepConfig, WorkflowConfig
from workflow.runners import (
    WorkflowContext,
    _dashboard_data_sorter,
    _stock_target_curator,
)


PROJECT_ROOT = Path(__file__).resolve().parent.parent
REPO_ROOT = PROJECT_ROOT.parent
SKILL2_ROOT = REPO_ROOT / "packages" / "selecting-skill"
SKILL1_SCRIPT = PROJECT_ROOT / "scripts" / "tin_catalog.py"


def _make_small_workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "目录"
    ws = workbook.create_sheet("测试进出口-月")
    ws.append(
        [
            "日期",
            "中国海关: 测试进口量: 总计: 月度",
            "中国海关: 测试出口量: 总计: 月度",
        ]
    )
    ws.append(["", "吨", "吨"])
    ws.append(["", "月度", "月度"])
    for index in range(6):
        ws.append([date(2024, 1 + index, 1), 100 + index, 110 + index])
    workbook.save(path)


def _write_smoke_config(tmp_path: Path) -> Path:
    config = {
        "workflow": {
            "name": "smoke_workflow",
            "project_root": str(SKILL2_ROOT).replace("\\", "/"),
            "python": str(Path(sys.executable)).replace("\\", "/"),
            "provider": "mock",
            "rules": str(SKILL2_ROOT / "input" / "selection_rules_minimal.txt").replace("\\", "/"),
            "run_root": str(tmp_path / "runs").replace("\\", "/"),
            "final_output": str(tmp_path / "final.xlsx").replace("\\", "/"),
            "steps": [
                {
                    "id": "excel_catalog_curator",
                    "name": "Catalog",
                    "runner": "excel_catalog_curator",
                    "params": {
                        "script": str(SKILL1_SCRIPT).replace("\\", "/"),
                        "output_name": "cataloged.xlsx",
                        "pickle_name": "dataset.pkl",
                    },
                },
                {
                    "id": "financial_variable_curation",
                    "name": "Curation",
                    "runner": "financial_variable_directory_mark",
                    "params": {
                        "output_name": "directory_marked.xlsx",
                        "database_name": "curation.db",
                        "artifacts_dir": "skill2_artifacts",
                        "timeout_ms": 600000,
                    },
                },
                {
                    "id": "calculate_net_exports",
                    "name": "NetExports",
                    "runner": "net_export_calculator",
                    "in_place": True,
                },
                {
                    "id": "apply_tin_selection_rules",
                    "name": "ApplyTinRules",
                    "runner": "tin_selection_rules",
                    "in_place": True,
                },
            ],
        }
    }
    path = tmp_path / "workflow.yaml"
    path.write_text(yaml.safe_dump(config, allow_unicode=True), encoding="utf-8")
    return path


def test_workflow_config_loads_project_config() -> None:
    entry_path = REPO_ROOT / "scripts" / "process_all.py"
    spec = importlib.util.spec_from_file_location("repo_process_all", entry_path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    config = WorkflowConfig.from_yaml(module.build_workflow_yaml("tin"))
    assert config.name == "tin_chain_excel_catalog_to_financial_curation"
    assert [step.id for step in config.steps] == [
        "excel_catalog_curator",
        "financial_variable_curation",
        "calculate_net_exports",
        "generate_stock_targets",
        "dashboard_data_sorter",
    ]
    assert config.steps[0].runner == "excel_catalog_curator"


def _context() -> WorkflowContext:
    return WorkflowContext(
        project_root=PROJECT_ROOT,
        python=Path(sys.executable),
        provider="mock",
        rules=None,
    )


def test_stock_target_curator_runner_writes_targets(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    _make_small_workbook(source)
    step = StepConfig(
        id="generate_stock_targets",
        name="Skill3",
        runner="stock_target_curator",
        params={
            "industry": "tin",
            "output_path": str(tmp_path / "个股标的_锡.xlsx"),
        },
    )

    result = _stock_target_curator(
        step,
        source,
        tmp_path / "output.xlsx",
        tmp_path / "run",
        _context(),
    )

    assert result.output == source
    assert result.metadata["stock_targets"] == str(tmp_path / "个股标的_锡.xlsx")
    workbook = load_workbook(tmp_path / "个股标的_锡.xlsx", read_only=True)
    assert "个股标的" in workbook.sheetnames
    workbook.close()


def test_dashboard_data_sorter_runner_sorts_catalog(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "指标目录"
    ws.append(
        [
            "#",
            "Sheet Name",
            "Freq",
            "Col",
            "Indicator Name",
            "Unit",
            "板块",
            "大类",
            "子类",
            "数据性质",
            "是否选中",
            "状态说明",
            "指标标签",
        ]
    )
    ws.append(["B-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""])
    ws.append([1, "", "日度", 1, "B 指标", "", "", "价格", "现货价格", "水平值", "否", "", ""])
    ws.append(["A-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""])
    ws.append([1, "", "日度", 1, "A 指标", "", "", "价格", "现货价格", "水平值", "是", "", ""])
    workbook.save(source)

    step = StepConfig(
        id="dashboard_data_sorter",
        name="Skill4",
        runner="dashboard_data_sorter",
    )
    result = _dashboard_data_sorter(
        step,
        source,
        tmp_path / "output.xlsx",
        tmp_path / "run",
        _context(),
    )

    assert result.output == source
    workbook = load_workbook(source, read_only=True)
    ws = workbook["指标目录"]
    names = [
        row[4]
        for row in ws.iter_rows(min_row=1, values_only=True)
        if row[4] not in (None, "", "Indicator Name")
    ]
    assert names == ["A 指标", "B 指标"]
    workbook.close()


@pytest.mark.skipif(
    not SKILL1_SCRIPT.exists(),
    reason="skill source projects are not available on this machine",
)
def test_workflow_end_to_end_mock(tmp_path: Path) -> None:
    source = tmp_path / "new_data.xlsx"
    _make_small_workbook(source)
    config = WorkflowConfig.from_yaml(_write_smoke_config(tmp_path))

    summary = run_workflow(
        config,
        source,
        provider="mock",
        output=tmp_path / "final.xlsx",
        run_root=tmp_path / "runs",
    )

    assert summary["status"] == "COMPLETED"
    assert summary["final_output"] == str((tmp_path / "final.xlsx").resolve())
    final = tmp_path / "final.xlsx"
    assert final.exists()
    workbook = load_workbook(final, read_only=True)
    assert "净出口计算" in workbook.sheetnames
    first = workbook[workbook.sheetnames[0]]
    headers = [cell.value for cell in next(first.iter_rows(min_row=1, max_row=1))]
    assert "是否选中" in headers
    assert "指标标签" in headers
    selected = 0
    for row in first.iter_rows(min_row=2, max_col=13):
        if row[11].value == "是":
            selected += 1
    assert selected >= 3
    workbook.close()
