from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from composite_metric_stage import (  # noqa: E402
    _build_balance_groups,
    _build_market_groups,
    _build_trade_groups,
    _market_key,
    _ratio_series,
    run_composite_stage,
)


HEADER = [
    "#",
    "Sheet Name",
    "Freq",
    "Col",
    "Indicator Name",
    "Unit",
    "板块",
    "大类",
    "子类",
    "数据性质",
    "是否选中",
    "状态说明",
    "指标标签",
]


def make_row(
    sheet,
    title,
    *,
    freq="日度",
    col=1,
    unit="手",
    major="价格",
    sub="成交量",
    selected="是",
    tags=None,
):
    return [
        sheet,
        "",
        freq,
        "",
        "",
        "",
        "",
        "其他",
        "其他",
        "",
        "否",
        "",
        "",
    ], [
        col,
        "",
        freq,
        col,
        title,
        unit,
        "板块",
        major,
        sub,
        "水平值",
        selected,
        "",
        json.dumps(tags or {}, ensure_ascii=False),
    ]


def _write_data_sheet(wb, name: str, values: list[tuple[str, float]]) -> None:
    ws = wb.create_sheet(name)
    ws.append(["日期", "值"])
    ws.append(["", ""])
    ws.append(["", ""])
    for date_key, value in values:
        ws.append([date_key, value])


def _make_workbook(path: Path, directory_rows: list[list]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "指标目录"
    ws.append(HEADER)
    for row in directory_rows:
        ws.append(row)
    wb.save(path)
    return wb


def _read_tags(path: Path, title: str) -> dict:
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[4] == title:
            tags = json.loads(str(row[12]))
            wb.close()
            return tags
    wb.close()
    raise AssertionError(f"title not found: {title}")


def test_market_key_separates_contract_frequency_unit_status() -> None:
    base = {
        "title": "SHFE: 锡: 主力合约: 成交量: 日度",
        "frequency": "日度",
        "unit": "手",
        "nature": "实际",
        "tags": {},
    }
    assert _market_key(base) != _market_key({**base, "title": "SHFE: 锡: 01合约: 成交量: 日度"})
    assert _market_key(base) != _market_key(
        {**base, "title": "SHFE: 锡: 主力合约: 成交量: 周度", "frequency": "周度"}
    )
    assert _market_key(base) != _market_key({**base, "unit": "万吨"})
    assert _market_key(base) != _market_key({**base, "nature": "预测"})


def test_market_activity_complete_and_ratio(tmp_path) -> None:
    path = tmp_path / "input.xlsx"
    header_vol, vol_row = make_row("vol", "SHFE: 锡: 主力合约: 成交量: 日度", sub="成交量")
    header_oi, oi_row = make_row("oi", "SHFE: 锡: 主力合约: 持仓量: 日度", sub="持仓量")
    wb = _make_workbook(path, [header_vol, vol_row, header_oi, oi_row])
    _write_data_sheet(wb, "vol", [("2024-01-01", 100.0), ("2024-01-02", 200.0)])
    _write_data_sheet(wb, "oi", [("2024-01-01", 500.0), ("2024-01-02", 600.0)])
    wb.save(path)

    output = tmp_path / "output.xlsx"
    report = run_composite_stage(path, output, tmp_path / "report.json")

    assert report["stats"]["market_activity_complete"] == 1
    assert report["groups"][0]["derived_metric"] == "oi_volume_ratio"
    assert report["groups"][0]["value_count"] == 2
    tags = _read_tags(output, "SHFE: 锡: 主力合约: 成交量: 日度")
    assert tags["composite"]["composite_type"] == "market_activity"
    assert tags["composite"]["composite_role"] == "volume"
    assert tags["composite"]["composite_status"] == "MATCHED"


def test_market_activity_incomplete_when_only_volume(tmp_path) -> None:
    header_vol, vol_row = make_row("vol", "SHFE: 锡: 主力合约: 成交量: 日度", sub="成交量")
    path = tmp_path / "input.xlsx"
    _make_workbook(path, [header_vol, vol_row])

    output = tmp_path / "output.xlsx"
    report = run_composite_stage(path, output, tmp_path / "report.json")

    assert report["stats"]["market_activity_incomplete"] == 1
    assert report["groups"][0]["composite_status"] == "INCOMPLETE"


def test_volume_zero_does_not_generate_inf(tmp_path) -> None:
    path = tmp_path / "input.xlsx"
    header_vol, vol_row = make_row("vol", "SHFE: 锡: 主力合约: 成交量: 日度", sub="成交量")
    header_oi, oi_row = make_row("oi", "SHFE: 锡: 主力合约: 持仓量: 日度", sub="持仓量")
    wb = _make_workbook(path, [header_vol, vol_row, header_oi, oi_row])
    _write_data_sheet(wb, "vol", [("2024-01-01", 0.0), ("2024-01-02", 200.0)])
    _write_data_sheet(wb, "oi", [("2024-01-01", 500.0), ("2024-01-02", 600.0)])
    wb.save(path)

    output = tmp_path / "output.xlsx"
    report = run_composite_stage(path, output, tmp_path / "report.json")
    group = report["groups"][0]
    assert group["value_count"] == 1
    assert "volume_zero" in group["reason"]
    assert "inf" not in json.dumps(report)


def test_trade_complete_with_reconciliation_warning(tmp_path) -> None:
    path = tmp_path / "input.xlsx"
    header_imp, imp_row = make_row(
        "imp",
        "中国海关: 碳酸锂进口量: 总计: 月度",
        freq="月度",
        unit="吨",
        major="进出口",
        sub="进口",
    )
    header_exp, exp_row = make_row(
        "exp",
        "中国海关: 碳酸锂出口量: 总计: 月度",
        freq="月度",
        unit="吨",
        major="进出口",
        sub="出口",
    )
    header_net, net_row = make_row(
        "net",
        "中国海关: 碳酸锂净出口: 月度",
        freq="月度",
        unit="吨",
        major="进出口",
        sub="净出口",
    )
    wb = _make_workbook(path, [header_imp, imp_row, header_exp, exp_row, header_net, net_row])
    _write_data_sheet(wb, "imp", [("2024-01-01", 100.0)])
    _write_data_sheet(wb, "exp", [("2024-01-01", 200.0)])
    _write_data_sheet(wb, "net", [("2024-01-01", 500.0)])
    wb.save(path)

    output = tmp_path / "output.xlsx"
    report = run_composite_stage(path, output, tmp_path / "report.json")

    assert report["stats"]["trade_complete"] == 1
    assert "NET_EXPORT_RECONCILIATION_WARNING" in report["groups"][0]["reason"]


def test_trade_different_product_or_unit_not_matched() -> None:
    rec_a = {
        "title": "中国海关: 碳酸锂进口量: 总计: 月度",
        "frequency": "月度",
        "unit": "吨",
        "major": "进出口",
        "sub": "进口",
        "selected": "是",
        "tags": {},
    }
    rec_b = {
        "title": "中国海关: 磷酸铁锂出口量: 总计: 月度",
        "frequency": "月度",
        "unit": "吨",
        "major": "进出口",
        "sub": "出口",
        "selected": "是",
        "tags": {},
    }
    rec_c = {
        "title": "中国海关: 碳酸锂出口量: 总计: 月度",
        "frequency": "月度",
        "unit": "千克",
        "major": "进出口",
        "sub": "出口",
        "selected": "是",
        "tags": {},
    }
    groups = _build_trade_groups([rec_a, rec_b, rec_c])
    assert len(groups) == 3
    assert all(len(v["import"]) + len(v["export"]) == 1 for v in groups.values())


def test_balance_groups_do_not_enter_trade() -> None:
    records = [
        {
            "title": "SMM: 锂矿供需平衡: 产量: 月度",
            "frequency": "月度",
            "unit": "吨",
            "major": "平衡",
            "sub": "平衡",
            "selected": "是",
            "tags": {},
        },
        {
            "title": "SMM: 锂矿供需平衡: 进口量: 月度",
            "frequency": "月度",
            "unit": "吨",
            "major": "平衡",
            "sub": "平衡",
            "selected": "是",
            "tags": {},
        },
    ]
    assert len(_build_trade_groups(records)) == 0
    assert len(_build_balance_groups(records)) == 1


def test_ratio_series_skips_zero_volume() -> None:
    wb = Workbook()
    _write_data_sheet(wb, "vol", [("2024-01-01", 0.0), ("2024-01-02", 100.0)])
    _write_data_sheet(wb, "oi", [("2024-01-01", 500.0), ("2024-01-02", 600.0)])
    volume_rec = {"sheet": "vol", "col": 1}
    position_rec = {"sheet": "oi", "col": 1}
    rows, reasons = _ratio_series(wb, volume_rec, position_rec)
    assert [row["date"] for row in rows] == ["2024-01-02"]
    assert any(reason.startswith("volume_zero") for reason in reasons)
