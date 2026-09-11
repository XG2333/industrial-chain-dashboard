# -*- coding: utf-8 -*-
"""仅重算"二次筛选是否保留"列（保留现有"是否选中"），用最新二次筛选规则。

用法：python recompute_secondary_only.py <input.xlsx> <output.xlsx>
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from secondary_selection import secondary_selection_decision  # noqa: E402


def main() -> int:
    if len(sys.argv) != 3:
        print("usage: recompute_secondary_only.py <input.xlsx> <output.xlsx>")
        return 2
    input_path, output_path = Path(sys.argv[1]), Path(sys.argv[2])

    wb = load_workbook(input_path)
    ws = wb["指标目录"]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    name_idx = header.index("Indicator Name")
    freq_idx = header.index("Freq")
    sector_idx = header.index("板块")
    selected_idx = header.index("是否选中")
    final_idx = header.index("二次筛选是否保留")

    changed = 0
    kept = 0
    for row in ws.iter_rows(min_row=2):
        title = str(row[name_idx].value).strip() if row[name_idx].value is not None else ""
        freq = str(row[freq_idx].value).strip() if row[freq_idx].value is not None else ""
        sector = str(row[sector_idx].value).strip() if row[sector_idx].value is not None else ""
        selected = str(row[selected_idx].value).strip() if row[selected_idx].value is not None else ""
        keep, reason = secondary_selection_decision(title, freq, sector, selected == "是")
        new_val = "是" if keep else "否"
        if str(row[final_idx].value).strip() != new_val:
            changed += 1
        row[final_idx].value = new_val
        if keep:
            kept += 1

    wb.save(output_path)
    print(f"total rows: {ws.max_row - 1}, kept={kept}, changed={changed}")
    print(f"output: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
