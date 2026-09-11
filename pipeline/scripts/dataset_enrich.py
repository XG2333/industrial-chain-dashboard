# -*- coding: utf-8 -*-
"""Phase 3b：Production-context Dataset enrichment（schema_version 1.0 → 2.0）。

把 classification/human_confirmed 从"只有名称+人工 expected"升级为可复现
production ai_disambiguation 输入结构的 Dataset：
- 从原 run 的 final_output.xlsx（该 run 实际运行 indicator_tags 后的产物）
  恢复 指标标签（候选分类组合/候选大类/候选子类）与 current 大类/子类 ——
  HISTORICAL_EXACT（run 内真实生产候选，非当前规则重算）。
- 禁止：用人工 expected 反推 candidates；把正确答案塞进候选。
- 每条 item metadata 记录 context_status / candidate_provenance /
  source_run_id / source_artifact / reconstruction_version / schema_version=2.0，
  并保留 preserved_v1 快照（不覆盖历史证据）。
- 输出 eligible 分层（production_equivalent / historical_reconstructed /
  context_limited / ineligible）+ expected_in_candidates + root-cause 分布。

用法：
  python scripts/dataset_enrich.py --human-cases <csv> --run-dir <run>
      [--final-xlsx <path>] [--experiment-diff <case_diff.csv>]
      [--dry-run|--apply|--verify]
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]

try:
    from dotenv import load_dotenv

    load_dotenv(PROJECT_ROOT / ".env")
except Exception:  # noqa: BLE001
    pass

sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

SCHEMA_VERSION = "2.0"
RECONSTRUCTION_VERSION = "2.0"

CONTEXT_HISTORICAL_EXACT = "HISTORICAL_EXACT"
CONTEXT_HISTORICAL_RECONSTRUCTED = "HISTORICAL_RECONSTRUCTED"
CONTEXT_CURRENT_RULES = "CURRENT_RULE_RECONSTRUCTED"
CONTEXT_UNAVAILABLE = "UNAVAILABLE"

ALLOWED_MAJORS = {
    "价格", "成本利润", "进出口", "库存", "需求", "供给", "平衡", "其他",
    "供需-进出口", "供需-库存", "供需-需求", "供需-供给", "供需-平衡",
}


def load_directory_tags(xlsx_path: str | Path) -> dict[str, dict]:
    """final_output 指标目录 → {name: {tags, major, sub}}（候选生产真源）。"""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["指标目录"]
    header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    i_name = header.index("Indicator Name")
    i_tag = header.index("指标标签") if "指标标签" in header else None
    i_major = header.index("大类") if "大类" in header else None
    i_sub = header.index("子类") if "子类" in header else None
    out: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = str(row[i_name] or "").strip()
        if not name:
            continue
        tags: dict = {}
        if i_tag is not None and row[i_tag] is not None:
            try:
                tags = json.loads(str(row[i_tag]).strip()) or {}
            except Exception:  # noqa: BLE001
                tags = {}
        out[name] = {
            "tags": tags if isinstance(tags, dict) else {},
            "major": str(row[i_major] or "").strip() if i_major is not None else "",
            "sub": str(row[i_sub] or "").strip() if i_sub is not None else "",
        }
    wb.close()
    return out


def load_experiment_diff(csv_path: str | Path) -> dict[str, dict]:
    """Phase 3 case_diff.csv → {case_id: {v1_output, v1_correct}}（root cause 用）。"""
    out: dict[str, dict] = {}
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            out[row["case_id"]] = {
                "v1_output": json.loads(row.get("v1_output") or "{}"),
                "v1_correct": row.get("v1_correct") == "True",
            }
    return out


def enrich_case(hc: dict, dir_info: dict, v1_result: dict | None) -> dict:
    """单条 case 重建（candidates 只来自 run 产物，禁止用 expected 反推）。"""
    name = hc.get("original_name", "")
    info = dir_info.get(name, {})
    tags = info.get("tags", {})
    combos = tags.get("候选分类组合")
    majors = tags.get("候选大类")
    subs = tags.get("候选子类")
    has_candidates = bool(combos or majors or subs)

    expected = {
        k: str(hc.get(k) or "").strip()
        for k in ("human_major", "human_sub", "human_selected")
        if str(hc.get(k) or "").strip()
    }
    expected_out = {}
    for k, field in (("human_major", "大类"), ("human_sub", "子类"), ("human_selected", "是否选中")):
        if str(hc.get(k) or "").strip():
            expected_out[field] = str(hc.get(k)).strip()

    context_status = CONTEXT_HISTORICAL_EXACT if has_candidates else CONTEXT_UNAVAILABLE
    # 当前实现只从 run 产物恢复 → 无 candidates 即为 UNAVAILABLE（不伪造）

    exp_major = expected.get("human_major", "")
    exp_sub = expected.get("human_sub", "")
    expected_in_candidates = None
    if has_candidates and exp_major:
        expected_in_candidates = any(
            str(c.get("大类") or "").strip() == exp_major
            and (not exp_sub or str(c.get("子类") or "").strip() == exp_sub)
            for c in (combos or [])
            if isinstance(c, dict)
        ) or (
            exp_major in [str(m).strip() for m in (majors or [])]
        )

    # eligibility 分层
    if not has_candidates:
        eligibility = "context_limited"
    elif context_status == CONTEXT_HISTORICAL_EXACT:
        eligibility = "production_equivalent"
    else:
        eligibility = "historical_reconstructed"

    # root cause（结合 v1 实验输出；无实验输出时标记 UNKNOWN）
    root_cause = "CORRECT"
    if not expected_out.get("大类"):
        root_cause = "EXPECTED_LABEL_MISMATCH"
    elif not has_candidates:
        root_cause = "NO_CANDIDATES"
    elif not expected_in_candidates:
        root_cause = "WRONG_CANDIDATES"
    elif exp_major not in ALLOWED_MAJORS:
        root_cause = "EXPECTED_LABEL_MISMATCH"
    elif v1_result is None:
        root_cause = "UNKNOWN_NO_EXPERIMENT"
    elif v1_result.get("v1_correct"):
        root_cause = "CORRECT"
    elif not v1_result.get("v1_output"):
        root_cause = "OUTPUT_PARSE_ERROR"
    else:
        root_cause = "PROMPT_SELECTION_ERROR"

    item = {
        "id": f"golden_{hc.get('industry', 'unknown')}_{hc.get('commodity', 'UNKNOWN')}_{hc.get('physical_variable_id')}",
        "input": {
            "physical_variable_id": hc.get("physical_variable_id", ""),
            "original_name": name,
            "industry": hc.get("industry", "unknown"),
            "commodity": hc.get("commodity", "UNKNOWN"),
            "sheet": hc.get("sheet", ""),
            "unit": hc.get("unit", ""),
            "frequency": hc.get("frequency", ""),
            "current_major": info.get("major", ""),
            "current_sub": info.get("sub", ""),
            "candidates": combos or [],
            "current_tags": {k: v for k, v in tags.items() if k in ("候选大类", "候选子类")},
        },
        "expected_output": expected_out,
        "metadata": {
            "source": "human_review",
            "source_case_id": hc.get("source_case_id", ""),
            "physical_variable_id": hc.get("physical_variable_id", ""),
            "audit_crosscheck": hc.get("audit_crosscheck", "-"),
            "original_run_id": hc.get("workflow_run_id", ""),
            "schema_version": SCHEMA_VERSION,
            "reconstruction_version": RECONSTRUCTION_VERSION,
            "context_status": context_status,
            "candidate_provenance": "run_artifact:indicator_tags:final_output.xlsx" if has_candidates else "",
            "source_artifact": "final/final_output.xlsx" if has_candidates else "",
            "expected_in_candidates": expected_in_candidates,
            "eligibility": eligibility,
            "root_cause": root_cause,
            "preserved_v1": {
                "schema_version": "1.0",
                "expected_output": expected_out,
            },
        },
    }
    return item


def enrich_all(human_cases: list[dict], dir_info: dict, v1_results: dict[str, dict] | None,
               industry: str = "unknown", commodity: str = "UNKNOWN") -> list[dict]:
    out = []
    for hc in human_cases:
        case_id = f"golden_{industry}_{commodity}_{hc.get('physical_variable_id')}"
        v1 = (v1_results or {}).get(case_id)
        hc = {**hc, "industry": industry, "commodity": commodity}
        out.append(enrich_case(hc, dir_info, v1))
    return out


def summarize(items: list[dict]) -> dict:
    layers = ["production_equivalent", "historical_reconstructed", "context_limited", "ineligible"]
    elig = {k: 0 for k in layers}
    roots: dict[str, int] = {}
    eic = {"true": 0, "false": 0, "na": 0}
    for it in items:
        m = it["metadata"]
        elig[m.get("eligibility", "ineligible")] = elig.get(m.get("eligibility", "ineligible"), 0) + 1
        roots[m.get("root_cause", "OTHER")] = roots.get(m.get("root_cause", "OTHER"), 0) + 1
        eic_val = m.get("expected_in_candidates")
        eic["true" if eic_val is True else ("false" if eic_val is False else "na")] += 1
    return {"eligible_layers": elig, "root_cause_distribution": dict(sorted(roots.items())),
            "expected_in_candidates": eic}


def make_provider(client):
    from evaluation_dataset_export import LangfuseDatasetProvider

    return LangfuseDatasetProvider(client)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Phase 3b dataset enrichment")
    parser.add_argument("--human-cases", required=True)
    parser.add_argument("--run-dir", required=True, help="原 run（source_run_id 与输出目录）")
    parser.add_argument("--final-xlsx", default=None, help="默认 <run>/final/final_output.xlsx")
    parser.add_argument("--experiment-diff", default=None, help="Phase 3 case_diff.csv（root cause）")
    parser.add_argument("--industry", default="tin", help="数据集行业（case_id 构成）")
    parser.add_argument("--commodity", default="TIN", help="数据集品类（case_id 构成）")
    parser.add_argument("--dataset-name", default="classification/human_confirmed")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--verify", action="store_true")
    args = parser.parse_args(argv)
    if args.apply and args.dry_run:
        parser.error("--apply 与 --dry-run 互斥")

    from evaluation_dataset_export import load_human_cases

    human_cases = load_human_cases(Path(args.human_cases))
    run_dir = Path(args.run_dir).resolve()
    final_xlsx = Path(args.final_xlsx) if args.final_xlsx else run_dir / "final" / "final_output.xlsx"
    dir_info = load_directory_tags(final_xlsx)
    v1_results = load_experiment_diff(Path(args.experiment_diff)) if args.experiment_diff else None
    items = enrich_all(human_cases, dir_info, v1_results,
                       industry=args.industry, commodity=args.commodity)
    summary = summarize(items)
    summary["total"] = len(items)
    summary["schema_version"] = SCHEMA_VERSION
    summary["source_run_id"] = run_dir.name
    summary["mode"] = "dry_run" if not args.apply else "apply"

    out_dir = run_dir / "audit" / "evaluation"
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "enriched_items.jsonl").write_text(
        "\n".join(json.dumps(i, ensure_ascii=False) for i in items) + "\n", encoding="utf-8"
    )

    if args.apply:
        from langfuse import Langfuse

        client = Langfuse(
            public_key=os.getenv("LANGFUSE_PUBLIC_KEY", ""),
            secret_key=os.getenv("LANGFUSE_SECRET_KEY", ""),
            base_url=os.getenv("LANGFUSE_BASE_URL", "https://us.cloud.langfuse.com"),
            environment=os.getenv("LANGFUSE_TRACING_ENVIRONMENT", "development"),
            debug=False,
        )
        provider = make_provider(client)
        stats = {"created": 0, "updated": 0, "unchanged": 0, "failed": 0}
        existing = provider.get_dataset_items(args.dataset_name) or {}
        for it in items:
            try:
                prev = existing.get(it["id"])
                if prev is None:
                    stats["created"] += 1
                else:
                    stats["updated"] += 1
                provider.upsert_item(args.dataset_name, it)
            except Exception as exc:  # noqa: BLE001
                stats["failed"] += 1
                it["error"] = str(exc)[:200]
        summary["dataset_write"] = stats
        client.flush()

    if args.verify:
        from langfuse import Langfuse

        client = Langfuse(
            public_key=os.getenv("LANGFUSE_PUBLIC_KEY", ""),
            secret_key=os.getenv("LANGFUSE_SECRET_KEY", ""),
            base_url=os.getenv("LANGFUSE_BASE_URL", "https://us.cloud.langfuse.com"),
            environment=os.getenv("LANGFUSE_TRACING_ENVIRONMENT", "development"),
            debug=False,
        )
        provider = make_provider(client)
        existing = provider.get_dataset_items(args.dataset_name) or {}
        v2 = sum(1 for it in existing.values() if (it.get("metadata") or {}).get("schema_version") == "2.0")
        summary["verify"] = {"items": len(existing), "schema_v2": v2}

    (out_dir / "enrichment_report.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
