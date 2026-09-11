# -*- coding: utf-8 -*-
"""Generic net-export calculation shared by all industry chains."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from apply_indicator_tags import apply_indicator_tags
from import_export_rules import (
    compute_net_series,
    find_missing_net_export_pairs,
    net_title,
)


def run_net_exports_on_workbook(wb) -> int:
    """在已加载的 wb 上执行净出口计算（共享 wb，不 load/save）。

    供合并链（workflow_merged_stage.py）复用；calculate_net_exports 包装
    load/save。返回计算出的净出口对数；无缺失对时返回 0 且不修改 wb。
    """
    ws = wb[wb.sheetnames[0]]
    current_sheet = ""
    directory_rows = []

    for row in ws.iter_rows(min_row=2):
        c0 = str(row[0].value).strip() if row[0].value is not None else ""
        if not c0 or "sht" in c0:
            continue
        if not c0.isdigit():
            current_sheet = c0
            continue
        c3 = str(row[3].value).strip() if row[3].value is not None else ""
        if not c3.isdigit():
            continue
        directory_rows.append(
            {
                "sheet": current_sheet,
                "col": int(c3),
                "title": str(row[4].value) if row[4].value is not None else "",
                "unit": str(row[5].value) if row[5].value is not None else "",
                "sector": str(row[6].value) if row[6].value is not None else "",
                "frequency": str(row[2].value) if row[2].value is not None else "",
                "major": str(row[7].value) if row[7].value is not None else "",
                "sub": str(row[8].value) if row[8].value is not None else "",
            }
        )

    missing = find_missing_net_export_pairs(directory_rows)
    if not missing:
        print("No missing net exports")
        return 0

    sheet_name = "净出口计算"
    if sheet_name not in wb.sheetnames:
        ws_calc = wb.create_sheet(sheet_name)
        ws_calc.cell(1, 1, "日期")
    else:
        ws_calc = wb[sheet_name]

    existing_cols = ws_calc.max_column
    seq_max = 0
    for row in ws.iter_rows(min_row=2, max_col=1):
        if row[0].value is not None and str(row[0].value).isdigit():
            seq_max = max(seq_max, int(row[0].value))

    directory_start = ws.max_row + 1
    row_idx = directory_start
    first_major = str(missing[0]["import_row"].get("major") or "进出口")
    first_sector = str(missing[0]["import_row"].get("sector") or "其他")
    ws.cell(row_idx, 1, sheet_name)
    ws.cell(row_idx, 2, "")
    ws.cell(row_idx, 3, "月度")
    ws.cell(row_idx, 7, first_sector)
    row_idx += 1

    computed_pairs = 0
    for pair in missing:
        imp = pair["import_row"]
        exp = pair["export_row"]
        series = compute_net_series(
            wb,
            pair["import_sheet"],
            imp["col"],
            pair["export_sheet"],
            exp["col"],
        )
        if not series:
            print(f"Skip {pair['key']}: no aligned data")
            continue

        col = existing_cols + 1
        existing_cols += 1
        title = net_title(pair["key"], pair["dimension"], "月度")
        unit = imp.get("unit") or exp.get("unit") or "吨"

        ws_calc.cell(1, col + 1, title)
        ws_calc.cell(2, col + 1, unit)
        ws_calc.cell(3, col + 1, "月度")
        for i, point in enumerate(series):
            ws_calc.cell(4 + i, 1, point["date"])
            ws_calc.cell(4 + i, col + 1, point["value"])

        seq_max += 1
        ws.cell(row_idx, 1, seq_max)
        ws.cell(row_idx, 2, "")
        ws.cell(row_idx, 3, "月度")
        ws.cell(row_idx, 4, col)
        ws.cell(row_idx, 5, title)
        ws.cell(row_idx, 6, unit)
        ws.cell(row_idx, 7, imp.get("sector") or first_sector)
        ws.cell(row_idx, 8, imp.get("major") or first_major)
        ws.cell(row_idx, 9, "净出口")
        ws.cell(row_idx, 10, "流量")
        ws.cell(row_idx, 11, 0.8)
        ws.cell(row_idx, 12, "是")
        ws.cell(row_idx, 13, "确定性净出口计算")
        row_idx += 1
        computed_pairs = computed_pairs + 1
        print(f"Computed {title}: {len(series)} points")

    print(f"Net exports added: {computed_pairs} pairs")
    return computed_pairs


def calculate_net_exports(path: Path) -> Path:
    path = Path(path).resolve()
    tmp = path.with_name(f"{path.stem}__tmp_net.xlsx")
    shutil.copy2(path, tmp)

    wb = load_workbook(tmp, data_only=False)
    count = run_net_exports_on_workbook(wb)
    if count == 0 and "净出口计算" not in wb.sheetnames:
        # 无缺失对且未创建计算 sheet：文件未被修改，直接返回
        wb.close()
        tmp.unlink(missing_ok=True)
        return path
    wb.save(tmp)
    wb.close()
    apply_indicator_tags(tmp)
    try:
        tmp.replace(path)
    except PermissionError:
        print(f"Permission error writing {path}; temp file kept at {tmp}")
        raise
    print(f"Net exports added to {path}: {count} pairs")
    return path


if __name__ == "__main__":
    paths = [Path(arg) for arg in sys.argv[1:] if not arg.startswith("-")]
    if not paths:
        print("Usage: python calculate_net_exports.py <file.xlsx>")
        sys.exit(1)
    for path in paths:
        calculate_net_exports(path)
