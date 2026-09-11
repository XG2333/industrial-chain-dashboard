# -*- coding: utf-8 -*-
"""修复锡工作簿「指标目录」的 Sheet Name 列（显式补全真实数据 sheet）。

与 fix_silicon_catalog_sheets.py 同法：数据 sheet 名称行（tin_catalog.discover，
权威读取）反向匹配目录行（Indicator Name + Col + Freq + Unit），把 Sheet Name
列写为真实 sheet 名，使 server 逐行直读、不再依赖被重排工具打乱的分组行。

用法：
  python fix_tin_catalog_sheets.py [--report <path>]
"""
import argparse
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from tin_catalog import discover

DATA_DIR = Path(__file__).resolve().parents[2] / "本地可视化dashboard" / "data"
FILE = DATA_DIR / "锡产业链数据_workflow_ai.xlsx"
CAT_SHEET = "指标目录"


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", nargs="?", const=True, default="")
    args = parser.parse_args()

    if not FILE.exists():
        print("missing:", FILE)
        return 1

    backup = DATA_DIR / (FILE.stem + f"_backup_{Path(__file__).stem}.xlsx")
    shutil.copy2(FILE, backup)
    print(f"backup -> {backup.name}")

    discovered = discover(FILE)
    by_name: dict[str, list] = {}
    sheet_names = set()
    for s in discovered:
        sn = s["sn"]
        sheet_names.add(sn)
        for ind in s["inds"]:
            nm = str(ind["name"]).strip()
            by_name.setdefault(nm, []).append(
                (sn, ind["col"], str(ind["unit"] or "").strip(), str(ind["freq"] or "").strip())
            )
    print(f"data sheets={len(discovered)} indicator labels={sum(len(v) for v in by_name.values())}")

    wb = load_workbook(FILE, data_only=False)  # data_only=False: 保留公式
    if CAT_SHEET not in wb.sheetnames:
        print("catalog sheet not found:", CAT_SHEET)
        return 1
    ws = wb[CAT_SHEET]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    for want in ("Sheet Name", "Freq", "Col", "Indicator Name", "Unit"):
        if want not in header:
            print("catalog header missing:", want)
            return 1
    hi = {name: i for i, name in enumerate(header)}
    c_sheet, c_freq, c_col, c_name, c_unit = (
        hi["Sheet Name"], hi["Freq"], hi["Col"], hi["Indicator Name"], hi["Unit"]
    )

    cur_group = None
    changes: list[tuple[int, str, str, str]] = []
    filled = multi_manual = not_found = skipped_ref = 0

    for r_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        vals = ["" if c.value is None else str(c.value).strip() for c in row]
        name = vals[c_name] if len(vals) > c_name else ""
        if not name:
            continue
        col = vals[c_col] if len(vals) > c_col else ""
        if not col.isdigit():
            c0 = vals[0] if vals else ""
            if c0 and not c0.isdigit() and "(" not in c0 and "[" not in c0 and "统计" not in c0:
                cur_group = c0
            continue
        freq = vals[c_freq] if len(vals) > c_freq else ""
        unit = vals[c_unit] if len(vals) > c_unit else ""
        col0 = int(col)

        if name.startswith("'"):
            skipped_ref += 1
            continue

        cands = by_name.get(name)
        if not cands:
            not_found += 1
            continue
        by_col = [c for c in cands if c[1] == col0]
        if len(by_col) == 1:
            cands = by_col
        else:
            filt = [c for c in cands if (not freq or c[3] == freq) and (not unit or c[2] == unit)]
            cands = filt if len(filt) == 1 else (by_col or cands)
        if len(cands) != 1:
            multi_manual += 1
            continue
        real_sheet, real_col, _, _ = cands[0]
        if real_sheet not in sheet_names:
            not_found += 1
            continue

        old = vals[c_sheet] if len(vals) > c_sheet else ""
        if old != real_sheet:
            ws.cell(row=r_idx, column=c_sheet + 1, value=real_sheet)
            filled += 1
            if old:
                changes.append((r_idx, old, real_sheet, name))
            elif cur_group and cur_group != real_sheet:
                changes.append((r_idx, f"<组:{cur_group}>", real_sheet, name))
    wb.save(FILE)
    print(f"filled={filled} multi_manual={multi_manual} not_found={not_found} ref_skipped={skipped_ref}")

    # 自检：全部指标行 (sheet,col) 与数据 sheet 名称行标签一致
    verify_err = 0
    for r_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        vals = ["" if c.value is None else str(c.value).strip() for c in row]
        name = vals[c_name] if len(vals) > c_name else ""
        col = vals[c_col] if len(vals) > c_col else ""
        sh = vals[c_sheet] if len(vals) > c_sheet else ""
        if not name or not col.isdigit() or name.startswith("'"):
            continue
        if not sh:
            verify_err += 1
            continue
        col0 = int(col)
        if not any(c[0] == sh and c[1] == col0 for c in by_name.get(name, [])):
            verify_err += 1
            print(f"  VERIFY FAIL r{r_idx}: sheet={sh} col={col0} name={name}")

    lines = [
        f"锡指标目录 Sheet Name 修复报告 {FILE.name}",
        f"备份: {backup.name}",
        f"指标行填表: {filled} | 多候选需人工: {multi_manual} | 未找到: {not_found} | 引用行跳过: {skipped_ref}",
        f"自检: {'PASS' if verify_err == 0 else f'FAIL x{verify_err}'}",
        f"错位/变更行数: {len(changes)}",
        "",
        "=== 错位行清单 (row | 原分组 -> 真实sheet | 指标) ===",
    ]
    for r_idx, old, real, name in changes:
        lines.append(f"r{r_idx} | {old} -> {real} | {name[:80]}")
    out = "\n".join(lines)
    if args.report is not False:
        rep_path = (
            Path(args.report)
            if isinstance(args.report, str) and args.report
            else Path(__file__).resolve().parent / "tin_sheetfix_report.txt"
        )
        rep_path.write_text(out, encoding="utf-8")
        print("report ->", rep_path)
    else:
        print(out)
    print("DONE verify_err=", verify_err)
    return 0 if verify_err == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
