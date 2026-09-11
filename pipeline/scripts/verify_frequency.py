# -*- coding: utf-8 -*-
"""Check and correct indicator frequency from actual time intervals.

Deterministic only; no AI. For each indicator, infer frequency from the median
interval between consecutive dates and compare it with the directory Freq.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import warnings
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parent))

from catalog_utils import find_header_rows, resolve_annualized_frequency


warnings.simplefilter("ignore", UserWarning)

FREQ_ORDER = ("日度", "周度", "月度", "季度", "年度")


def is_indicator_row(values: list) -> bool:
    if not values or values[0] is None:
        return False
    raw = str(values[0]).strip()
    return bool(re.fullmatch(r"\d+", raw)) and bool(values[3]) and bool(values[4])


def infer_freq_from_dates(dates: list[pd.Timestamp]) -> str | None:
    if len(dates) < 2:
        return None
    diffs = sorted(
        (dates[i + 1] - dates[i]).days
        for i in range(len(dates) - 1)
        if (dates[i + 1] - dates[i]).days > 0
    )
    if not diffs:
        return None
    median = diffs[len(diffs) // 2]
    if median <= 2:
        return "日度"
    if median <= 10:
        return "周度"
    if median <= 45:
        return "月度"
    if median <= 120:
        return "季度"
    return "年度"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Verify indicator frequency from data intervals.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--report", default=None)
    parser.add_argument(
        "--data-file",
        default=None,
        help="完整数据文件（含各数据 sheets）。频率推断需要读数据序列，"
        "写入仍在 --input 文件。默认与 --input 相同。",
    )
    args = parser.parse_args(argv)

    input_path = Path(args.input).resolve()
    output_path = Path(args.output).resolve()
    report_path = Path(args.report).resolve() if args.report else output_path.with_suffix(".md")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(input_path, output_path)

    data_source = Path(args.data_file).resolve() if args.data_file else input_path
    wb = load_workbook(output_path)
    ws = wb.worksheets[0]

    headers = [str(c.value or "").strip() for c in ws[1]]
    tag_col = None
    for idx, header in enumerate(headers, 1):
        if header == "指标标签":
            tag_col = idx
            break

    # 预扫描目录：收集每个 sheet 需要读取的列（日期列 0 + 各指标列及其前一列），
    # 用 pd.ExcelFile + usecols 按需解析，避免全文件全列读取
    sheet_cols: dict[str, set[int]] = {}
    for row in ws.iter_rows(min_row=2):
        values = [c.value for c in row]
        col0 = str(values[0]).strip() if values[0] is not None else ""
        if col0.startswith("[统计]"):
            continue
        if not is_indicator_row(values):
            if col0 and not re.fullmatch(r"\d+", col0) and len(values) > 6 and values[6] is not None:
                current_sheet = col0
            continue
        col = int(str(values[3]).strip())
        cols = sheet_cols.setdefault(current_sheet, set())
        cols.add(0)
        cols.add(col)
        if col > 0:
            cols.add(col - 1)
    xl = pd.ExcelFile(data_source, engine="openpyxl")
    df_cache: dict[tuple[str, tuple], object] = {}

    current_sheet = ""
    corrected = 0
    unchanged = 0
    skipped = 0
    discrepancies = 0
    examples: list[str] = []

    for row in ws.iter_rows(min_row=2):
        values = [c.value for c in row]
        col0 = str(values[0]).strip() if values[0] is not None else ""
        if col0.startswith("[统计]"):
            continue
        if not is_indicator_row(values):
            if col0 and not re.fullmatch(r"\d+", col0) and len(values) > 6 and values[6] is not None:
                current_sheet = col0
            continue

        col = int(str(values[3]).strip())
        name = str(values[4])
        old_freq = str(values[2] or "").strip()
        if current_sheet not in sheet_cols:
            skipped += 1
            continue
        usecols = sorted(sheet_cols[current_sheet])
        key = (current_sheet, tuple(usecols))
        df = df_cache.get(key)
        if df is None:
            try:
                df = pd.read_excel(xl, sheet_name=current_sheet, header=None, dtype=str, usecols=usecols)
            except ValueError:
                df = None
            df_cache[key] = df
        if df is None or col not in usecols or col >= len(usecols):
            skipped += 1
            continue

        hdr_row, unit_row = find_header_rows(df)
        date_col = 0
        if col > 0 and col - 1 in usecols and str(df.iloc[hdr_row, usecols.index(col - 1)]).strip() == "指标名称":
            date_col = col - 1
        dates = pd.to_datetime(df.iloc[:, usecols.index(date_col)], errors="coerce")
        numeric = pd.to_numeric(df.iloc[:, usecols.index(col)], errors="coerce")
        mask = dates.notna() & numeric.notna()
        if int(mask.sum()) < 2:
            skipped += 1
            continue
        ordered = sorted(dates[mask].dropna().unique().tolist())
        inferred = infer_freq_from_dates(ordered)
        if inferred is None:
            skipped += 1
            continue
        annualized = resolve_annualized_frequency(name or "")
        if annualized:
            inferred = annualized

        if not old_freq or old_freq == "未识别":
            if inferred and inferred != old_freq:
                row[2].value = inferred
                corrected += 1
                if tag_col and values[tag_col - 1] is not None:
                    raw = str(values[tag_col - 1]).strip()
                    if raw:
                        try:
                            tags = json.loads(raw)
                            tags["频率"] = inferred
                            row[tag_col - 1].value = json.dumps(tags, ensure_ascii=False)
                        except Exception:
                            pass
                if len(examples) < 50:
                    examples.append(f"- {current_sheet} | {name} | {old_freq or '空'} -> {inferred}")
            else:
                unchanged += 1
        elif inferred and inferred != old_freq:
            discrepancies += 1
            if len(examples) < 50:
                examples.append(f"- {current_sheet} | {name} | 当前={old_freq} 实测间隔推断={inferred}（待复核）")
        else:
            unchanged += 1

    wb.save(output_path)

    report_lines = [
        "# 指标频率校验报告",
        "",
        f"- 输入：`{input_path.name}`",
        f"- 输出：`{output_path.name}`",
        f"- 修正条数：{corrected}",
        f"- 未变更条数：{unchanged}",
        f"- 跳过条数：{skipped}",
        f"- 待复核不一致条数：{discrepancies}",
        "",
        "## 修正示例",
        "",
    ]
    report_lines.extend(examples or ["- 无"])
    report_path.write_text("\n".join(report_lines), encoding="utf-8")

    print(f"Frequency verified: corrected={corrected} unchanged={unchanged} skipped={skipped} discrepancies={discrepancies}")
    print(f"Output: {output_path}")
    print(f"Report: {report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
