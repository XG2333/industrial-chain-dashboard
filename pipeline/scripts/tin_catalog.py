# -*- coding: utf-8 -*-
"""Skill1 catalog generator for the tin industry.

The deterministic rules live in tin_rules.py; this module only owns Excel
scanning, catalog layout, and pickle creation.
"""

from __future__ import annotations

import pickle
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink


sys.path.insert(0, str(Path(__file__).resolve().parent))

from catalog_utils import (
    column_freq,
    find_freq_row,
    find_header_rows,
)
from tin_rules import FREQ_EN, classify_category, classify_sector, detect_freq, detect_indicator_freq


def discover(wp):
    xl = pd.ExcelFile(wp, engine="openpyxl")
    catalog = []
    for sn in xl.sheet_names:
        if len(sn.strip()) < 3:
            continue
        df = pd.read_excel(xl, sheet_name=sn, header=None, dtype=str, nrows=8)
        r, c = df.shape
        if r < 3 or c < 2:
            continue
        sheet_freq = detect_freq(sn)
        hdr_row, unit_row = find_header_rows(df)
        freq_row = find_freq_row(df)
        inds = []
        for ci in range(1, df.shape[1]):
            try:
                n = str(df.iloc[hdr_row, ci]).strip() if pd.notna(df.iloc[hdr_row, ci]) else ""
                if not n or n.lower() == "nan":
                    continue
                u = str(df.iloc[unit_row, ci]).strip() if (
                    unit_row < r and pd.notna(df.iloc[unit_row, ci])
                ) else ""
                if n in ("指标名称", "指标", "指标Id") or u in ("单位", "频率", "指标Id"):
                    continue
                inds.append(dict(col=ci, name=n, unit=u, freq=column_freq(df, freq_row, ci, n, sheet_freq)))
            except Exception:
                pass
        if inds:
            catalog.append(dict(sn=sn, freq=sheet_freq, cnt=len(inds), inds=inds))
    return catalog


def build_catalog(wp, catalog=None):
    if catalog is None:
        catalog = discover(wp)
    wb = load_workbook(wp)
    first_ws = wb.worksheets[0]
    first_a1 = str(first_ws.cell(row=1, column=1).value or "")
    if first_a1 == "#":
        ws = first_ws
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet("指标目录", 0)

    hf = PatternFill("solid", fgColor="4472C4")
    gf = PatternFill("solid", fgColor="5B9BD5")
    sf_p = PatternFill("solid", fgColor="D6E4F0")
    lf = PatternFill("solid", fgColor="F2F2F2")
    cf_font = Font(name="Arial", size=10)
    tb = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    ac = Alignment(horizontal="center", vertical="center")
    for i, w in enumerate([5, 32, 8, 8, 64, 12, 12]):
        ws.column_dimensions[get_column_letter(i + 1)].width = w

    headers = ["#", "Sheet Name", "Freq", "Col", "Indicator Name", "Unit", "板块"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
        c.fill = hf
        c.alignment = ac
        c.border = tb

    row = 2
    seq = 0
    for freq_label in ["日度", "周度", "月度", "季度", "年度"]:
        sf = [s for s in catalog if s["freq"] == freq_label]
        if not sf:
            continue
        ti = sum(s["cnt"] for s in sf)
        for ci in range(1, 8):
            c = ws.cell(row=row, column=ci)
            c.font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
            c.fill = gf
            c.border = tb
            if ci == 1:
                c.value = f"{freq_label} ({len(sf)} sht {ti} ind)"
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        row += 1

        for s in sf:
            sn = s["sn"]
            sec = classify_sector(sn)
            for ci in range(1, 8):
                c = ws.cell(row=row, column=ci)
                c.font = Font(name="Arial", bold=True, size=10)
                c.fill = sf_p
                c.border = tb
                if ci == 1:
                    c.value = sn
                    hl = Hyperlink(ref=c.coordinate, location=f"'{sn}'!A1")
                    c.hyperlink = hl
                    c.font = Font(name="Arial", bold=True, size=10, color="0563C1", underline="single")
                elif ci == 3:
                    c.value = s["freq"]
                elif ci == 7:
                    c.value = sec
                    c.font = Font(name="Arial", bold=True, size=10)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            row += 1

            for ind in s["inds"]:
                seq += 1
                vals = [seq, "", ind["freq"], ind["col"], ind["name"], ind["unit"], sec]
                for ci, v in enumerate(vals, 1):
                    c = ws.cell(row=row, column=ci, value=v)
                    c.font = cf_font
                    c.border = tb
                    if ci in (3, 7):
                        c.alignment = ac
                    if ci == 5:
                        cl = get_column_letter(ind["col"])
                        hl = Hyperlink(ref=c.coordinate, location=f"'{sn}'!{cl}1")
                        c.hyperlink = hl
                        c.font = Font(name="Arial", size=10, color="0563C1", underline="single")
                    if seq % 2 == 0:
                        c.fill = lf
                row += 1

    sr = row + 1
    total_inds = sum(s["cnt"] for s in catalog)
    ws.cell(row=sr, column=1, value="[统计]").font = Font(name="Arial", bold=True, size=11)
    ws.cell(row=sr, column=2, value=f"Sheet 数: {len(catalog)}").font = Font(name="Arial", size=10)
    ws.cell(row=sr, column=4, value=f"指标数: {total_inds}").font = Font(name="Arial", size=10)
    wb.save(wp)
    return len(catalog), total_inds


def build_tin_pickle(wp, out_path, catalog=None):
    if catalog is None:
        catalog = discover(wp)
    records = []
    xl = pd.ExcelFile(wp, engine="openpyxl")
    for s in catalog:
        sn = s["sn"]
        sec = classify_sector(sn)
        cat = classify_category(sn)
        df = pd.read_excel(xl, sheet_name=sn, header=None, dtype=str)
        hdr_row, unit_row = find_header_rows(df)
        data_start = max(hdr_row + 2, unit_row + 1, 3)
        if data_start < df.shape[0]:
            fc = str(df.iloc[data_start, 0]).strip() if pd.notna(df.iloc[data_start, 0]) else ""
            if "频率" in fc:
                data_start += 1
        dates = pd.to_datetime(df.iloc[data_start:, 0], errors="coerce")
        for ind in s["inds"]:
            ci = ind["col"]
            values = pd.to_numeric(df.iloc[data_start:, ci], errors="coerce")
            mask = dates.notna() & values.notna()
            rows = []
            for d, v in zip(dates[mask], values[mask]):
                rows.append({"date": d.strftime("%Y-%m-%d"), "value": float(v)})
            rows.sort(key=lambda x: x["date"], reverse=True)
            unit = ind["unit"]
            div = 1.0
            if unit == "元/吨" and rows and max(x["value"] for x in rows) >= 10000:
                unit = "万元/吨"
                div = 10000.0
            if div != 1.0:
                rows = [{"date": x["date"], "value": round(x["value"] / div, 4)} for x in rows]
            fl = ind.get("freq") or detect_freq(sn)
            freq_en = FREQ_EN.get(fl[0], "monthly") if fl else "monthly"
            records.append(
                {
                    "id": f"tin_{sn}_c{ci}",
                    "title": ind["name"],
                    "unit": unit,
                    "category": cat,
                    "section": sn,
                    "source": "SMM/海关/行业",
                    "freq": freq_en,
                    "data": rows,
                    "sector": sec,
                    "commodity": "TIN",
                }
            )
    with open(out_path, "wb") as f:
        pickle.dump(records, f)
    return len(records)
