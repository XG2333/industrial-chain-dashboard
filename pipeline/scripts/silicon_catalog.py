import pandas as pd
import pickle
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink


sys.path.insert(0, str(Path(__file__).resolve().parent))

from catalog_utils import column_freq, find_freq_row, find_header_rows


_FM = {"日": "daily", "周": "weekly", "月": "monthly", "季": "quarterly", "年": "yearly"}


_SECTOR_RULES = [
    ("硅石石英砂", ["硅石", "石英砂"]),
    ("工业硅", ["工业硅", "金属硅"]),
    ("有机硅", ["有机硅"]),
    ("铝合金", ["铝合金"]),
    ("多晶硅", ["多晶硅"]),
    ("硅片", ["硅片"]),
    ("电池片", ["电池片"]),
    ("组件", ["组件", "T型"]),
    (
        "光伏辅材",
        [
            "光伏玻璃",
            "光伏胶膜",
            "光伏EVA",
            "光伏背板",
            "逆变器",
            "焊带",
            "光伏边框",
            "光伏支架",
            "光伏网板",
            "光伏硅胶",
        ],
    ),
    ("硅基原料", ["硅粉", "三氯氢硅", "原料", "电价"]),
    ("下游需求", ["装机", "发电", "用电", "消纳", "消费", "需求"]),
    ("期货市场", ["期货", "成交", "持仓", "仓单", "基差", "月差"]),
]


_SECTORS = [
    "硅石石英砂",
    "工业硅",
    "有机硅",
    "铝合金",
    "多晶硅",
    "硅片",
    "电池片",
    "组件",
    "光伏辅材",
    "硅基原料",
    "下游需求",
    "期货市场",
    "其他",
]


def detect_freq(sn):
    for k, v in zip(["日", "周", "月", "季", "年"], ["日度", "周度", "月度", "季度", "年度"]):
        if k in sn:
            return v
    return "年度"


def classify_sector(sn):
    for sector, keywords in _SECTOR_RULES:
        if any(k in sn for k in keywords):
            return sector
    return "其他"


def discover(wp):
    xl = pd.ExcelFile(wp, engine="openpyxl")
    catalog = []
    for sn in xl.sheet_names:
        if len(sn.strip()) < 3:
            continue
        df = pd.read_excel(xl, sheet_name=sn, header=None, dtype=str, nrows=10)
        r, c = df.shape
        if r < 3 or c < 2:
            continue
        sheet_freq = detect_freq(sn)
        hdr_row, unit_row = find_header_rows(df)
        freq_row = find_freq_row(df)
        inds = []
        for ci in range(1, df.shape[1]):
            try:
                n = str(df.iloc[hdr_row, ci]).strip() if pd.notna(df.iloc[hdr_row, ci]) else ""
                if not n or n.lower() == "nan":
                    continue
                u = str(df.iloc[unit_row, ci]).strip() if (
                    unit_row < r and pd.notna(df.iloc[unit_row, ci])
                ) else ""
                if n in ("指标名称", "指标") or u in ("单位", "频率"):
                    continue
                inds.append(dict(col=ci, name=n, unit=u, freq=column_freq(df, freq_row, ci, n, sheet_freq)))
            except Exception:
                pass
        if inds:
            catalog.append(dict(sn=sn, freq=sheet_freq, cnt=len(inds), inds=inds))
    return catalog


def build_catalog(wp, catalog=None):
    if catalog is None:
        catalog = discover(wp)

    wb = load_workbook(wp)
    first_ws = wb.worksheets[0]
    first_a1 = str(first_ws.cell(row=1, column=1).value or "")
    if first_a1 == "#":
        ws = first_ws
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet("指标目录", 0)

    hf = PatternFill("solid", fgColor="4472C4")
    gf = PatternFill("solid", fgColor="5B9BD5")
    sf_p = PatternFill("solid", fgColor="D6E4F0")
    lf = PatternFill("solid", fgColor="F2F2F2")
    cf = Font(name="Arial", size=10)
    tb = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    ac = Alignment(horizontal="center", vertical="center")

    for i, w in enumerate([5, 28, 8, 8, 60, 12, 10]):
        ws.column_dimensions[get_column_letter(i + 1)].width = w

    headers = ["#", "Sheet Name", "Freq", "Col", "Indicator Name", "Unit", "板块"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
        c.fill = hf
        c.alignment = ac
        c.border = tb

    row = 2
    seq = 0
    for freq_label in ["日度", "周度", "月度", "季度", "年度"]:
        sf = [s for s in catalog if s["freq"] == freq_label]
        if not sf:
            continue
        ti = sum(s["cnt"] for s in sf)
        for ci in range(1, 8):
            c = ws.cell(row=row, column=ci)
            c.font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
            c.fill = gf
            c.border = tb
            if ci == 1:
                c.value = freq_label + " (" + str(len(sf)) + " sht " + str(ti) + " ind)"
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        row += 1

        for s in sf:
            sn = s["sn"]
            sec = classify_sector(sn)
            for ci in range(1, 8):
                c = ws.cell(row=row, column=ci)
                c.font = Font(name="Arial", bold=True, size=10)
                c.fill = sf_p
                c.border = tb
                if ci == 1:
                    c.value = sn
                    hl = Hyperlink(ref=c.coordinate, location="'" + sn + "'!A1")
                    c.hyperlink = hl
                    c.font = Font(name="Arial", bold=True, size=10, color="0563C1", underline="single")
                elif ci == 3:
                    c.value = s["freq"]
                elif ci == 7:
                    c.value = sec
                    c.font = Font(name="Arial", bold=True, size=10)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            row += 1

            for ind in s["inds"]:
                seq += 1
                vals = [seq, "", s["freq"], ind["col"], ind["name"], ind["unit"], sec]
                for ci, v in enumerate(vals, 1):
                    c = ws.cell(row=row, column=ci, value=v)
                    c.font = cf
                    c.border = tb
                    if ci == 3 or ci == 7:
                        c.alignment = ac
                    if ci == 5:
                        cl = get_column_letter(ind["col"])
                        hl = Hyperlink(ref=c.coordinate, location="'" + sn + "'!" + cl + "1")
                        c.hyperlink = hl
                        c.font = Font(name="Arial", size=10, color="0563C1", underline="single")
                    if seq % 2 == 0:
                        c.fill = lf
                row += 1

    sr = row + 1
    total_inds = sum(s["cnt"] for s in catalog)
    ws.cell(row=sr, column=1, value="[统计]").font = Font(name="Arial", bold=True, size=11)
    ws.cell(row=sr, column=2, value="Sheet 数: " + str(len(catalog))).font = Font(name="Arial", size=10)
    ws.cell(row=sr, column=4, value="指标数: " + str(total_inds)).font = Font(name="Arial", size=10)

    wb.save(wp)
    return len(catalog), total_inds


def classify_category(sn):
    if "成本" in sn or "利润" in sn:
        return "四、成本利润"
    if any(k in sn for k in ["进出口", "进口", "出口", "盈亏"]):
        return "三、供需-进出口"
    if any(k in sn for k in ["库存", "仓单"]):
        return "三、供需-库存"
    if any(k in sn for k in ["需求", "消费", "装机", "用电", "发电"]):
        return "二、供需-需求"
    if "平衡" in sn:
        return "三、供需-平衡"
    if any(k in sn for k in ["产量", "产能", "储量", "开工", "平衡", "排产", "预测", "加工费"]):
        return "二、供需-供给"
    if any(k in sn for k in ["成交", "持仓", "交易"]):
        return "五、量价"
    if any(k in sn for k in ["价格", "价差", "基差", "月差"]):
        return "一、价格"
    return "其他"


def build_tin_pickle(wp, out_path, catalog=None):
    if catalog is None:
        catalog = discover(wp)
    records = []
    for s in catalog:
        sn = s["sn"]
        sec = classify_sector(sn)
        cat = classify_category(sn)
        df = pd.read_excel(wp, sheet_name=sn, header=None, dtype=str)
        hdr_row, unit_row = find_header_rows(df)
        data_start = max(hdr_row + 2, unit_row + 1, 3)
        if data_start < df.shape[0]:
            fc = str(df.iloc[data_start, 0]).strip() if pd.notna(df.iloc[data_start, 0]) else ""
            if "频率" in fc:
                data_start += 1
        dates = pd.to_datetime(df.iloc[data_start:, 0], errors="coerce")
        for ind in s["inds"]:
            ci = ind["col"]
            values = pd.to_numeric(df.iloc[data_start:, ci], errors="coerce")
            mask = dates.notna() & values.notna()
            rows = []
            for d, v in zip(dates[mask], values[mask]):
                rows.append({"date": d.strftime("%Y-%m-%d"), "value": float(v)})
            rows.sort(key=lambda x: x["date"], reverse=True)
            unit = ind["unit"]
            div = 1.0
            if unit == "元/吨" and rows and max(x["value"] for x in rows) >= 1000:
                unit = "万元/吨"
                div = 10000.0
            if div != 1.0:
                rows = [{"date": x["date"], "value": round(x["value"] / div, 4)} for x in rows]
            fl = ind.get("freq") or detect_freq(sn)
            freq_en = _FM.get(fl[0], "monthly") if fl else "monthly"
            records.append(
                {
                    "id": "silicon_" + sn + "_c" + str(ci),
                    "title": ind["name"],
                    "unit": unit,
                    "category": cat,
                    "section": sn,
                    "source": "SMM/海关/行业",
                    "freq": freq_en,
                    "data": rows,
                    "sector": sec,
                    "commodity": "SILICON",
                }
            )
    with open(out_path, "wb") as f:
        pickle.dump(records, f)
    return len(records)
