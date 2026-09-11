# -*- coding: utf-8 -*-
"""AI-assisted curation for all industry rule modules.

Workflow:
1. Read the existing directory sheet produced by deterministic rules.
2. Keep high-confidence deterministic results unchanged.
3. Send only uncertain/low-confidence/unit-conflict indicators to DeepSeek in batches.
4. Parse structured JSON, apply unit consistency checks, cache results, and write back.
"""

from __future__ import annotations

import argparse
import concurrent.futures
import hashlib
import json
import os
import re
import shutil
import sqlite3
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


try:
    from dotenv import load_dotenv

    load_dotenv(Path(r"<local-documents>\Selecting skill\.env"))
except Exception:
    pass

from openai import OpenAI

# 统一 observability adapter（唯一 Langfuse 接触层，best-effort）。
# 直接加载 workflow/observability.py，避免触发 workflow/__init__（其 import cli）。
import importlib.util  # noqa: E402

_OBS_SPEC = importlib.util.spec_from_file_location(
    "workflow_observability",
    str(Path(__file__).resolve().parents[1] / "workflow" / "observability.py"),
)
observability = importlib.util.module_from_spec(_OBS_SPEC)
assert _OBS_SPEC and _OBS_SPEC.loader
_OBS_SPEC.loader.exec_module(observability)


PROMPT_VERSION = "ai_assisted_curation_v1"

# 系统提示模板（与调用解耦，供 observability 计算稳定 prompt_hash）
AI_DISAMBIGUATION_SYSTEM_PROMPT = (
    "你是行业指标分类助手。根据指标名称、所在Sheet、单位、频率和当前规则结果，"
    "只对明显不合理的项做修正。如果指标标签中包含“候选分类组合”，"
    "必须从该组合中选择一组合法的大类和子类；如果包含“候选大类”或“候选子类”，"
    "必须从对应候选中选择最合适的一项，不能自行另选。必须输出严格JSON："
    '{"results":[{"variable_id":"...","major":"...","sub":"...","nature":"...",'
    '"tags":{"产品":"...","研究主题":"...","指标类型":"...","规格":"...","工艺属性":"...",'
    '"地域":"...","统计口径":"...","状态":"...","频率":"...","单位":"..."},'
    '"review":false,"reason":"..."}]}'
)
PROMPT_HASH = None  # 懒计算：prompt_fingerprint(AI_DISAMBIGUATION_SYSTEM_PROMPT)
CRITICAL_TAG_FIELDS = ("产品", "研究主题", "指标类型")
DEFAULT_BATCH_SIZE = int(os.getenv("AI_ASSISTED_BATCH_SIZE", os.getenv("DEEPSEEK_BATCH_SIZE", "20")))
DEFAULT_MAX_CONCURRENCY = int(os.getenv("AI_ASSISTED_MAX_CONCURRENCY", os.getenv("DEEPSEEK_MAX_CONCURRENCY", "4")))
DEFAULT_MAX_TOKENS = int(os.getenv("AI_ASSISTED_MAX_TOKENS", os.getenv("DEEPSEEK_MAX_TOKENS", "32000")))
DEFAULT_TIMEOUT_SECONDS = float(os.getenv("AI_ASSISTED_TIMEOUT_SECONDS", os.getenv("DEEPSEEK_TIMEOUT_SECONDS", "120")))
DEFAULT_MAX_RETRIES = int(os.getenv("AI_ASSISTED_MAX_RETRIES", os.getenv("DEEPSEEK_MAX_RETRIES", "3")))

ALLOWED_MAJORS = {
    "价格",
    "成本利润",
    "进出口",
    "库存",
    "需求",
    "供给",
    "平衡",
    "其他",
    "供需-进出口",
    "供需-库存",
    "供需-需求",
    "供需-供给",
    "供需-进出口",
    "供需-平衡",
}

PRICE_UNITS = (
    "元/吨",
    "美元/吨",
    "元/千克",
    "元/片",
    "元/瓦",
    "元/kWh",
    "元/平方米",
    "元/套",
    "元/个",
)

QUANTITY_UNITS = (
    "吨",
    "万吨",
    "GW",
    "万千瓦",
    "亿千瓦时",
    "手",
    "台",
    "片",
    "瓦",
    "平方米",
    "千克",
)


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


def unit_conflict(major: str, sub: str, unit: str) -> str | None:
    unit = str(unit or "").strip()
    if not unit or unit in ("nan", "未指定"):
        return None
    if unit in PRICE_UNITS or unit.startswith(("元", "美元")):
        if major not in ("价格", "成本利润"):
            return f"价格/成本单位与大类不一致：{major}"
        return None
    if unit in QUANTITY_UNITS:
        if major == "价格":
            return f"数量单位与价格大类不一致：{unit}"
        return None
    if "%" in unit and major not in ("供需-供给", "供需-需求", "供需-库存", "成本利润", "其他"):
        return f"比率单位与大类不一致：{unit}"
    if "天" in unit and major not in ("供需-库存", "其他"):
        return f"天数单位与大类不一致：{unit}"
    return None


def is_indicator_row(values: list) -> bool:
    if not values or values[0] is None:
        return False
    raw = str(values[0]).strip()
    return bool(re.fullmatch(r"\d+", raw)) and bool(values[3]) and bool(values[4])


def load_rows(ws):
    rows = []
    current_sheet = ""
    current_sector = ""
    tag_col = None
    headers = [str(c.value or "").strip() for c in ws[1]]
    for idx, header in enumerate(headers, 1):
        if header == "指标标签":
            tag_col = idx
            break

    for row in ws.iter_rows(min_row=2):
        values = [c.value for c in row]
        col0 = str(values[0]).strip() if values[0] is not None else ""
        if col0.startswith("[统计]"):
            continue
        if not is_indicator_row(values):
            if col0 and not re.fullmatch(r"\d+", col0) and len(values) > 6 and values[6] is not None:
                current_sheet = col0
                current_sector = str(values[6]).strip()
            continue

        major = str(values[7] or "").strip()
        sub = str(values[8] or "").strip()
        unit = str(values[5] or "").strip()
        raw_tags = str(values[tag_col - 1]).strip() if tag_col and values[tag_col - 1] is not None else ""
        tags = {}
        if raw_tags:
            try:
                tags = json.loads(raw_tags)
            except Exception:
                tags = {}
        rows.append(
            {
                "sheet": current_sheet,
                "sector": str(values[6] or "").strip() or current_sector,
                "name": str(values[4]),
                "unit": unit,
                "freq": str(values[2] or "").strip(),
                "major": major,
                "sub": sub,
                "nature": str(values[9] or "").strip(),
                "status": str(values[12] or "").strip() if len(values) > 12 else "",
                "tags": tags,
                "before_major": major,
                "before_sub": sub,
                "before_nature": str(values[9] or "").strip(),
                "before_tags": dict(tags),
                "_row": row,
                "_tag_col": tag_col,
            }
        )
    return rows


def is_uncertain(rec: dict) -> tuple[bool, str]:
    if "确定性净出口计算" in rec.get("status", ""):
        return False, "确定性净出口计算已覆盖"
    if rec["major"] in ("", "其他") or rec["sub"] in ("", "其他"):
        return True, "大类或子类未识别"
    if any(rec["tags"].get(field) in ("未识别", "未指定") for field in CRITICAL_TAG_FIELDS):
        return True, "关键标签未识别"
    conflict = unit_conflict(rec["major"], rec["sub"], rec["unit"])
    if conflict:
        return True, conflict
    return False, ""


def build_item(rec: dict) -> dict:
    vid = hashlib.sha256(
        f"{rec['sheet']}|{rec['name']}|{rec['unit']}".encode("utf-8")
    ).hexdigest()[:16]
    return {
        "variable_id": vid,
        "name": rec["name"],
        "sheet": rec["sheet"],
        "unit": rec["unit"],
        "frequency": rec["freq"],
        "sector": rec["sector"],
        "current_major": rec["major"],
        "current_sub": rec["sub"],
        "current_tags": rec["tags"],
    }


def cache_key(item: dict, provider: str, model: str) -> str:
    payload = {
        "prompt_version": PROMPT_VERSION,
        "provider": provider,
        "model": model,
        "item": item,
    }
    return hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()


def init_cache(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(str(db_path))
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS ai_classification_cache (
            cache_key TEXT PRIMARY KEY,
            payload_json TEXT,
            provider TEXT,
            model TEXT,
            created_at TEXT
        )
        """
    )
    return con


def cache_get(con: sqlite3.Connection, key: str):
    cur = con.execute(
        "SELECT payload_json FROM ai_classification_cache WHERE cache_key=?",
        (key,),
    )
    row = cur.fetchone()
    if not row:
        return None
    payload = json.loads(row[0])
    return payload if payload else None


def cache_set(con: sqlite3.Connection, key: str, payload: dict, provider: str, model: str) -> None:
    con.execute(
        """
        INSERT OR REPLACE INTO ai_classification_cache
        (cache_key, payload_json, provider, model, created_at)
        VALUES (?, ?, ?, ?, ?)
        """,
        (key, json.dumps(payload, ensure_ascii=False), provider, model, utcnow()),
    )
    con.commit()


def _extract_results_from_content(content: str) -> list[dict]:
    """Parse DeepSeek JSON robustly, salvaging valid objects when output is truncated."""
    if not content:
        return []
    text = content.strip()
    # Remove markdown code fences if present.
    fence = re.search(r"```(?:json)?\s*(.*?)\s*```", text, re.S)
    if fence:
        text = fence.group(1).strip()

    def parse_full(candidate: str):
        try:
            payload = json.loads(candidate)
        except Exception:
            return None
        results = payload.get("results")
        return results if isinstance(results, list) else None

    results = parse_full(text)
    if results is not None:
        return results

    # Scan from the first '[' after "results". The scan ignores braces and quotes
    # inside JSON strings, so it can recover all complete objects before a
    # truncated trailing object.
    results = []
    array_start = text.find("[")
    if array_start < 0:
        return []
    i = array_start
    n = len(text)
    while i < n:
        if text[i] != "{":
            i += 1
            continue
        start = i
        depth = 0
        in_string = False
        escaped = False
        closed = False
        while i < n:
            ch = text[i]
            if in_string:
                if escaped:
                    escaped = False
                elif ch == "\\":
                    escaped = True
                elif ch == '"':
                    in_string = False
            else:
                if ch == '"':
                    in_string = True
                elif ch == "{":
                    depth += 1
                elif ch == "}":
                    depth -= 1
                    if depth == 0:
                        closed = True
                        i += 1
                        break
            i += 1
        if not closed:
            break
        candidate = text[start:i]
        try:
            obj = json.loads(candidate)
        except Exception:
            i = start + 1
            continue
        if isinstance(obj, dict):
            results.append(obj)
        i += 1
    return results


def call_deepseek_batch(
    items: list[dict],
    provider: str,
    model: str,
    *,
    timeout: float = DEFAULT_TIMEOUT_SECONDS,
    max_retries: int = DEFAULT_MAX_RETRIES,
    max_tokens: int = DEFAULT_MAX_TOKENS,
    span=None,
) -> list[dict]:
    api_key = os.getenv("DEEPSEEK_API_KEY", "")
    base_url = os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com")
    if not api_key:
        raise RuntimeError("DEEPSEEK_API_KEY is not set.")

    client = OpenAI(api_key=api_key, base_url=base_url, timeout=timeout, max_retries=0)
    user_payload = {
        "provider": provider,
        "model": model,
        "prompt_version": PROMPT_VERSION,
        "items": items,
    }
    messages = [
        {"role": "system", "content": AI_DISAMBIGUATION_SYSTEM_PROMPT},
        {"role": "user", "content": json.dumps(user_payload, ensure_ascii=False, default=str)},
    ]

    # prompt_hash：系统提示模板的稳定 SHA-256（不含 items，随模板版本变化）
    global PROMPT_HASH
    if PROMPT_HASH is None:
        PROMPT_HASH = observability.prompt_fingerprint(AI_DISAMBIGUATION_SYSTEM_PROMPT)

    last_error: Exception | None = None
    for attempt in range(1, max_retries + 1):
        try:
            response = observability.completion(
                lambda: client.chat.completions.create(
                    model=model,
                    messages=messages,
                    response_format={"type": "json_object"},
                    temperature=0,
                    max_tokens=max_tokens,
                ),
                model=model,
                messages=messages,
                prompt_name="ai_disambiguation",
                prompt_version=PROMPT_VERSION,
                prompt_hash=PROMPT_HASH,
                extra_metadata={"items_count": len(items)},
                span=span,
            )
            content = response.choices[0].message.content or "{}"
            results = _extract_results_from_content(content)
            if isinstance(results, list) and results:
                return results
            raise ValueError("response missing non-empty results list")
        except Exception as exc:
            last_error = exc
            if attempt < max_retries:
                time.sleep(min(30, 1.5 ** attempt))
    raise RuntimeError(f"DeepSeek batch failed: {last_error}")


def apply_ai_result(rec: dict, result: dict) -> None:
    raw_major = str(result.get("major") or "").strip()
    raw_sub = str(result.get("sub") or "").strip()
    if raw_major == "供应":
        raw_major = "供给"
    if raw_sub == "供应":
        raw_sub = "供给"
    current_tags = rec.get("before_tags") or rec.get("tags") or {}
    candidate_majors = current_tags.get("候选大类") or []
    candidate_subs = current_tags.get("候选子类") or []
    candidate_combos = current_tags.get("候选分类组合") or []

    def major_allowed(value: str) -> bool:
        if value in ALLOWED_MAJORS:
            return True
        if value == "价格-价格" or "价格" in value:
            return "价格" in (candidate_majors or []) or not candidate_majors
        return False

    selected_major = None
    if major_allowed(raw_major):
        if raw_major == "价格-价格" or ("价格" in raw_major and raw_major not in ALLOWED_MAJORS):
            selected_major = "价格"
        elif raw_major.startswith("供需-"):
            selected_major = raw_major
        else:
            selected_major = raw_major

    if candidate_combos:
        combo_matches = [
            c for c in candidate_combos
            if str(c.get("大类", "")) == selected_major and str(c.get("子类", "")) == raw_sub
        ]
        if not combo_matches:
            return
    elif candidate_majors and selected_major not in candidate_majors:
        return
    if candidate_subs and raw_sub and raw_sub not in candidate_subs:
        return

    if selected_major:
        rec["major"] = selected_major
    if raw_sub:
        rec["sub"] = raw_sub
    elif rec["major"] == "价格" and candidate_subs and len(candidate_subs) == 1:
        rec["sub"] = candidate_subs[0]
    rec["nature"] = str(result.get("nature") or rec["nature"])
    tags = result.get("tags")
    if isinstance(tags, dict):
        for key, value in list(tags.items()):
            if value == "供应":
                tags[key] = "供给"
        rec["tags"] = tags
    rec["_ai_reason"] = str(result.get("reason") or "AI辅助判断")
    rec["_ai_review"] = bool(result.get("review", False))


def tag_diff(before: dict, after: dict) -> dict:
    diffs: dict[str, dict] = {}
    for key in sorted(set(before) | set(after)):
        old = before.get(key)
        new = after.get(key)
        if old != new:
            diffs[key] = {"before": old, "after": new}
    return diffs


def has_multiple_candidates(rec: dict) -> bool:
    tags = rec.get("before_tags") or rec.get("tags") or {}
    return bool(
        len(tags.get("候选大类") or []) > 1
        or len(tags.get("候选子类") or []) > 1
        or len(tags.get("候选分类组合") or []) > 1
    )


def candidate_combo_key(rec: dict) -> tuple:
    combos = rec.get("candidate_combos") or []
    return tuple(
        sorted(
            (
                str(c.get("大类", "")),
                str(c.get("子类", "")),
            )
            for c in combos
        )
    )


def write_rows(ws, records: list[dict]) -> None:
    for rec in records:
        row = rec["_row"]
        row[7].value = rec["major"]
        row[8].value = rec["sub"]
        row[9].value = rec["nature"]
        tag_col = rec["_tag_col"]
        if tag_col:
            row[tag_col - 1].value = json.dumps(rec["tags"], ensure_ascii=False)
        conflict = unit_conflict(rec["major"], rec["sub"], rec["unit"])
        old_status = str(row[12].value or "").strip()
        note = rec.get("_ai_reason") or rec.get("_reason") or ""
        if conflict:
            note = f"{note}；单位冲突：{conflict}".strip("；")
        if note:
            if old_status and note not in old_status:
                row[12].value = f"{old_status}；{note}"
            else:
                row[12].value = note


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="AI-assisted curation for industry directories.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--industry", choices=["lithium", "silicon", "tin"], required=True)
    parser.add_argument("--provider", default="deepseek")
    parser.add_argument("--model", default=None)
    parser.add_argument("--batch-size", type=int, default=DEFAULT_BATCH_SIZE)
    parser.add_argument("--max-concurrency", type=int, default=DEFAULT_MAX_CONCURRENCY)
    parser.add_argument("--max-tokens", type=int, default=DEFAULT_MAX_TOKENS)
    parser.add_argument("--timeout", type=float, default=DEFAULT_TIMEOUT_SECONDS)
    parser.add_argument("--max-retries", type=int, default=DEFAULT_MAX_RETRIES)
    parser.add_argument("--database", default=None)
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--output", default=None)
    parser.add_argument("--report", default=None)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--force-refresh", action="store_true")
    args = parser.parse_args(argv)

    input_path = Path(args.input).resolve()
    output_path = Path(args.output).resolve() if args.output else input_path.with_name(f"{input_path.stem}_ai.xlsx")
    db_path = (
        Path(args.database).resolve()
        if args.database
        else Path(__file__).resolve().parents[1] / "output" / "ai_curation.db"
    )
    report_path = Path(args.report).resolve() if args.report else output_path.with_suffix(".md").with_name(f"{output_path.stem}_report.md")
    model = args.model or os.getenv("DEEPSEEK_MODEL", "deepseek-v4-flash")

    wb = load_workbook(input_path)
    ws = wb.worksheets[0]
    records = load_rows(ws)
    for rec in records:
        uncertain, reason = is_uncertain(rec)
        rec["_uncertain"] = uncertain
        rec["_reason"] = reason if uncertain else ""

    uncertain = [r for r in records if r["_uncertain"]]
    if args.limit is not None:
        uncertain = uncertain[: args.limit]

    if args.dry_run:
        print(
            json.dumps(
                {
                    "industry": args.industry,
                    "total": len(records),
                    "uncertain": len(uncertain),
                    "batch_size": args.batch_size,
                    "estimated_batches": (len(uncertain) + args.batch_size - 1) // args.batch_size
                    if uncertain
                    else 0,
                    "max_concurrency": args.max_concurrency,
                    "max_tokens": args.max_tokens,
                    "estimated_tokens": len(uncertain) * 220,
                    "dry_run": True,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
        return 0

    con = init_cache(db_path)
    provider = args.provider
    ai_results: dict[str, dict] = {}
    cache_hits = 0
    api_calls = 0

    stage_stats = observability.StageStats()
    stage_stats.items_total = len(uncertain)
    with observability.stage_span("ai_disambiguation") as _span:
        pending_batches = []
        for start in range(0, len(uncertain), args.batch_size):
            batch = uncertain[start : start + args.batch_size]
            items = [build_item(rec) for rec in batch]
            keys = [cache_key(item, provider, model) for item in items]
            missing_items = []
            missing_keys = []
            missing_recs = []
            for rec, item, key in zip(batch, items, keys):
                cached = None if args.force_refresh else cache_get(con, key)
                if cached is not None:
                    cache_hits += 1
                    stage_stats.hit()
                    ai_results[item["variable_id"]] = cached
                else:
                    missing_items.append(item)
                    missing_keys.append(key)
                    missing_recs.append(rec)

            if missing_items:
                pending_batches.append((missing_items, missing_keys, missing_recs))

        max_workers = max(1, min(args.max_concurrency, len(pending_batches)))
        if pending_batches:
            print(f"  AI pending batches: {len(pending_batches)}")
            with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
                futures = {
                    executor.submit(
                        call_deepseek_batch,
                        missing_items,
                        provider,
                        model,
                        timeout=args.timeout,
                        max_retries=args.max_retries,
                        max_tokens=args.max_tokens,
                        span=_span,
                    ): (missing_items, missing_keys, missing_recs)
                    for missing_items, missing_keys, missing_recs in pending_batches
                }
                completed = 0
                for future in concurrent.futures.as_completed(futures):
                    missing_items, missing_keys, missing_recs = futures[future]
                    try:
                        results = future.result()
                        api_calls += 1
                        stage_stats.call()
                    except Exception as exc:
                        print(f"  AI batch failed: {exc}", file=sys.stderr)
                        results = []
                    completed += 1
                    print(f"  AI batch {completed}/{len(pending_batches)}: {len(results)} results")
                    by_id = {str(r.get("variable_id")): r for r in results}
                    for item, key, rec in zip(missing_items, missing_keys, missing_recs):
                        result = by_id.get(item["variable_id"], {})
                        if result:
                            cache_set(con, key, result, provider, model)
                        ai_results[item["variable_id"]] = result
        else:
            print("  AI pending batches: 0 (all cached)")
        if _span is not None:
            try:
                _span.update(metadata=stage_stats.to_metadata())
            except Exception:
                pass

    applied = 0
    conflicts = 0
    changed_records: list[dict] = []
    human_review_records: list[dict] = []
    for rec in records:
        if not rec["_uncertain"]:
            continue
        item = build_item(rec)
        result = ai_results.get(item["variable_id"])
        before_review = (
            rec["major"],
            rec["sub"],
            rec["nature"],
            dict(rec["before_tags"]),
        )
        if result:
            before_snapshot = (
                rec["major"],
                rec["sub"],
                rec["nature"],
                json.dumps(rec["tags"], ensure_ascii=False, sort_keys=True),
            )
            apply_ai_result(rec, result)
            applied += 1
            after_snapshot = (
                rec["major"],
                rec["sub"],
                rec["nature"],
                json.dumps(rec["tags"], ensure_ascii=False, sort_keys=True),
            )
            if before_snapshot != after_snapshot:
                changed_records.append(
                    {
                        "sheet": rec["sheet"],
                        "name": rec["name"],
                        "unit": rec["unit"],
                        "freq": rec["freq"],
                        "before_major": before_snapshot[0],
                        "before_sub": before_snapshot[1],
                        "before_nature": before_snapshot[2],
                        "before_tags": rec["before_tags"],
                        "after_major": after_snapshot[0],
                        "after_sub": after_snapshot[1],
                        "after_nature": after_snapshot[2],
                        "after_tags": rec["tags"],
                        "reason": rec.get("_ai_reason", ""),
                    }
                )
        if has_multiple_candidates(rec):
            before_tags = rec.get("before_tags") or {}
            human_review_records.append(
                {
                    "sheet": rec["sheet"],
                    "name": rec["name"],
                    "unit": rec["unit"],
                    "freq": rec["freq"],
                    "candidate_majors": before_tags.get("候选大类") or [],
                    "candidate_subs": before_tags.get("候选子类") or [],
                    "candidate_combos": before_tags.get("候选分类组合") or [],
                    "before_major": before_review[0],
                    "before_sub": before_review[1],
                    "before_nature": before_review[2],
                    "after_major": rec["major"],
                    "after_sub": rec["sub"],
                    "after_nature": rec["nature"],
                    "ai_reason": rec.get("_ai_reason") or rec.get("_reason") or "",
                }
            )
        if unit_conflict(rec["major"], rec["sub"], rec["unit"]):
            conflicts += 1

    write_rows(ws, records)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    con.close()

    report_lines = [
        "# AI 辅助判断审计报告",
        "",
        f"- 产业：{args.industry}",
        f"- 输入：`{input_path.name}`",
        f"- 输出：`{output_path.name}`",
        f"- 指标总数：{len(records)}",
        f"- 触发 AI 数量：{len(uncertain)}",
        f"- 实际应用 AI 结果：{applied}",
        f"- AI 实际修改指标数：{len(changed_records)}",
        f"- 人工复核清单数量：{len(human_review_records)}",
        f"- 缓存命中：{cache_hits}",
        f"- API 调用批次：{api_calls}",
        f"- 单位冲突数：{conflicts}",
        f"- Provider：{provider} / {model}",
        "",
        "## 触发原因统计",
        "",
    ]
    reason_counter: dict[str, int] = {}
    for rec in uncertain:
        reason = rec.get("_reason") or "未知"
        reason_counter[reason] = reason_counter.get(reason, 0) + 1
    for reason, count in sorted(reason_counter.items(), key=lambda x: -x[1]):
        report_lines.append(f"- {reason}: {count}")
    report_lines += ["", "## AI 实际修改明细", ""]
    report_lines.append("| # | Sheet | 指标名称 | 单位 | 修改前 | 修改后 | 标签变更 | AI原因 |")
    report_lines.append("|---|---|---|---|---|---|---|---|")
    for idx, rec in enumerate(changed_records, 1):
        before_cls = f"{rec['before_major']}/{rec['before_sub']}/{rec['before_nature']}"
        after_cls = f"{rec['after_major']}/{rec['after_sub']}/{rec['after_nature']}"
        diff = json.dumps(
            tag_diff(rec["before_tags"], rec["after_tags"]),
            ensure_ascii=False,
        )
        report_lines.append(
            f"| {idx} | {rec['sheet']} | {rec['name']} | {rec['unit']} "
            f"| {before_cls} | {after_cls} | {diff} | {rec['reason']} |"
        )

    grouped_human_review: dict[tuple, list[dict]] = {}
    for rec in human_review_records:
        key = candidate_combo_key(rec)
        grouped_human_review.setdefault(key, []).append(rec)
    grouped_human_review_items = sorted(
        grouped_human_review.items(),
        key=lambda item: (-len(item[1]), item[0]),
    )

    report_lines += ["", "## 人工二次复核清单", ""]
    report_lines.append(f"- 多候选指标总数：{len(human_review_records)}")
    report_lines.append(f"- 候选组合分组数：{len(grouped_human_review_items)}")
    for group_idx, (combos_key, recs) in enumerate(grouped_human_review_items, 1):
        recs = sorted(recs, key=lambda x: (x["sheet"], x["name"]))
        combo_text = "、".join(f"{major}/{sub}" for major, sub in combos_key) or "未形成候选组合"
        outcome_counter: dict[str, int] = {}
        for rec in recs:
            after_cls = f"{rec['after_major']}/{rec['after_sub']}/{rec['after_nature']}"
            outcome_counter[after_cls] = outcome_counter.get(after_cls, 0) + 1

        report_lines += ["", f"### 组 {group_idx}：{combo_text}", ""]
        report_lines.append(f"- 指标数量：{len(recs)}")
        report_lines.append(f"- 候选组合：`{combo_text}`")
        report_lines.append("- AI 结果分布：")
        for after_cls, count in sorted(outcome_counter.items(), key=lambda x: -x[1]):
            report_lines.append(f"  - {after_cls}：{count}")

        report_lines.append("")
        report_lines.append(
            f"<details><summary>展开 {len(recs)} 条指标明细</summary>"
        )
        report_lines.append("")
        report_lines.append("| Sheet | 指标名称 | 单位 | 频率 | 确定性结果 | AI结果 | AI原因 | 人工复核结论 |")
        report_lines.append("|---|---|---|---|---|---|---|---|")
        for rec in recs:
            before_cls = f"{rec['before_major']}/{rec['before_sub']}/{rec['before_nature']}"
            after_cls = f"{rec['after_major']}/{rec['after_sub']}/{rec['after_nature']}"
            report_lines.append(
                f"| {rec['sheet']} | {rec['name']} | {rec['unit']} | {rec['freq']} "
                f"| {before_cls} | {after_cls} | {rec['ai_reason']} |  |"
            )
        report_lines.append("")
        report_lines.append("</details>")

    report_path.write_text("\n".join(report_lines), encoding="utf-8")

    print(f"Total={len(records)} uncertain={len(uncertain)} applied={applied} changed={len(changed_records)} cache_hits={cache_hits} api_calls={api_calls} conflicts={conflicts}")
    print(f"Output: {output_path}")
    print(f"Report: {report_path}")
    observability.flush()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
