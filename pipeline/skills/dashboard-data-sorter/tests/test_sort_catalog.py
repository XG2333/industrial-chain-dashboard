# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

from sort_catalog import (  # noqa: E402
    CHINA_REGIONS,
    FOREIGN_REGIONS,
    _business_group_reason,
    _product_family_key,
    _region_key,
    _source_detail,
    build_output,
    classify_rows,
    data_sort_key,
    normalize_indicator_title,
    sort_catalog_workbook,
)


HEADER = [
    "#",
    "Sheet Name",
    "Freq",
    "Col",
    "Indicator Name",
    "Unit",
    "板块",
    "大类",
    "子类",
    "数据性质",
    "是否选中",
    "状态说明",
    "指标标签",
]


def make_row(
    num=1,
    freq="日度",
    col=1,
    name="测试指标",
    sector="",
    major="价格",
    sub="现货价格",
    nature="水平值",
    selected="是",
    tags="",
):
    return [
        num,
        "",
        freq,
        col,
        name,
        "",
        sector,
        major,
        sub,
        nature,
        selected,
        "",
        tags,
    ]


def write_workbook(path: Path, rows: list[list]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "指标目录"
    for row in rows:
        ws.append(row)
    wb.create_sheet("数据")
    wb.save(path)
    return path


def read_catalog(path: Path) -> list[list]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["指标目录"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def data_names(catalog: list[list]) -> list[str]:
    return [row[4] for row in catalog if row[4] not in (None, "", "Indicator Name")]


def test_normalize_matches_chart_grouping() -> None:
    assert normalize_indicator_title("GFEX: 工业硅: 主力合约: 成交量: 日度") == "工业硅主力合约"
    assert normalize_indicator_title("GFEX: 工业硅: 主力合约: 持仓量: 日度") == "工业硅主力合约"
    assert normalize_indicator_title("工业硅主力合约成交持仓比") == "工业硅主力合约"
    assert normalize_indicator_title("SMM: 电池级碳酸锂 - 平均价: 日度") == "碳酸锂"
    assert normalize_indicator_title("SMM: 木片 - 平均价: 周度") == "木片"


def test_selected_then_unselected_within_sheet_and_stats_at_bottom(tmp_path) -> None:
    rows = [
        HEADER,
        ["日度 (1 sht 1 ind)", "", "", "", "", "", "", "其他", "其他", "", "否", "", ""],
        ["B-sheet", "", "日度", "", "", "", "板块B", "其他", "其他", "", "否", "", ""],
        make_row(1, name="B 成交量", major="价格", sub="成交量", selected="是"),
        make_row(2, name="B 持仓", major="价格", sub="持仓", selected="是"),
        ["A-sheet", "", "日度", "", "", "", "板块A", "其他", "其他", "", "否", "", ""],
        make_row(1, name="A 已选", major="价格", sub="现货价格", selected="是"),
        make_row(2, name="A 未选", major="价格", sub="现货价格", selected="否"),
        ["[统计]", "", "", "", "'Sheet 数: 9'!A1", "", "", "其他", "其他", "", "否", "", ""],
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    catalog = read_catalog(output)
    names = [row[4] for row in catalog]
    # 已选优先是全局的（所有选中行按 data_sort_key 排在未选中行之前）：
    # A 已选（现货价格）与 B 成交量/持仓（均选中）按 key 相邻，A 未选最后
    assert names == [
        "Indicator Name",
        None,  # A-sheet 表头
        "A 已选",
        None,  # B-sheet 表头
        "B 成交量",
        "B 持仓",
        "A 未选",
        None,  # 统计行
        "'Sheet 数: 9'!A1",
    ]


def test_frequency_order_ignores_missing_half_year(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, freq="年度", name="Z 年度"),
        make_row(2, freq="月度", name="M 月度"),
        make_row(3, freq="日度", name="D 日度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    catalog = read_catalog(output)
    names = data_names(catalog)
    assert names == ["D 日度", "M 月度", "Z 年度"]


def test_china_first_then_current_predicted(tmp_path) -> None:
    def tagged(state: str, region: str) -> str:
        return json.dumps({"状态": state, "地域": region}, ensure_ascii=False)

    rows = [
        HEADER,
        ["R-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, name="平均价", tags=tagged("实际", "其他")),
        make_row(2, name="平均价", tags=tagged("预测", "中国")),
        make_row(3, name="平均价", tags=tagged("实际", "中国")),
        make_row(4, name="平均价", tags=tagged("预测", "其他")),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    catalog = read_catalog(output)
    tags = [row[12] for row in catalog if row[4] == "平均价"]
    # 地区顺序优先于当期/预测：无地区（其他）先于中国，地区内 当期→预测。
    assert [json.loads(t)["地域"] for t in tags] == ["其他", "其他", "中国", "中国"]
    assert [json.loads(t)["状态"] for t in tags] == ["实际", "预测", "实际", "预测"]


def test_import_export_rows_stay_separate(tmp_path) -> None:
    rows = [
        HEADER,
        ["I-sheet", "", "月度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, name="出口量", major="进出口", sub="出口", selected="是"),
        make_row(2, name="进口量", major="进出口", sub="进口", selected="是"),
        make_row(3, name="净出口", major="进出口", sub="净出口", selected="是"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    catalog = read_catalog(output)
    names = data_names(catalog)
    subs = [row[8] for row in catalog if row[4] not in (None, "", "Indicator Name")]
    assert names == ["净出口", "进口量", "出口量"]
    assert subs == ["净出口", "进口", "出口"]


def test_generic_spread_before_month_spread(tmp_path) -> None:
    rows = [
        HEADER,
        ["P-sheet", "", "日度", "", "", "", "", "价格", "其他", "", "否", "", ""],
        make_row(1, name="Z 价差指标", sub="价差"),
        make_row(2, name="A 月差指标", sub="月差"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    catalog = read_catalog(output)
    subs = [row[8] for row in catalog if row[4] not in (None, "", "Indicator Name")]
    assert subs == ["价差", "月差"]


def test_supports_directory_marked_with_confidence_column(tmp_path) -> None:
    header = HEADER[:10] + ["置信度"] + HEADER[10:]
    rows = [
        header,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "", "否", "", ""],
        [
            1,
            "",
            "日度",
            1,
            "B 指标",
            "",
            "",
            "价格",
            "现货价格",
            "水平值",
            0.2,
            "否",
            "",
            "",
        ],
        [
            2,
            "",
            "日度",
            2,
            "A 指标",
            "",
            "",
            "价格",
            "现货价格",
            "水平值",
            0.9,
            "是",
            "",
            json.dumps({"状态": "实际", "地域": "中国"}, ensure_ascii=False),
        ],
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    catalog = read_catalog(output)
    names = data_names(catalog)
    assert names == ["A 指标", "B 指标"]
    selected = [row[11] for row in catalog if row[4] == "A 指标"]
    assert selected == ["是"]


def test_restores_sheet_jump_hyperlinks(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, col=1, name="B 指标"),
        make_row(2, col=2, name="A 指标"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    workbook = load_workbook(output, data_only=False)
    ws = workbook["指标目录"]
    links = {}
    sheet_header_link = None
    for row in ws.iter_rows(min_row=1):
        if row[0].value == "S-sheet":
            sheet_header_link = row[0].hyperlink.location if row[0].hyperlink else None
        name = row[4].value
        if name in ("A 指标", "B 指标"):
            links[name] = row[4].hyperlink.location if row[4].hyperlink else None
    assert sheet_header_link == "'S-sheet'!A1"
    assert links["A 指标"] == "'S-sheet'!B1"
    assert links["B 指标"] == "'S-sheet'!A1"
    workbook.close()


def test_preserves_existing_hyperlinks(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, col=1, name="Only 指标"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    workbook = load_workbook(source, data_only=False)
    ws = workbook["指标目录"]
    ws["E3"].hyperlink = "https://example.com/jump"
    workbook.save(source)
    workbook.close()

    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    workbook = load_workbook(output, data_only=False)
    ws = workbook["指标目录"]
    links = [
        row[4].hyperlink.target
        for row in ws.iter_rows(min_row=1)
        if row[4].value == "Only 指标"
    ]
    assert links == ["https://example.com/jump"]
    workbook.close()


def test_china_before_foreign_within_subclass(tmp_path) -> None:
    def tagged(region: str) -> str:
        return json.dumps({"地域": region}, ensure_ascii=False)

    rows = [
        HEADER,
        ["S-sheet", "", "月度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, name="测试进口量", major="进出口", sub="进口", tags=tagged("日本")),
        make_row(2, name="测试进口量", major="进出口", sub="进口", tags=tagged("中国")),
        make_row(3, name="测试进口量", major="进出口", sub="进口", tags=tagged("全球")),
        make_row(4, name="测试进口量", major="进出口", sub="进口", tags=tagged("未指定")),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    catalog = read_catalog(output)
    regions = [
        json.loads(row[12])["地域"]
        for row in catalog
        if row[4] == "测试进口量"
    ]
    # 地区顺序：无地区（未指定）排前，其次中国/省份，全球/海外/外国最后。
    assert regions[0] == "未指定"
    assert regions[1] == "中国"
    assert set(regions) == {"中国", "未指定", "全球", "日本"}


def test_sector_node_order_applied_as_final_rule(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, sector="磷酸铁锂", name="SMM: 磷酸铁锂 - 平均价: 日度"),
        make_row(2, sector="磷酸铁锂", name="SMM: 磷酸铁 - 平均价: 日度"),
        make_row(3, sector="磷酸铁锂", name="SMM: 磷酸一铵 - 平均价: 日度"),
        make_row(4, sector="磷酸铁锂", name="SMM: 无法识别指标 - 平均价: 日度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    catalog = read_catalog(output)
    names = data_names(catalog)
    assert names == [
        "SMM: 磷酸一铵 - 平均价: 日度",
        "SMM: 磷酸铁 - 平均价: 日度",
        "SMM: 磷酸铁锂 - 平均价: 日度",
        "SMM: 无法识别指标 - 平均价: 日度",
    ]


def test_parenthesized_china_does_not_override_foreign_source(tmp_path) -> None:
    assert _region_key({}, "SMM: 澳大利亚锂辉石精矿（CIF中国）现货 - 平均价: 日度") == 2
    assert _region_key({}, "SMM: 锂辉石（中国现货 Li2O: 3%-4%） - 平均价: 日度") == 0
    assert _region_key({}, "SMM: 磷锂铝石（中国现货）（Li₂O: 6%-7%） - 平均价: 日度") == 2


def test_parallel_graphite_branches_keep_display_order(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, sector="负极材料", name="SMM: 人造石墨 - 平均价: 日度"),
        make_row(2, sector="负极材料", name="SMM: 天然石墨 - 平均价: 日度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    catalog = read_catalog(output)
    names = data_names(catalog)
    assert names == ["SMM: 天然石墨 - 平均价: 日度", "SMM: 人造石墨 - 平均价: 日度"]


def test_metric_policy_and_peer_reason(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "月度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, sector="负极材料", major="供给", sub="产量", name="负极材料产量"),
        make_row(2, sector="负极材料", major="成本利润", sub="成本", name="负极材料成本"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output, add_sort_reason_column=True)

    workbook = load_workbook(output, read_only=True, data_only=True)
    ws = workbook["指标目录"]
    rows_out = list(ws.iter_rows(values_only=True))
    header = [str(v or "") for v in rows_out[0]]
    assert "指标名称归一化" in header
    assert "排序说明" in header
    norm_idx = header.index("指标名称归一化")
    reason_idx = header.index("排序说明")
    assert reason_idx == norm_idx + 1
    reasons = [r[reason_idx] for r in rows_out[1:] if r[4] not in (None, "", "Indicator Name")]
    assert any("产量/产能" in str(r) for r in reasons)
    assert any("同级" in str(r) for r in reasons)
    workbook.close()


def test_sort_output_is_stable(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, sector="磷酸铁锂", name="SMM: 磷酸铁锂 - 平均价: 日度"),
        make_row(2, sector="磷酸铁锂", name="SMM: 磷酸铁 - 平均价: 日度"),
        make_row(3, sector="磷酸铁锂", name="SMM: 磷酸一铵 - 平均价: 日度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    sort_catalog_workbook(source, first)
    sort_catalog_workbook(source, second)
    assert data_names(read_catalog(first)) == data_names(read_catalog(second))


def test_futures_policy_sorts_main_contract_before_month_contracts(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, major="价格", sub="期货价格", name="SHFE: 锡: 二月合约: 收盘价: 日度"),
        make_row(2, major="价格", sub="期货价格", name="SHFE: 锡: 一月合约: 收盘价: 日度"),
        make_row(3, major="价格", sub="期货价格", name="SHFE: 锡: 主力合约: 收盘价: 日度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    catalog = read_catalog(output)
    names = data_names(catalog)
    assert names == [
        "SHFE: 锡: 主力合约: 收盘价: 日度",
        "SHFE: 锡: 一月合约: 收盘价: 日度",
        "SHFE: 锡: 二月合约: 收盘价: 日度",
    ]


def test_region_words_sorted_with_lexical_tie_break() -> None:
    # 地区词表排序必须带字典序 tie-break（先长度降序、同长字典序降序），
    # 否则同长词顺序依赖 set/hash（PYTHONHASHSEED），导致排序结果运行间不稳定
    words = sorted(
        set(CHINA_REGIONS) | set(FOREIGN_REGIONS) | {"全球", "海外", "中国"},
        key=lambda w: (len(w), w),
        reverse=True,
    )
    for i in range(len(words) - 1):
        if len(words[i]) == len(words[i + 1]):
            assert words[i] > words[i + 1], f"同长词未按字典序 tie-break: {words[i]} vs {words[i+1]}"


def test_source_detail_same_length_region_deterministic() -> None:
    # "巴西"与"中国"同长：_source_detail 必须稳定返回同一词（修复前随 hash seed 变化）
    name = "SMM: 巴西锂辉石精矿（CIF中国）现货 - 平均价: 日度"
    results = {_source_detail(name) for _ in range(20)}
    assert len(results) == 1, results
    (result,) = results
    words = sorted(
        set(CHINA_REGIONS) | set(FOREIGN_REGIONS) | {"全球", "海外", "中国"},
        key=lambda w: (len(w), w),
        reverse=True,
    )
    first = next(w for w in words if w in name)
    assert result == first


def test_spread_policy_sorts_basis_before_month_spread(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, major="价格", sub="月差", name="SMM: 沪锡月差"),
        make_row(2, major="价格", sub="基差", name="SMM: 沪锡基差"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    catalog = read_catalog(output)
    names = data_names(catalog)
    assert names == ["SMM: 沪锡基差", "SMM: 沪锡月差"]


def _key_rec(major: str, sub: str, policy: str, node_priority: int, name: str, row_idx: int, sector: str = "") -> dict:
    return {
        "name": name,
        "major": major,
        "sub": sub,
        "sector": sector,
        "freq": "日度",
        "nature": "",
        "tags": {},
        "industry": {
            "policy": policy,
            "matched": True,
            "fallback": False,
            "conflict": False,
            "node_priority": node_priority,
            "branch_priority": 0,
        },
        "row_idx": row_idx,
    }


def test_subcategory_priority_not_overridden_by_industry_stage() -> None:
    spot = _key_rec("价格", "现货价格", "spot_price", 999, "现货A", 1)
    futures = _key_rec("价格", "期货价格", "futures_price", 0, "期货主力", 2)
    assert data_sort_key(spot) < data_sort_key(futures)


def test_full_price_subcategory_order() -> None:
    subs = ["价格", "现货价格", "期货价格", "现货价差", "基差", "月差"]
    keys = [
        data_sort_key(_key_rec("价格", sub, "default", 999 - index, f"指标{index}", index))
        for index, sub in enumerate(subs)
    ]
    assert keys == sorted(keys)


def test_same_rank_sectors_grouped_by_sub_rank() -> None:
    # 同 rank 板块必须有次级键（SECTOR_SUB_RANK）：磷酸铁锂(2,0) 与 磷化工链(2,1)，
    # 所有磷酸铁锂段（无论大类）都排在磷化工链段前，避免两板块段按大类交错
    lfp_price = _key_rec("价格", "现货价格", "spot_price", 999, "磷酸铁锂价格", 1, sector="磷酸铁锂")
    lfp_inventory = _key_rec("库存", "库存", "inventory", 0, "磷酸铁锂库存", 2, sector="磷酸铁锂")
    phos_price = _key_rec("价格", "现货价格", "spot_price", 999, "磷化工价格", 3, sector="磷化工链")
    assert data_sort_key(lfp_price) < data_sort_key(phos_price)
    assert data_sort_key(lfp_inventory) < data_sort_key(phos_price)
    # 碳酸锂(1,0) < 氢氧化锂(1,1) < 其他锂盐(1,2) 不变
    carbonate = _key_rec("价格", "现货价格", "spot_price", 999, "碳酸锂价格", 4, sector="碳酸锂")
    hydroxide = _key_rec("价格", "现货价格", "spot_price", 999, "氢氧化锂价格", 5, sector="氢氧化锂")
    other_salt = _key_rec("价格", "现货价格", "spot_price", 999, "其他锂盐价格", 6, sector="其他锂盐")
    assert data_sort_key(carbonate) < data_sort_key(hydroxide) < data_sort_key(other_salt)


def test_sector_priority_overrides_major() -> None:
    # 排序一级为板块（SECTOR_RANK：锂矿→锂盐→磷酸铁锂→三元正极→…），
    # 大类（major）不再独立成层：锂矿板块的价格指标排在电池电芯板块的库存指标之前
    price_lithium = _key_rec("价格", "现货价格", "spot_price", 999, "价格指标", 1, sector="锂矿")
    inventory_cell = _key_rec("库存", "库存", "inventory", 0, "库存指标", 2, sector="电池电芯")
    assert data_sort_key(price_lithium) < data_sort_key(inventory_cell)
    # 同板块内仍按子类顺序（现货价格在期货价格之前）
    spot = _key_rec("价格", "现货价格", "spot_price", 999, "现货A", 1, sector="锂矿")
    futures = _key_rec("价格", "期货价格", "futures_price", 0, "期货主力", 2, sector="锂矿")
    assert data_sort_key(spot) < data_sort_key(futures)


def test_incomplete_trade_rows_do_not_lead_sections(tmp_path) -> None:
    # 未配对（INCOMPLETE）的进出口行不参与复合组排序（role=9999），
    # 不再以 role=0 插队到板块最前：锂矿价格段应排在锂矿进出口段之前
    incomplete_tags = json.dumps(
        {
            "composite": {
                "composite_type": "trade",
                "composite_key": "中国海关|锂精矿月度|总计|月度|吨|实际",
                "composite_role": "import",
                "composite_status": "INCOMPLETE",
            }
        },
        ensure_ascii=False,
    )
    rows = [
        HEADER,
        ["S-锂矿价格", "", "日度", "", "", "", "", "锂矿", "价格", "现货价格", "", "是", "", ""],
        make_row(
            1,
            name="SMM: 澳大利亚锂辉石精矿（CIF中国）现货 - 平均价: 日度",
            sector="锂矿", major="价格", sub="现货价格", freq="日度",
        ),
        ["S-锂矿进出口", "", "月度", "", "", "", "", "锂矿", "进出口", "进口", "", "是", "", ""],
        make_row(
            2,
            name="中国海关: 锂精矿月度进口量: 总计: 月度",
            sector="锂矿", major="进出口", sub="进口", freq="月度",
            tags=incomplete_tags,
        ),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    catalog = read_catalog(output)
    names = data_names(catalog)
    assert names[0] == "SMM: 澳大利亚锂辉石精矿（CIF中国）现货 - 平均价: 日度"
    assert names[1] == "中国海关: 锂精矿月度进口量: 总计: 月度"


def test_policy_only_applies_inside_same_subcategory_bucket() -> None:
    spot = _key_rec("价格", "现货价格", "spot_price", 999, "现货A", 1)
    futures = _key_rec("价格", "期货价格", "futures_price", 0, "期货主力", 2)
    assert data_sort_key(spot) < data_sort_key(futures)
    assert data_sort_key(futures)[4] > data_sort_key(spot)[4]


def test_major_order_within_sector() -> None:
    # 板块内大类顺序为硬顺序（MAJOR_ORDER）：价格 -> 成本利润 -> 库存 -> 供给
    # -> 需求 -> 进出口 -> 平衡；大类内再按子类顺序
    price = _key_rec("价格", "现货价格", "spot_price", 999, "锂辉石精矿现货", 1, sector="锂矿")
    inventory = _key_rec("库存", "库存", "inventory", 0, "锂矿港口库存", 2, sector="锂矿")
    trade = _key_rec("进出口", "进口", "trade", 0, "锂精矿月度进口量", 3, sector="锂矿")
    assert data_sort_key(price) < data_sort_key(inventory)
    assert data_sort_key(inventory) < data_sort_key(trade)
    # 大类层在复合组层之前：同板块不同大类的顺序不受 composite role 影响
    assert data_sort_key(price)[1] < data_sort_key(inventory)[1] < data_sort_key(trade)[1]


def test_sheet_segments_respect_price_subcategory_order(tmp_path) -> None:
    rows = [
        HEADER,
        ["期货+月差 Sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, major="价格", sub="期货价格", name="期货指标"),
        make_row(2, major="价格", sub="月差", name="月差指标"),
        ["现货 Sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, major="价格", sub="现货价格", name="现货指标"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    catalog = read_catalog(output)
    sheet_headers = [
        row[0]
        for row in catalog
        if isinstance(row[0], str) and row[0] not in ("#", "Sheet Name")
    ]
    assert sheet_headers[0] == "现货 Sheet"
    assert sheet_headers[1] == "期货+月差 Sheet"


def test_self_loop_explanation_is_suppressed() -> None:
    rec = {
        "name": "磷酸铁",
        "industry": {
            "matched": True,
            "conflict": False,
            "policy": "spot_price",
            "node_id": "iron_phosphate",
            "display_name": "磷酸铁",
            "stage": "midstream",
            "relation_type": "strict",
        },
    }
    group = {"recs": [rec]}
    reason = _business_group_reason(group, group, None)
    assert "之间" not in reason
    assert "上游承接" not in reason
    assert "同一产业环节" in reason


def test_metal_lithium_not_labeled_as_lithium_salt(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, sector="其他锂盐", name="SMM: 金属锂 - 平均价: 日度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output, add_sort_reason_column=True)
    reason = _read_first_reason(output)
    assert "金属锂" in reason
    assert "锂盐基础产品" not in reason


def test_cmc_sbr_nmp_not_labeled_as_battery_aux_salt(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, sector="其他锂盐", name="SMM: NMP - 平均价: 日度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output, add_sort_reason_column=True)
    reason = _read_first_reason(output)
    assert "NMP" in reason
    assert "电池辅材锂盐" not in reason


def test_solid_electrolyte_not_labeled_as_electrolyte_additive(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, sector="电解液产业链", name="SMM: LPSC固态电解质 - 平均价: 日度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output, add_sort_reason_column=True)
    reason = _read_first_reason(output)
    assert "LPSC" in reason or "固态电解质" in reason
    assert "电解液原料添加剂" not in reason


def test_parallel_explanation_does_not_claim_upstream(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, sector="负极材料", name="SMM: 天然石墨 - 平均价: 日度"),
        make_row(2, sector="负极材料", name="SMM: 人造石墨 - 平均价: 日度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output, add_sort_reason_column=True)
    reasons = _read_all_reasons(output)
    assert all("平行" in r for r in reasons)
    assert all("不存在直接上下游关系" in r for r in reasons)


def _read_first_reason(path) -> str:
    reasons = _read_all_reasons(path)
    return reasons[0] if reasons else ""


def _read_all_reasons(path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    ws = workbook["指标目录"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(v or "") for v in rows[0]]
    reason_idx = header.index("排序说明")
    reasons = [
        str(r[reason_idx] or "")
        for r in rows[1:]
        if r[4] not in (None, "", "Indicator Name")
    ]
    workbook.close()
    return reasons


def test_natural_graphite_chain_order(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, sector="负极材料", name="SMM: 天然石墨（中端）-平均价: 日度"),
        make_row(2, sector="负极材料", name="SMM: 球形石墨（中国）（99.95%: 15-20μm） - 平均价: 日度"),
        make_row(3, sector="负极材料", name="SMM: 鳞片石墨（内蒙古）（-194） - 平均价: 日度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)
    names = data_names(read_catalog(output))
    assert "鳞片石墨" in names[0] and "球形石墨" in names[1] and "天然石墨" in names[2]


def test_artificial_graphite_chain_order(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, sector="负极材料", name="SMM: 人造石墨 - 平均价: 日度"),
        make_row(2, sector="负极材料", name="SMM: 负极石墨化（箱体式炉轻料） - 平均价: 日度"),
        make_row(3, sector="负极材料", name="SMM: 石油焦 - 平均价: 日度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)
    names = data_names(read_catalog(output))
    assert "石油焦" in names[0] and "石墨化" in names[1] and "人造石墨" in names[2]


def test_electrolyte_inputs_are_parallel(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, sector="电解液产业链", name="SMM: DMC - 平均价: 日度"),
        make_row(2, sector="电解液产业链", name="SMM: 六氟磷酸锂 - 平均价: 日度"),
        make_row(3, sector="电解液产业链", name="SMM: FEC - 平均价: 日度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output, add_sort_reason_column=True)
    reasons = _read_all_reasons(output)
    assert all("平行投入分支" in r and "不互为上下游" in r for r in reasons)


def test_cell_before_pack(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "周度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, sector="电池电芯", name="SMM: Pack 电池组 - 平均价: 周度"),
        make_row(2, sector="电池电芯", name="SMM: 方形磷酸铁锂电芯（174Ah）- 平均价: 周度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)
    names = data_names(read_catalog(output))
    assert "电芯" in names[0] and "Pack" in names[1]


def test_lithium_concentrate_family_order(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "周度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, sector="锂矿", name="SMM: 锂辉石精矿(CIF中国) -平均价: 周度"),
        make_row(2, sector="锂矿", name="SMM: 津巴布韦锂辉石精矿 6%（CIF中国）现货 -平均价: 周度"),
        make_row(3, sector="锂矿", name="SMM: 津巴布韦锂辉石精矿 5%（CIF中国）现货 -平均价: 周度"),
        make_row(4, sector="锂矿", name="SMM: 津巴布韦锂辉石精矿 5.5%（CIF中国）现货 -平均价: 周度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)
    names = data_names(read_catalog(output))
    assert "锂辉石精矿(CIF中国)" in names[0]
    assert "5%" in names[1] and "5.5%" in names[2] and "6%" in names[3]


def test_product_family_does_not_merge_different_minerals() -> None:
    assert _product_family_key("SMM: 津巴布韦锂辉石精矿 5%（CIF中国）现货 -平均价: 周度") == "锂辉石精矿"
    assert _product_family_key("SMM: 江西锂云母精矿 2.5%（中国现货） -平均价: 周度") == "锂云母精矿"


def _composite_tag(ctype: str, ckey: str, role: str, status: str = "MATCHED") -> str:
    return json.dumps(
        {
            "composite": {
                "composite_type": ctype,
                "composite_key": ckey,
                "composite_role": role,
                "composite_status": status,
            }
        },
        ensure_ascii=False,
    )


def test_market_activity_composite_order(tmp_path) -> None:
    rows = [
        HEADER,
        ["M-sheet", "", "日度", "", "", "", "", "价格", "其他", "", "否", "", ""],
        make_row(1, name="GFEX: 碳酸锂: 主力合约: 持仓量: 日度", sub="持仓", tags=_composite_tag("market_activity", "K1", "open_interest")),
        make_row(2, name="GFEX: 碳酸锂: 主力合约: 成交持仓比: 日度", sub="成交持仓比", tags=_composite_tag("market_activity", "K1", "oi_volume_ratio")),
        make_row(3, name="GFEX: 碳酸锂: 主力合约: 成交量: 日度", sub="成交量", tags=_composite_tag("market_activity", "K1", "volume")),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    names = data_names(read_catalog(output))
    assert "成交量" in names[0] and "持仓量" in names[1] and "成交持仓比" in names[2]


def test_trade_composite_order(tmp_path) -> None:
    rows = [
        HEADER,
        ["T-sheet", "", "月度", "", "", "", "", "进出口", "其他", "", "否", "", ""],
        make_row(1, name="中国海关: 碳酸锂出口量: 总计: 月度", sub="出口", tags=_composite_tag("trade", "T1", "export")),
        make_row(2, name="中国海关: 碳酸锂净出口: 总计: 月度", sub="净出口", tags=_composite_tag("trade", "T1", "net_export")),
        make_row(3, name="中国海关: 碳酸锂进口量: 总计: 月度", sub="进口", tags=_composite_tag("trade", "T1", "import")),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    names = data_names(read_catalog(output))
    assert "进口量" in names[0] and "出口量" in names[1] and "净出口" in names[2]


def test_balance_composite_five_group_order(tmp_path) -> None:
    rows = [
        HEADER,
        ["B-sheet", "", "月度", "", "", "", "", "平衡", "平衡", "", "否", "", ""],
        make_row(1, name="SMM: 锂矿供需平衡: 出口量: 月度", tags=_composite_tag("balance", "B1", "export")),
        make_row(2, name="SMM: 锂矿供需平衡: 平衡: 月度", tags=_composite_tag("balance", "B1", "balance")),
        make_row(3, name="SMM: 锂矿供需平衡: 产量: 月度", tags=_composite_tag("balance", "B1", "production")),
        make_row(4, name="SMM: 锂矿供需平衡: 需求量: 月度", tags=_composite_tag("balance", "B1", "demand")),
        make_row(5, name="SMM: 锂矿供需平衡: 进口量: 月度", tags=_composite_tag("balance", "B1", "import")),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    names = data_names(read_catalog(output))
    assert "产量" in names[0] and "进口量" in names[1] and "出口量" in names[2]
    assert "需求量" in names[3] and "平衡" in names[4]


def test_composite_actual_and_forecast_do_not_mix(tmp_path) -> None:
    rows = [
        HEADER,
        ["B-sheet", "", "月度", "", "", "", "", "平衡", "平衡", "", "否", "", ""],
        make_row(1, name="SMM: 碳酸锂供需预测: 产量: 月度", tags=_composite_tag("balance", "PRED", "production")),
        make_row(2, name="SMM: 碳酸锂供需: 产量: 月度", tags=_composite_tag("balance", "ACT", "production")),
        make_row(3, name="SMM: 碳酸锂供需预测: 需求: 月度", tags=_composite_tag("balance", "PRED", "demand")),
        make_row(4, name="SMM: 碳酸锂供需: 需求: 月度", tags=_composite_tag("balance", "ACT", "demand")),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    names = data_names(read_catalog(output))
    assert names.index("SMM: 碳酸锂供需: 产量: 月度") < names.index("SMM: 碳酸锂供需预测: 产量: 月度")
    assert names.index("SMM: 碳酸锂供需: 需求: 月度") < names.index("SMM: 碳酸锂供需预测: 需求: 月度")


def test_composite_contiguity_violation_detected(tmp_path) -> None:
    rows = [
        HEADER,
        ["M-sheet", "", "日度", "", "", "", "", "价格", "其他", "", "否", "", ""],
        make_row(1, name="GFEX: 碳酸锂: 主力合约: 成交量: 日度", sub="成交量", tags=_composite_tag("market_activity", "K1", "volume")),
        make_row(2, name="GFEX: 碳酸锂: 主力合约: 持仓量: 日度", sub="持仓", tags=_composite_tag("market_activity", "K1", "open_interest")),
        make_row(3, name="GFEX: 碳酸锂: 主力合约: 成交持仓比: 日度", sub="成交持仓比", tags=_composite_tag("market_activity", "K1", "oi_volume_ratio")),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    names = data_names(read_catalog(output))
    assert "成交量" in names[0] and "持仓量" in names[1] and "成交持仓比" in names[2]


def test_composite_sort_reason(tmp_path) -> None:
    rows = [
        HEADER,
        ["T-sheet", "", "月度", "", "", "", "", "进出口", "其他", "", "否", "", ""],
        make_row(1, name="中国海关: 碳酸锂进口量: 总计: 月度", sub="进口", tags=_composite_tag("trade", "T1", "import")),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output, add_sort_reason_column=True)

    reasons = _read_all_reasons(output)
    assert any("进出口复合组" in r and "进口→出口→净出口" in r for r in reasons)


def test_removes_old_merged_ranges_before_rewriting(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        ["", "", "", "", "", "", "", "其他", "其他", "", "否", "", ""],
        ["", "", "", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, name="B 指标"),
        make_row(2, name="A 指标"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    workbook = load_workbook(source)
    ws = workbook["指标目录"]
    ws.merge_cells("A3:G3")
    ws.merge_cells("A4:G4")
    workbook.save(source)
    workbook.close()

    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output)

    workbook = load_workbook(output)
    ws = workbook["指标目录"]
    assert len(ws.merged_cells.ranges) == 0
    names = [row[4].value for row in ws.iter_rows(min_row=2) if row[4].value]
    assert "A 指标" in names
    assert "B 指标" in names
    workbook.close()


def test_adds_normalized_name_column_after_indicator_name(tmp_path) -> None:
    rows = [
        HEADER,
        ["S-sheet", "", "日度", "", "", "", "", "其他", "其他", "", "否", "", ""],
        make_row(1, name="电池级碳酸锂 - 平均价: 日度"),
        make_row(2, name="电池级碳酸锂 -平均价: 周度"),
    ]
    source = write_workbook(tmp_path / "input.xlsx", rows)
    output = tmp_path / "output.xlsx"
    sort_catalog_workbook(source, output, add_normalized_column=True)

    catalog = read_catalog(output)
    assert catalog[0][4] == "Indicator Name"
    assert catalog[0][5] == "指标名称归一化"
    assert catalog[0][6] == "Unit"
    for row in catalog[1:]:
        if row[4] in (None, "", "S-sheet"):
            continue
        assert row[5] == normalize_indicator_title(row[4])
