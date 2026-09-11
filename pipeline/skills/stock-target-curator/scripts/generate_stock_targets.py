# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = ["序号", "产业链", "细分板块", "所属环节", "股票名称", "股票代码", "总市值(亿)", "AI筛选原因"]


def load_config(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def build_rows(config: dict, industry: str) -> list[dict]:
    industry_meta = config["industries"].get(industry)
    if not industry_meta:
        raise ValueError(f"Unsupported industry: {industry}")

    rows: list[dict] = []
    seen: set[tuple[str, str]] = set()
    for stock in config.get("stocks", []):
        if stock.get("industry") != industry:
            continue
        segment = str(stock.get("segment") or "").strip()
        code = str(stock.get("code") or "").strip()
        key = (segment, code)
        if not segment or not code or key in seen:
            continue
        seen.add(key)
        try:
            cap = float(str(stock.get("market_cap") or 0) or 0)
        except (TypeError, ValueError):
            cap = 0.0
        rows.append(
            {
                "industry": industry_meta["name"],
                "segment": segment,
                "stage": str(stock.get("stage") or "").strip(),
                "name": str(stock.get("name") or "").strip(),
                "code": code,
                "cap": cap,
                "reason": str(stock.get("note") or "").strip(),
            }
        )

    # 板块内按总市值从高到低排序（市值相同按代码升序，保持稳定）
    rows.sort(key=lambda row: (row["segment"], -row["cap"], row["code"]))
    return rows


def check_segment_limits(config: dict, industry: str, rows: list[dict]) -> list[str]:
    requirements = config.get("requirements", {}).get(industry, {})
    min_per_segment = int(requirements.get("min_per_segment", 0))
    max_per_segment = requirements.get("max_per_segment")
    if max_per_segment is not None:
        max_per_segment = int(max_per_segment)

    counts = Counter(row["segment"] for row in rows)
    segments = config["industries"][industry]["segments"]
    warnings: list[str] = []
    for segment in segments:
        count = counts.get(segment, 0)
        if count < min_per_segment:
            warnings.append(f"{segment}: {count} < min {min_per_segment}")
        if max_per_segment is not None and count > max_per_segment:
            warnings.append(f"{segment}: {count} > max {max_per_segment}")
    return warnings


def write_xlsx(rows: list[dict], output: Path) -> Path:
    output = output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "个股标的"

    header_fill = PatternFill("solid", fgColor="D9E2F3")
    header_font = Font(bold=True)
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, row in enumerate(rows, 1):
        values = [
            index,
            row["industry"],
            row["segment"],
            row["stage"],
            row["name"],
            row["code"],
            row["cap"],
            row["reason"],
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=index + 1, column=col, value=value)

    widths = [8, 14, 18, 18, 18, 14, 12, 30]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(output)
    return output


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Generate stock targets for an industry chain.")
    parser.add_argument("--industry", choices=["silicon", "tin", "lithium"], required=True)
    parser.add_argument(
        "--config",
        default=Path(__file__).resolve().parents[1] / "config" / "stock_targets.json",
    )
    parser.add_argument("--output", required=True)
    parser.add_argument("--strict", action=argparse.BooleanOptionalAction, default=True)
    args = parser.parse_args(argv)

    config = load_config(Path(args.config))
    rows = build_rows(config, args.industry)
    if not rows:
        print(f"No stock targets found for {args.industry}", file=__import__("sys").stderr)
        return 2

    warnings = check_segment_limits(config, args.industry, rows)
    for warning in warnings:
        print(f"Segment limit warning: {warning}", file=__import__("sys").stderr)
    if warnings and args.strict:
        return 2

    output = write_xlsx(rows, Path(args.output))
    workbook = load_workbook(output, read_only=True)
    ws = workbook["个股标的"]
    print(f"Industry={args.industry} stocks={len(rows)} output={output}")
    print(f"Sheet={ws.title} rows={ws.max_row - 1}")
    workbook.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
